# -*- coding: utf-8 -*-
"""天津桥梁定额自动匹配（纯逻辑模块，不依赖 Streamlit）。

- parse_quota_library: 解析定额库 —— 导定额.xlsx（单表 38 页分页计价表，983 条定额
  + 材料行，含分页断行合并/名称修补）；兼容旧 天津桥梁.xlsx（sheet1 定额库 + sheet2
  套定额规则样例）。
- parse_rule_template: 解析 单位工程.xlsx 套定额模板（8 页分页计价表：
  分部→分项→子组→明细行，特殊编号 1-6-71 换/Z01308@1/借2-33）。
- parse_quantity_excel: 解析海泰式数量总表（分项材料表：分项名/适用位置/备注/数量/材料列）。
- match_quotas: 规则驱动（分项→子分部→样例组合）+ 文本打分回退。
- ai_explain: DeepSeek 生成多候选的匹配理由/差异说明（仅解释，不改变入选结果）。
- build_result_excel: 输出结果工作簿。
"""
import io
import json
import re
import unicodedata
from pathlib import Path

# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------


class BridgeMatchError(Exception):
    """定额库/数量表格式不符合预期。"""


_UNIT_MAP = {
    "m3": "m3", "立方": "m3", "方": "m3", "m³": "m3",
    "m2": "m2", "平米": "m2", "平方米": "m2", "m²": "m2",
    "米": "m", "延米": "m", "m": "m",
    "吨": "t", "t": "t", "kg": "kg", "千克": "kg",
    "套": "套", "个": "个", "cm3": "cm3",
}
_COUNT_UNITS = {"套", "个", "只", "根", "块", "座", "处", "孔", "跨", "条", "道", "对"}


def _norm_text(s):
    """全角→半角、粘→黏、φ 变体统一、去全部空白。"""
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = s.replace("粘", "黏").replace("ф", "φ").replace("Φ", "φ").replace("ϕ", "φ")
    s = re.sub(r"\s+", "", s)
    return s


def _norm_code(s):
    s = _norm_text(s).lower()
    s = re.sub(r"[—－]", "-", s)
    return s


def _unit_parts(unit):
    """定额单位 → (基类, 扩大倍数)："10m"→("m",10)、"100m2"→("m2",100)、"t"→("t",1)。"""
    u = _norm_text(unit).lower()
    u = u.replace("³", "3").replace("²", "2")
    m = re.match(r"^(\d+(?:\.\d+)?)\s*(.*)$", u)
    if m:
        return _UNIT_MAP.get(m.group(2), m.group(2)), float(m.group(1))
    return _UNIT_MAP.get(u, u), 1.0


# ---------------------------------------------------------------------------
# 解析：定额库（导定额.xlsx 单表分页 / 天津桥梁.xlsx 双表）+ 套定额模板（单位工程.xlsx）
# ---------------------------------------------------------------------------

_HEADER_LABELS = {"序号": "seq", "编号": "code", "定额编号": "code", "项目名称": "name",
                  "名称": "name", "单位": "unit", "单价": "price", "工程量": "qty"}

_STD_CODE = re.compile(r"^\d+-\d+-\d+$")  # 标准定额子目编号
_Z_CODE = re.compile(r"^Z\d+")            # 材料行编号（Z 开头，不参与匹配）

# 名称截断修补：分页断行合并后按编号整名替换（依据旧库权威全名；
# 其中 1-6-44/3-6-17 断行处为汉字接汉字，合并丢空格，按旧库补回）
_LIB_NAME_PATCH = {
    "1-6-44": "大型机械场外运输 履带式拖拉机 75kW以内",
    "3-1-44": "冲击式钻机钻孔 φ≤1500 h≤40m 黏土",
    "3-3-27": "现浇混凝土墩(台)身 C30轻型桥台",
    "3-3-52": "现浇混凝土小型构件 C50混凝土垫块",
    "3-5-28": "金属波纹管安装 直径(mm) φ84",
    "3-6-17": "安装混凝土板梁 安装板梁 水上 L≤25m",
    "3-7-53": "现浇上部混凝土模板 立柱、端柱、灯柱",
}
# 断行价格修补：3-3-52 页断 '7853'+'5' 只能合并出 7853.5，
# 但源表真实值为 7853.05（PDF 转表丢 '.0'），按旧库权威值修正
_LIB_PRICE_PATCH = {"3-3-52": 7853.05}

# 单位工程.xlsx 模板结构词汇：分部/分项/子组（无编号行）
_TPL_DIVS = ("下部结构", "上部结构", "附属结构")
_TPL_SUBS = ("灌注桩", "承台", "支撑梁", "系梁", "桥台", "台帽", "台盖梁", "桥墩",
             "墩帽", "墩盖梁", "现浇箱梁", "预制小箱梁", "预制板梁", "钢箱梁",
             "支座", "伸缩缝", "人行道板", "防撞护栏", "隔声屏障", "搭板", "泄水孔",
             "沥青混凝土铺装", "防水混凝土铺装", "防水层", "土方", "护坡")
_TPL_SUBGROUPS = ("成孔", "混凝土", "钢筋", "施工措施", "钢绞线",
                  "板式橡胶支座", "四氟板式橡胶支座", "细粒式沥青混凝土",
                  "中粒式沥青混凝土", "粘层")
# 模板断行合并后的名称修补（3-2-5 页断合并丢空格）
_TPL_NAME_FIX = {
    "非预应力钢筋制作与安装螺纹钢 φ14以外": "非预应力钢筋制作与安装 螺纹钢 φ14以外",
}


def _find_columns(cells, labels):
    """按表头标签找列位置 → {字段: 列号}。"""
    found = {}
    for j, c in enumerate(cells):
        key = _HEADER_LABELS.get(_norm_text(c))
        if key and key not in found:
            found[key] = j
    return found


def _to_float(v):
    try:
        return float(str(v).replace(",", "").replace("元", "").strip())
    except (TypeError, ValueError):
        return None


def _is_cjk(ch):
    return "一" <= ch <= "鿿"


def _join_name(a, b):
    """断行名称合并：')' 后、数字/汉字后接字母数字时补空格，其余直接拼接。"""
    if not a:
        return b
    if not b:
        return a
    pa, fb = a[-1], b[0]
    if pa == ")":
        return a + " " + b
    if (pa.isdigit() or _is_cjk(pa)) and fb.isascii() and fb.isalnum():
        return a + " " + b
    return a + b


def _merge_fragment(last, code, name, unit, price_str):
    """分页断行的下半段合并进上一行（名称/编号/单位/单价拼接）。

    编号：仅在拼接结果是合法 STD/Z 编号或上一行以 '-' 结尾时合并（孤儿数字丢弃）；
    单价：上一段含 '.' 直接拼接（'4745.4'+'6'），否则补 '.'（'7853'+'5'→7853.5）。
    """
    last["name"] = _join_name(last["name"], name)
    if code:
        combined = last["code"] + code
        if _STD_CODE.match(combined) or _Z_CODE.match(combined) or last["code"].endswith("-"):
            last["code"] = combined
    if unit:
        last["unit"] += unit
    if price_str:
        p = last["price_str"]
        last["price_str"] = p + price_str if p and "." in p else p + "." + price_str


_PAGE_SKIP_KEYS = ("小计", "合计", "计价表", "金额单位", "专业工程", "表6")


def _iter_paged_rows(rows, structure_names=frozenset()):
    """分页计价表行流：跳过表头/页脚/小计；断行合并到上一内容行。

    yield {"kind": "detail"|"material"|"structure", "seq","code","name","unit",
           "price_str","row"}。无编号名称行：名称在 structure_names → 结构行；
    否则视为上一页末行被页边切开的下半段，合并进上一行。
    """
    last = None
    for i, r in enumerate(rows):
        vals = [str(x).strip() if x is not None else "" for x in r]
        c0, c1, c2 = vals[0], vals[1], vals[2]
        c4 = vals[4] if len(vals) > 4 else ""
        c6 = vals[6] if len(vals) > 6 else ""
        if not c2:
            continue
        if any(k in c2 for k in _PAGE_SKIP_KEYS):
            continue
        if c0 == "序号" or c1 == "编号":
            continue
        first = c1.split()[0] if c1 else ""
        code_ok = bool(_STD_CODE.match(first) or c1.startswith(("Z", "借")))
        if code_ok:
            if last is not None:
                yield last
            last = {"seq": c0, "code": c1, "name": c2, "unit": c4,
                    "price_str": c6, "row": i,
                    "kind": "material" if c1.startswith("Z") else "detail"}
        elif structure_names and c2 in structure_names:
            if last is not None:
                yield last
            last = None
            yield {"kind": "structure", "name": c2, "row": i}
        elif last is not None:
            _merge_fragment(last, c1, c2, c4, c6)
    if last is not None:
        yield last


def _parse_lib_rows_old(rows, hdr):
    """旧库（天津桥梁.xlsx sheet1）明细行 → (entries, counts, materials)。"""
    cols = None
    for i, r in enumerate(rows[:15]):
        c = _find_columns(r, _HEADER_LABELS)
        if len(c) >= 3 and "name" in c:
            cols, hdr = c, i
            break
    if cols is None:
        raise BridgeMatchError("未在定额库中识别到表头（需含 编号/项目名称/单位/单价 列）。")
    # 两行表头（表7(1) 版：第二行是 单价/合价）：把缺的列从下一行补上
    if "price" not in cols and hdr + 1 < len(rows):
        c2 = _find_columns(rows[hdr + 1], _HEADER_LABELS)
        for k, v in c2.items():
            if k not in cols:
                cols[k] = v
    entries, counts = [], {}
    for i, r in enumerate(rows[hdr + 1:], start=hdr + 2):
        name = str(r[cols["name"]]).strip() if cols.get("name") is not None and r[cols["name"]] is not None else ""
        code = str(r[cols["code"]]).strip() if cols.get("code") is not None and r[cols["code"]] is not None else ""
        unit = str(r[cols["unit"]]).strip() if cols.get("unit") is not None and r[cols["unit"]] is not None else ""
        price = _to_float(r[cols["price"]]) if cols.get("price") is not None else None
        if not name:
            continue
        if any(k in name for k in ("合计", "小计", "总计")):
            continue
        if not _STD_CODE.match(code):
            continue  # 章节行（3-1）、Z 开头材料明细行不参与匹配
        entries.append({"code": _norm_code(code), "name": name, "unit": unit,
                        "price": price, "row": i})
        if code:
            counts[code] = counts.get(code, 0) + 1
    return entries, counts, []


def _parse_rule_rows_old(rows):
    """旧库 sheet2（自然单位含分部小计）→ 套定额规则样例。"""
    rcols, rheader = None, None
    for i, r in enumerate(rows[:10]):
        c = _find_columns(r, _HEADER_LABELS)
        if "name" in c and ("code" in c or "unit" in c):
            rcols, rheader = c, i
            break
    if rcols is None:
        return {}
    if "price" not in rcols and rheader + 1 < len(rows):
        c2 = _find_columns(rows[rheader + 1], _HEADER_LABELS)
        for k, v in c2.items():
            if k not in rcols:
                rcols[k] = v
    rules = {}
    parent_set = {"下部结构", "上部结构", "桥面系", "其他"}
    stack, group = [], None
    for i, r in enumerate(rows[rheader + 1:], start=rheader + 2):
        name = str(r[rcols["name"]]).strip() if rcols.get("name") is not None and r[rcols["name"]] is not None else ""
        code = str(r[rcols["code"]]).strip() if rcols.get("code") is not None and r[rcols["code"]] is not None else ""
        unit = str(r[rcols["unit"]]).strip() if rcols.get("unit") is not None and r[rcols["unit"]] is not None else ""
        if not name:
            continue
        if any(k in name for k in ("合计", "小计", "总计")):
            continue
        if not code:
            # 分部/子分部行
            nm = _norm_text(name)
            if nm in parent_set:
                stack, group = [name], name
            elif stack and stack[-1] in parent_set:
                stack.append(name)
                group = name
            else:
                stack, group = [name], name
            continue
        # 明细行（带编号）
        if not code[0].isdigit():
            continue
        qty = _to_float(r[rcols["qty"]]) if rcols.get("qty") is not None else None
        rules.setdefault(group, []).append(
            {"code": _norm_code(code), "name": name, "unit": unit, "qty": qty})
    return rules


def _parse_lib_rows_new(rows):
    """导定额.xlsx 单表分页 → (entries, counts, materials)。"""
    entries, counts, materials = [], {}, []
    for row in _iter_paged_rows(rows):
        if row["kind"] == "structure":
            continue  # 库表无结构行
        name, code, unit = row["name"], _norm_code(row["code"]), row["unit"]
        price = _to_float(row["price_str"]) if row["price_str"] else None
        if code in _LIB_NAME_PATCH:
            name = _LIB_NAME_PATCH[code]
        if code in _LIB_PRICE_PATCH:
            price = _LIB_PRICE_PATCH[code]
        if row["kind"] == "material":
            materials.append({"code": code, "name": name, "unit": unit,
                              "price": price, "row": row["row"]})
            continue
        if not _STD_CODE.match(code):
            continue  # 借/换 等特殊编号不进入库条目
        entries.append({"code": code, "name": name, "unit": unit,
                        "price": price, "row": row["row"]})
        counts[code] = counts.get(code, 0) + 1
    return entries, counts, materials


def parse_quota_library(data: bytes) -> dict:
    """解析定额库 Excel。

    - 导定额.xlsx（新库）：单表分页计价表 → 983 条定额 + 材料行；
      分页断行合并、名称/单价修补（依据旧库权威值）。
    - 天津桥梁.xlsx（旧库，向后兼容）：sheet1 定额库 + sheet2 套定额规则样例。
    返回 {"entries": [{code,name,unit,price,row}], "rules": {子分部名: [定额行...]},
          "materials": [{code,name,unit,price,row}], "duplicates": {编号: 条数}}
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise BridgeMatchError(f"无法读取定额库文件：{e}") from e

    sheets = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        hdr = None
        for i, r in enumerate(rows[:15]):
            c = _find_columns(r, _HEADER_LABELS)
            if len(c) >= 3 and "name" in c:
                hdr = i
                break
        t = _norm_text(ws.title)
        sheets.append({"title": t, "rows": rows, "hdr": hdr,
                       "is_rule": "自然单位" in t or "分部小计" in t})

    rule_sheet = next((s for s in sheets if s["is_rule"]), None)
    if rule_sheet is not None:
        # 旧格式：sheet1 定额库 + sheet2 规则样例
        lib_sheet = next((s for s in sheets
                          if not s["is_rule"] and s["hdr"] is not None), sheets[0])
        entries, counts, materials = _parse_lib_rows_old(lib_sheet["rows"], lib_sheet["hdr"])
        rules = _parse_rule_rows_old(rule_sheet["rows"])
    else:
        # 新格式：导定额.xlsx 单表多页
        src = next((s for s in sheets if s["hdr"] is not None),
                   sheets[0] if sheets else None)
        if src is None:
            raise BridgeMatchError("未在定额库中识别到表头（需含 编号/项目名称/单位/单价 列）。")
        entries, counts, materials = _parse_lib_rows_new(src["rows"])
        rules = {}
    if len(entries) < 50:
        raise BridgeMatchError(f"定额库条目过少（{len(entries)} 条），请确认选择的是天津桥梁定额库。")
    return {"entries": entries, "rules": rules, "materials": materials,
            "duplicates": {k: v for k, v in counts.items() if v > 1}}


def parse_rule_template(data: bytes) -> dict:
    """解析 单位工程.xlsx 套定额模板（8 页分页计价表）。

    结构：分部（下部结构/上部结构/附属结构）→ 分项（灌注桩/承台/…/护坡）
    → 子组（成孔/混凝土/钢筋/施工措施/钢绞线…）→ 明细行。
    返回 {"rules": {规则组名: [{"code","name","unit","qty","huan"}...]}, "subs": [分项名...],
          "divs": [{"div": 分部名, "groups": [规则组名...]}]}（分项按模板页序归到分部，
          供导出计价表对齐模板分部顺序——土方/护坡在附属结构末尾）。
    特殊编号：'1-6-71 换'→剥换后缀记 huan 标志；'Z01308@1' 材料行原样保留；
    '借2-33' 借用定额（库内无此子目）原样保留。
    """
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise BridgeMatchError(f"无法读取套定额模板文件：{e}") from e
    rows = None
    for ws in wb.worksheets:
        for i, r in enumerate(list(ws.iter_rows(values_only=True))[:15]):
            c = _find_columns(r, _HEADER_LABELS)
            if len(c) >= 3 and "name" in c:
                rows = list(ws.iter_rows(values_only=True))
                break
        if rows is not None:
            break
    if rows is None:
        raise BridgeMatchError("未在套定额模板中识别到表头（需含 编号/项目名称/单位 列）。")

    structure_names = frozenset(_TPL_DIVS) | frozenset(_TPL_SUBS) | frozenset(_TPL_SUBGROUPS)
    rules, subs, divs = {}, [], []
    sub = group = div = subgrp = None
    for row in _iter_paged_rows(rows, structure_names):
        if row["kind"] == "structure":
            nm = row["name"]
            if nm in _TPL_DIVS:
                sub = subgrp = None  # 分部行（下部结构/上部结构/附属结构）
                div = {"div": nm, "groups": []}
                divs.append(div)
            elif nm in _TPL_SUBS and nm != sub:
                sub, group, subgrp = nm, _map_group(nm), None
                subs.append(nm)
                if group and div is not None and group not in div["groups"]:
                    div["groups"].append(group)
            elif nm in _TPL_SUBGROUPS:
                subgrp = nm  # 子组行（成孔/混凝土/钢筋/施工措施/钢绞线…）
            continue  # 子组行不产出条目
        name = _TPL_NAME_FIX.get(row["name"], row["name"]).replace("\x7f", "")
        code = row["code"]
        huan = "换" in code
        rec = {"code": code.split()[0] if huan else code, "name": name,
               "unit": row["unit"], "qty": None, "row": row["row"],
               "subgroup": subgrp}
        if huan:
            rec["huan"] = True
        rules.setdefault(group or "其他", []).append(rec)
    # 单位工程.xlsx 末页「土方」自成一部（位于附属结构之后）——把解析时归到
    # 附属结构末尾的 土方/护坡 组提升为独立「土方」分部（上传的自定义模板同此处理）
    if divs and divs[-1]["div"] == "附属结构":
        gs = divs[-1]["groups"]
        lift = [g for g in gs if g in ("土方", "护坡")]
        if lift:
            divs[-1]["groups"] = [g for g in gs if g not in lift]
            divs.append({"div": "土方", "groups": lift})
    return {"rules": rules, "subs": subs, "divs": divs}


# 内置单位工程.xlsx 模板的完整层级（分部→分项→子分项→子组），
# 供导出计价表严格对齐模板排布。土方是第 4 个分部（位于最后）；
# 桥台/桥墩为父分项，下分 桥台/台帽/台盖梁、桥墩/墩帽/墩盖梁 子分项，
# 每个子分项下按 成孔/混凝土/钢筋/施工措施/钢绞线 等子组分明细。
_BUILTIN_TPL_HIERARCHY = [
    {"div": "下部结构", "sections": [
        {"name": "灌注桩", "group": "桩基础", "subs": ["成孔", "混凝土", "钢筋"]},
        {"name": "承台", "group": "承台", "subs": ["混凝土", "钢筋"]},
        {"name": "支撑梁", "group": "支撑梁", "subs": ["混凝土", "钢筋"]},
        {"name": "系梁", "group": "系梁", "subs": ["混凝土", "钢筋"]},
        {"name": "桥台", "children": [
            {"name": "桥台", "group": "桥台", "subs": ["混凝土", "钢筋"]},
            {"name": "台帽", "group": "台帽", "subs": ["混凝土", "钢筋"]},
            {"name": "台盖梁", "group": "台盖梁", "subs": ["混凝土", "钢筋"]}]},
        {"name": "桥墩", "children": [
            {"name": "桥墩", "group": "桥墩", "subs": ["混凝土", "钢筋"]},
            {"name": "墩帽", "group": "墩帽", "subs": ["混凝土", "钢筋"]},
            {"name": "墩盖梁", "group": "墩盖梁", "subs": ["混凝土", "钢筋"]}]},
    ]},
    {"div": "上部结构", "sections": [
        {"name": "现浇箱梁", "group": "预应力砼现浇箱梁",
         "subs": ["混凝土", "施工措施", "钢筋", "钢绞线"]},
        {"name": "预制小箱梁", "group": "预应力混凝土小箱梁",
         "subs": ["混凝土", "钢筋", "钢绞线"]},
        {"name": "预制板梁", "group": "预应力板梁",
         "subs": ["混凝土", "钢筋", "钢绞线"]},
        {"name": "钢箱梁", "group": "钢箱梁", "subs": []},
        # 梯道并入上部结构（用户 2026-10-27 口径）：与主梁同构，分 混凝土/钢筋 子组
        {"name": "梯道", "group": "梯道", "subs": ["混凝土", "钢筋"]},
    ]},
    {"div": "附属结构", "sections": [
        {"name": "支座", "group": "支座", "subs": ["板式橡胶支座", "四氟板式橡胶支座"]},
        {"name": "伸缩缝", "group": "伸缩缝", "subs": []},
        {"name": "人行道板", "group": "人行道板", "subs": ["混凝土", "钢筋"]},
        {"name": "防撞护栏", "group": "防撞护栏", "subs": ["混凝土", "钢筋"]},
        {"name": "隔声屏障", "group": "隔声屏障", "subs": []},
        {"name": "搭板", "group": "搭板", "subs": ["混凝土", "钢筋"]},
        {"name": "泄水孔", "group": "排水", "subs": []},
        {"name": "沥青混凝土铺装", "group": "沥青铺装",
         "subs": ["细粒式沥青混凝土", "中粒式沥青混凝土", "粘层"]},
        {"name": "防水混凝土铺装", "group": "防水铺装", "subs": ["混凝土", "钢筋"]},
        {"name": "防水层", "group": "防水层", "subs": []},
    ]},
    {"div": "土方", "sections": [
        {"name": "土方", "group": "土方", "subs": []},
        {"name": "护坡", "group": "护坡", "subs": []},
    ]},
]

_BUILTIN_TPL_PATH = Path(__file__).resolve().parent / "quota_base" / "单位工程.xlsx"
_builtin_tpl_cache = None


def load_builtin_rule_template():
    """内置套定额模板规则（quota_base/单位工程.xlsx，与页面 ② 上传的模板同格式）。

    进程内缓存解析结果；divs 用完整层级常量（分部→分项→子分项→子组）。
    文件缺失时抛 BridgeMatchError（页面提示重新放置文件）。
    """
    global _builtin_tpl_cache
    if _builtin_tpl_cache is None:
        if not _BUILTIN_TPL_PATH.exists():
            raise BridgeMatchError(f"内置套定额模板文件缺失：{_BUILTIN_TPL_PATH}")
        data = parse_rule_template(_BUILTIN_TPL_PATH.read_bytes())
        data["divs"] = _BUILTIN_TPL_HIERARCHY
        _builtin_tpl_cache = data
    return _builtin_tpl_cache


# ---------------------------------------------------------------------------
# 解析：海泰式数量总表（分项材料表）
# ---------------------------------------------------------------------------

_SKIP_SHEET_KEYS = ("使用说明", "工程数量表", "材料统计", "指标信息")

# 工程数量表里不分部位的材料行（值集中在首列，直接以材料名作为分项）
_MATRIX_STANDALONE = {"挖方", "填方", "试桩", "锥坡", "支座", "泄水孔", "伸缩缝",
                      "超声波检测管", "其他", "桩基类型"}

# 汇总矩阵类总表 sheet 名关键词（合计/汇总/总计/工程数量表/总表 → 权威矩阵，优先解析）
_MATRIX_SHEET_KEYS = ("合计", "汇总", "总计", "工程数量表", "总表")
# 矩阵中充当「分项」的父材料行（其子行按父名定分项，而非列部位）
_MATRIX_PARENT_FX = {"伸缩缝", "支座", "波纹管", "锚具", "声测管", "锥坡", "桥台锥坡",
                     "挖填方", "试桩", "其他"}
# 标签自带单位后缀的已知单位（支座 GBZYH400x86(CR) 的 (CR) 不是单位，不得剥离）
_UNIT_SUFFIXES = {"m", "m2", "m3", "kg", "t", "个", "套", "根", "米", "㎡", "m³",
                  "处", "L", "l", "块", "方"}


def _looks_like_matrix(rows):
    """表头区域含「工程材料」单元格 → 汇总矩阵式数量表（合计/工程数量表类总表）。"""
    for r in rows[:12]:
        if r and any("工程材料" in _norm_text(str(x)) for x in r if x):
            return True
    return False


def _parse_matrix_sheet(rows):
    """「合计/工程数量表」类汇总矩阵 sheet → 部位×材料矩阵条目
    [{fenxiang, mat, mat_unit, value}]，作为权威工程量来源。

    列头：两级（海泰：构造组+子部位 → 分项名「组+子部位」拼接）或
    三级（大清河：构造组+分组+子部位 → 分项名取最深子部位，用户 2026-10-27 口径），
    「小计/合计」列跳过，构造组行上方「搭板/合计」定搭板列与右边界。
    特殊父材料行（支座/波纹管/锚具/声测管/伸缩缝/挖填方/锥坡/其他…）：
    - 位置列「规格+数量」成对单元格（如 支座行 GBZYH400x86(CR)→64.0）→ 拆成规格条目；
    - 单单元格「名称：数值+单位」对（如 桥台锥坡 C30预制实心六角块：24.8㎡）→ 逐对拆分；
    - 挖方/填方 → 分项=材料=挖方/填方（不分部位）。
    """
    out = []
    hdr = None
    for i, r in enumerate(rows[:12]):
        if r and any("工程材料" in _norm_text(str(x)) for x in r if x):
            hdr = i
            break
    if hdr is None:
        return out
    r6 = rows[hdr]
    r7 = rows[hdr + 1] if hdr + 1 < len(rows) else None
    r8 = rows[hdr + 2] if hdr + 2 < len(rows) else None
    # r8 仅当是子部位表头行（位置列全为文本）才参与最深名取值；
    # 海泰两级列头时 hdr+2 已是数据行 → 退回「组+子部位」拼接
    r8_is_hdr = bool(r8) and not any(
        _to_float(r8[j]) is not None for j in range(4, min(len(r8), 32)))
    # hdr 上方一行为构造组行（上部构造/下部构造/…/搭板/合　计），
    # 它的「合计」列是矩阵右边界——其后是右侧「项/工程量/单价/综合单价」清单块，
    # 列头（盖梁/桩基…）与左侧重复，必须排除；搭板列在 r6/r7/r8 无子头，取该行名称。
    grand = None
    for k in range(max(0, hdr - 3), hdr):
        rk = rows[k]
        if any("合计" in _norm_text(str(x)) or "搭板" in _norm_text(str(x)) for x in rk):
            grand = rk
            break
    cols = []
    cstop = min(len(r6), 32)  # 矩阵右边界 = grand 行的「合计」列（其后为右侧清单块，须排除）
    if grand is not None:
        for j in range(4, min(len(grand), 32)):
            if str(grand[j] or "").strip() and "合计" in _norm_text(str(grand[j])):
                cstop = j
                break
    group = ""
    for j in range(4, cstop):
        gj = str(grand[j] or "").strip() if grand is not None and j < len(grand) else ""
        if gj and "搭板" in _norm_text(gj):
            name = "搭板"
        else:
            a = str(r6[j] or "").strip() if j < len(r6) else ""
            if a:
                group = a  # 合并单元格：组名向后继承
            b = str(r7[j] or "").strip() if r7 is not None and j < len(r7) else ""
            c = str(r8[j] or "").strip() if r8_is_hdr and j < len(r8) else ""
            if r8_is_hdr:
                name = c or b or a  # 三级 → 最深子部位
                name = re.sub(r"[（(][^）)]*[）)]", "", name).strip()
            else:
                name = f"{group} {b}".strip()  # 两级 → 组+子部位（海泰口径）
        name = _norm_text(name)
        if not name or any(k in name for k in ("小计", "合计")):
            continue
        cols.append((j, name))
    parent = ""
    for i in range(hdr + 2, len(rows)):
        r = rows[i]
        if not r:
            continue
        p0 = str(r[1] or "").strip() if len(r) > 1 else ""
        mat = str(r[2] or r[1] or "").strip() if len(r) > 2 else p0
        if p0:
            parent = _norm_text(p0)
        if not mat or any(k in _norm_text(mat) for k in ("小计", "合计")):
            continue
        unit = str(r[3] or "").strip() if len(r) > 3 else ""
        mat_n = _norm_text(mat)
        # 特殊行（规格+数量成对 / 名称：数量描述）：仅三级列头（大清河式）启用；
        # 海泰式两级列头的位置列文本是「共：30个」类汇总备注，由专项解析/明细 sheet 负责，
        # 且右侧「项/工程量/单价」清单块在合计列右界外，不参与扫描
        txt_cells, num_cells = [], []
        for j in range(4, min(len(r), cstop)):
            v = r[j]
            if v in (None, ""):
                continue
            fv = _to_float(v)
            if fv is not None:
                if fv != 0:
                    num_cells.append((j, fv))
            else:
                txt_cells.append((j, str(v).strip()))
        if txt_cells and r8_is_hdr and mat_n != "桩基类型" and "检测管" not in mat_n:
            cells = sorted([(j, "t", s) for j, s in txt_cells]
                           + [(j, "n", v) for j, v in num_cells])
            pairs, pending = [], None
            for _j, kind, val in cells:
                if kind == "t":
                    pending = val
                elif pending is not None:
                    pairs.append((pending, val))
                    pending = None
            if pairs:  # 「规格+数量」成对单元格
                for lab, v in pairs:
                    lab_n = _norm_text(lab)
                    if lab_n in ("挖方", "填方"):  # 分项即材料本身
                        out.append({"fenxiang": lab, "mat": lab,
                                    "mat_unit": unit or "m3", "value": v})
                        continue
                    lab2, lu = lab, unit
                    m = re.match(r"^(.*?)\s*[（(]([^（）()]+)[）)]\s*$", lab)
                    if m and _norm_text(m.group(2)) in _UNIT_SUFFIXES:
                        lab2, lu = m.group(1).strip(), m.group(2)
                    out.append({"fenxiang": parent or mat, "mat": lab2,
                                "mat_unit": _UNIT_MAP.get(_norm_text(lu), lu),
                                "value": v})
                continue
            # 单单元格「名称：数值+单位」对（桥台锥坡/其他/涂装描述行）
            for _j, s in txt_cells:
                for m in re.finditer(
                        r"([^：:，,；;。\s]{1,40})[：:]\s*(\d+(?:\.\d+)?)"
                        r"\s*(㎡|m³|m2|m3|个|套|根|米|处|块|方|[a-zA-Z]{1,4})?",
                        s):
                    out.append({"fenxiang": parent or mat, "mat": m.group(1).strip(),
                                "mat_unit": _UNIT_MAP.get(_norm_text(m.group(3) or ""),
                                                          m.group(3) or ""),
                                "value": float(m.group(2))})
            continue
        # 普通数值矩阵行：部位列 × 材料行交叉值 = 工程量
        fx = None
        if parent and parent in _MATRIX_PARENT_FX:
            fx = parent  # 伸缩缝等父材料行的子行：分项=父名（不分列部位）
        elif mat_n in _MATRIX_STANDALONE:
            fx = mat  # 挖方/填方等不分部位
        for j, col_name in cols:
            if len(r) <= j or r[j] in (None, ""):
                continue
            v = _to_float(r[j])
            if v is None or v == 0:
                continue
            out.append({"fenxiang": fx or col_name, "mat": mat,
                        "mat_unit": unit, "value": v})
    return out


def _looks_like_qinbin(rows):
    """秦滨矩阵式总表：表头行 col0 含「项目」且同行含「单位」列
    （三级列头 构造组/分组/子部位 + 大类/材料/单位三列行，如轻纺互通主线桥总表）。"""
    for r in rows[:8]:
        if not r:
            continue
        if "项目" in _norm_text(str(r[0] or "")):
            if any("单位" in _norm_text(str(x)) for x in r[1:12] if x):
                return True
    return False


def _parse_qinbin_matrix(rows):
    """秦滨矩阵式总表（轻纺互通主线桥式）→ (材料条目, 桩基参数条目)。

    列头三级（行1 构造组：上部构造/下部构造/搭板/合计；行2 分组：预应力箱梁/
    桥面系/桥墩/桥台/小计；行3 子部位：盖梁/墩柱/桩基/小计…）：子部位取最深名，
    下部构造列加「桥墩/桥台」分组前缀消歧（系梁承台、桩基在桥墩/桥台下重复出现），
    小计列与合计列（右边界）排除。
    数据行：col0=大类（块内继承）、col1=材料规格、col2=单位（块内首行，随大类
    换行重置；缺失按材料推断）。col1 为空 → 独立材料行（分项即材料）。
    特殊行：支座/波纹管/锚具等父材料行分项=父名；伸缩缝/排水系统/防抛网/备注行
    单格「名称：数值+单位」冒号/分号对拆分（含 管卡（kg)4208.12 无冒号变体）；
    拆桥/护栏切除/主梁悬臂改造/旧桥铺装拆除/防腐面积行 label 格→下一非空数值格
    配对（合并单元格跨度可跨空列）；桩基 label/value 两行按 4 列组
    （直径/根数/桩长40m以下合计/40m以上合计）转置成桩基参数（桩长=合计m÷根数）。
    """
    out, piles = [], []
    hdr = None
    for i, r in enumerate(rows[:8]):
        if not r:
            continue
        if "项目" in _norm_text(str(r[0] or "")):
            if any("单位" in _norm_text(str(x)) for x in r[1:12] if x):
                hdr = i
                break
    if hdr is None:
        return out, piles
    r0 = rows[hdr]
    r1 = rows[hdr + 1] if hdr + 1 < len(rows) else None
    r2 = rows[hdr + 2] if hdr + 2 < len(rows) else None
    cstop = None  # 合计列 = 矩阵右边界（其后不参与扫描）
    for j in range(3, min(len(r0), 40)):
        if r0[j] and "合计" in _norm_text(str(r0[j])):
            cstop = j
            break
    if cstop is None:
        cstop = min(len(r0), 40)
    cols = []
    group, sub = "", ""
    for j in range(3, cstop):
        a = str(r0[j] or "").strip() if j < len(r0) else ""
        b = str(r1[j] or "").strip() if r1 is not None and j < len(r1) else ""
        c = str(r2[j] or "").strip() if r2 is not None and j < len(r2) else ""
        if a:
            group = a
        if b:
            sub = b
        name = c or b or a
        name = re.sub(r"[（(][^）)]*[）)]", "", name).strip()
        name_n = _norm_text(name)
        if not name_n or any(k in name_n for k in ("小计", "合计")):
            continue
        if _norm_text(group) == "下部构造" and c:
            name = f"{sub} {name}"  # 桥墩 桩基 / 桥台 系梁承台（消歧+嵌套层级）
        cols.append((j, name))
    block, cur_unit = "", ""
    for i in range(hdr + 3, len(rows)):
        r = rows[i]
        if not r:
            continue
        c0 = str(r[0] or "").strip() if len(r) > 0 else ""
        c1 = str(r[1] or "").strip() if len(r) > 1 else ""
        uc = str(r[2] or "").strip() if len(r) > 2 else ""
        if c0:
            block = _norm_text(c0)
            cur_unit = uc  # 单位随大类块首行重置
        elif uc:
            cur_unit = uc
        mat = c1 or c0
        if not mat:
            continue
        mat_n = _norm_text(mat)
        if any(k in mat_n for k in ("小计", "合计")):
            continue
        txt_cells, num_cells = [], []
        for j in range(3, cstop):
            if j >= len(r):
                continue
            v = r[j]
            if v in (None, ""):
                continue
            fv = _to_float(v)
            if fv is not None:
                if fv != 0:
                    num_cells.append((j, "n", fv))
            else:
                txt_cells.append((j, "t", str(v).strip()))
        # 桩基参数转置块：label 行（直径m/根数/桩长40m以下合计m/桩长40m以上合计m
        # 每 4 列循环）后接 value 行（可隔空行）
        if (mat_n == "桩基" and not c1 and txt_cells and not num_cells
                and any("直径" in _norm_text(s) for _, _, s in txt_cells)):
            vrow = None
            for i2 in range(i + 1, min(i + 4, len(rows))):
                rr = rows[i2]
                if not rr:
                    continue
                if any(_to_float(x) is not None and _to_float(x) != 0
                       for x in rr[3:cstop] if x not in (None, "")):
                    vrow = rr
                    break
            if vrow is not None:
                labels = [j for j, _, _ in txt_cells]
                for g0 in range(0, len(labels) // 4 * 4, 4):
                    jd, jn, jl, jh = labels[g0:g0 + 4]
                    dia = _to_float(vrow[jd]) if jd < len(vrow) else None
                    cnt = _to_float(vrow[jn]) if jn < len(vrow) else None
                    low = (_to_float(vrow[jl]) if jl < len(vrow) else None) or 0.0
                    high = (_to_float(vrow[jh]) if jh < len(vrow) else None) or 0.0
                    if not dia or not cnt:
                        continue
                    remark = ""
                    if low > 0 and high > 0:
                        remark = "桩长含40m以下与40m以上，按平均桩长（合计m÷根数）"
                    piles.append({"kind": "桩基参数", "fenxiang": "桩基",
                                  "pos": "", "remark": remark, "count": cnt,
                                  "dia": dia, "depth": (low + high) / cnt,
                                  "row": i + 1})
            continue
        # 分项名：父材料行 → 父名；独立行（col1 空）→ 材料名；普通行 → 列部位
        if (block in _MATRIX_PARENT_FX
                or any(k in block for k in ("支座", "波纹管", "锚具", "声测管"))):
            fx = block
        elif not c1:
            fx = mat
        else:
            fx = None
        # 长文本行（伸缩缝/排水系统/防抛网/备注 等）：单格「名称：数值+单位」对
        if txt_cells and not num_cells:
            for _j, _, s in txt_cells:
                for m in re.finditer(
                        r"([^：:；;，,。\s]{1,40})[：:；;]\s*(\d+(?:\.\d+)?)"
                        r"\s*(㎡|m³|m2|m3|个|套|根|米|处|块|方|[a-zA-Z]{1,4})?",
                        s):
                    name = m.group(1).strip()
                    unit = m.group(3) or ""
                    pm = re.match(r"^(.*?)\s*[（(]([^（）()]+)[）)]\s*$", name)
                    if pm and _norm_text(pm.group(2)) in _UNIT_SUFFIXES:
                        name = pm.group(1).strip()
                        unit = unit or pm.group(2)
                    v = float(m.group(2))
                    if v == 0:
                        continue
                    fxn = ("超声波检测管"
                           if re.search(r"声测管|检测管", _norm_text(name))
                           else fx)
                    out.append({"fenxiang": fxn, "mat": name,
                                "mat_unit": _UNIT_MAP.get(_norm_text(unit), unit),
                                "value": v})
                # 无冒号变体：名称（单位）+数值（如 管卡（kg)4208.12、Q235钢材17370）
                for m in re.finditer(
                        r"([^：:；;，,。]*?)(?:[（(]([^（）()：:；;，,。\s]{1,6})[）)])?"
                        r"\s*(\d+(?:\.\d+)?)\s*(?=$|[;；。])",
                        s):
                    name = m.group(1).strip()
                    if not name or not m.group(3):
                        continue
                    unit = m.group(2) or ""
                    if not unit:  # 无单位：同行小件钢材按 kg 计
                        unit = ("kg" if re.search(r"钢|铁|管|栓|钉|丝|网|板",
                                                  _norm_text(name)) else "")
                    v = float(m.group(3))
                    if v == 0:
                        continue
                    out.append({"fenxiang": fx, "mat": name,
                                "mat_unit": _UNIT_MAP.get(_norm_text(unit), unit),
                                "value": v})
            continue
        # label 格→下一非空数值格配对（拆桥/护栏切除/主梁悬臂改造/旧桥铺装拆除/
        # 防腐面积 等；合并单元格跨度可跨空列，数值 0 不配对）
        if txt_cells and num_cells:
            cells = sorted(txt_cells + num_cells, key=lambda x: x[0])
            for k, (_j, kind, val) in enumerate(cells):
                if kind != "t" or k + 1 >= len(cells):
                    continue
                _nj, nkind, nval = cells[k + 1]
                if nkind != "n":
                    continue
                lab = val.strip().rstrip("：:")
                lu = cur_unit
                pm = re.match(r"^(.*?)\s*[（(]([^（）()]+)[）)]\s*$", lab)
                if pm and _norm_text(pm.group(2)) in _UNIT_SUFFIXES:
                    lab = pm.group(1).strip()
                    lu = pm.group(2)
                out.append({"fenxiang": fx, "mat": lab,
                            "mat_unit": _UNIT_MAP.get(_norm_text(lu), lu),
                            "value": nval})
            continue
        # 普通数值矩阵行：部位列 × 材料行交叉值 = 工程量
        mun = "m3" if mat_n in ("挖方", "填方") else cur_unit
        for j, col_name in cols:
            if j >= len(r) or r[j] in (None, ""):
                continue
            v = _to_float(r[j])
            if v is None or v == 0:
                continue
            out.append({"fenxiang": fx or col_name, "mat": mat,
                        "mat_unit": mun, "value": v})
    return out, piles


def parse_quantity_excel(data: bytes) -> list[dict]:
    """解析海泰式数量总表：分项材料表块（表头行 分项名/适用位置/备注/数量/材料列…）。

    返回 [{sheet, kind:"材料"|"桩基参数", fenxiang, pos, remark, count,
           mat, mat_unit, value, row}] + 桩基参数 {dia, depth, 根数=count}。
    """
    # 格式嗅探：openpyxl 先试，失败回退 xlrd
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        sheets = []
        for ws in wb.worksheets:
            sheets.append((ws.title, [list(r) for r in ws.iter_rows(values_only=True)]))
        xls_mode = False
    except Exception:
        import xlrd
        try:
            wb = xlrd.open_workbook(file_contents=data)
        except Exception as e:
            raise BridgeMatchError(f"无法读取数量表文件：{e}") from e
        sheets = []
        for ws in wb.sheets():
            sheets.append((ws.name,
                           [[ws.cell_value(i, j) for j in range(ws.ncols)] for i in range(ws.nrows)]))
        xls_mode = True

    items = []
    skipped = []
    # 汇总矩阵总表选择：优先 sheet 名含 合计/汇总/总计/工程数量表/总表，
    # 否则退回表头含「工程材料」或秦滨「项目+单位」的内容识别（人名等分表不优先）
    matrix_rows = matrix_sname = None
    for sname, rows in sheets:
        if not (_looks_like_matrix(rows) or _looks_like_qinbin(rows)):
            continue
        if any(k in _norm_text(sname) for k in _MATRIX_SHEET_KEYS):
            matrix_rows, matrix_sname = rows, sname
            break
        if matrix_rows is None and not any(k in _norm_text(sname) for k in _SKIP_SHEET_KEYS):
            matrix_rows, matrix_sname = rows, sname
    matrix_items = []
    if matrix_rows is not None:
        if _looks_like_qinbin(matrix_rows):
            matrix_items, qinbin_piles = _parse_qinbin_matrix(matrix_rows)
            for p in qinbin_piles:
                p["sheet"] = matrix_sname
                items.append(p)
        else:
            matrix_items = _parse_matrix_sheet(matrix_rows)  # 权威汇总矩阵，最后覆盖合并
        # 桩基类型汇总行（"桩径1m桩共：10 根; …桩长…合计：300 m; …"）
        # → 按桩径汇总的桩基参数条目（工程量以矩阵总表为准）
        for r in matrix_rows:
            if r and _norm_text(str(r[1] or "")) == "桩基类型":
                txt = str(r[4] or "") if len(r) > 4 else ""
                groups = {}
                for p in re.split(r"[;；]", txt):
                    m = re.search(r"桩径([\d.]+)m", p)
                    if not m:
                        continue
                    g = groups.setdefault(float(m.group(1)),
                                          {"dia": float(m.group(1)),
                                           "count": None, "total": 0.0})
                    if "桩共" in p:
                        mc = re.search(r"共[：:]\s*([\d.]+)\s*根", p)
                        if mc:
                            g["count"] = float(mc.group(1))
                    else:  # 桩长…合计：X m
                        mt = re.search(r"合计[：:]\s*([\d.]+)\s*m", p)
                        if mt:
                            g["total"] += float(mt.group(1))
                for g in groups.values():
                    if g["count"] and g["total"]:
                        items.append({"sheet": matrix_sname, "kind": "桩基参数",
                                      "fenxiang": "桩基类型", "pos": "", "remark": "",
                                      "count": g["count"], "dia": g["dia"],
                                      "depth": g["total"] / g["count"], "row": 0})
        # 超声波检测管备注汇总行（"Q235B：288.4 kg; 钢管φ57x3.5：25513.8 m; …"）
        # → 逐项拆成条目（矩阵总表为权威工程量，明细 sheet 同名条目下方去重）
        for r in matrix_rows:
            if not r:
                continue
            fx = str(r[1] or "").strip() if len(r) > 1 else ""
            if "检测管" not in _norm_text(fx):
                continue
            txt = ""
            for j in range(4, len(r)):
                if r[j]:
                    txt = str(r[j]).strip()
                    break
            for p in re.split(r"[;；]", txt):
                m = re.match(r"\s*([^：:]+?)\s*[：:]\s*([\d.]+)\s*([^；;]*)$", p)
                if not m:
                    continue
                mun = _norm_text(m.group(3).strip())
                mun = _UNIT_MAP.get(mun, mun)
                items.append({"sheet": matrix_sname, "kind": "材料",
                              "fenxiang": fx, "pos": "", "remark": "",
                              "count": None, "mat": m.group(1).strip(),
                              "mat_unit": mun, "value": float(m.group(2)), "row": 0})
    for sname, rows in sheets:
        if sname == matrix_sname:
            continue
        if any(k in _norm_text(sname) for k in _SKIP_SHEET_KEYS):
            skipped.append(sname)
            continue
        header_idx = [i for i, r in enumerate(rows)
                      if r and _norm_text(r[0]) == "分项名"]
        if not header_idx:
            continue
        header_idx.append(len(rows))  # 末哨兵
        for bi in range(len(header_idx) - 1):
            h = header_idx[bi]
            hend = header_idx[bi + 1]
            hdr = rows[h]
            # 单位行 = 表头下一行且首列为空
            unit_row = rows[h + 1] if h + 1 < hend and rows[h + 1] and not str(rows[h + 1][0] or "").strip() else None
            # 材料列（第 5 列起，即索引 4+）
            mat_cols = [(j, str(hdr[j]).strip(), _norm_text(str(hdr[j])))
                        for j in range(4, len(hdr)) if str(hdr[j] or "").strip()]
            pile_cols = {_norm_text(str(hdr[j])): j for j in range(4, len(hdr))}
            is_pile_block = ("桩径" in pile_cols) and ("桩长" in pile_cols)
            fenxiang = None
            for i in range(h + (2 if unit_row else 1), hend):
                r = rows[i]
                if not r:
                    continue
                c0 = str(r[0] or "").strip() if len(r) > 0 else ""
                if _norm_text(c0) == "分项名":
                    break  # 下一个块（安全兜底）
                if c0:
                    fenxiang = c0
                if not fenxiang or any(k in fenxiang for k in ("此行不填", "本行不填", "不填")):
                    continue
                pos = str(r[1] or "").strip() if len(r) > 1 else ""
                remark = str(r[2] or "").strip() if len(r) > 2 else ""
                count = _to_float(r[3]) if len(r) > 3 else None
                if is_pile_block:
                    dia = _to_float(r[pile_cols["桩径"]]) if "桩径" in pile_cols and len(r) > pile_cols["桩径"] else None
                    dep = _to_float(r[pile_cols["桩长"]]) if "桩长" in pile_cols and len(r) > pile_cols["桩长"] else None
                    if count is not None and dia is not None and dep is not None:
                        items.append({"sheet": sname, "kind": "桩基参数", "fenxiang": fenxiang,
                                      "pos": pos, "remark": remark, "count": count,
                                      "dia": dia, "depth": dep, "row": i + 1})
                    continue
                # 材料列取值
                row_has_value = False
                for j, mname, _ in mat_cols:
                    if len(r) <= j or r[j] in (None, ""):
                        continue
                    v = _to_float(r[j])
                    if v is None:
                        continue
                    if count is not None:
                        v = v * count  # 海泰口径：合计工程量 = 数量(倍数) × 材料量
                    mun = str(unit_row[j] or "").strip() if unit_row and len(unit_row) > j else ""
                    items.append({"sheet": sname, "kind": "材料", "fenxiang": fenxiang,
                                  "pos": pos, "remark": remark, "count": count,
                                  "mat": mname, "mat_unit": mun, "value": v, "row": i + 1})
                    row_has_value = True
                # 备注列带单位（如 泄水孔 的 套/m/处）
                if remark and _norm_text(remark) in _UNIT_MAP and count is not None:
                    items.append({"sheet": sname, "kind": "材料", "fenxiang": fenxiang,
                                  "pos": pos, "remark": remark, "count": None,
                                  "mat": f"{fenxiang}（{remark}）", "mat_unit": remark,
                                  "value": count, "row": i + 1})
                    row_has_value = True
                _ = row_has_value
    # 桩基参数以工程数量表桩基类型汇总为准（按桩径合并根数/桩长）
    if matrix_sname and any(it.get("kind") == "桩基参数" and it.get("sheet") == matrix_sname
                            for it in items):
        items = [it for it in items
                 if it.get("kind") != "桩基参数" or it.get("sheet") == matrix_sname]
    # 超声波检测管以矩阵总表备注汇总行为准（长度权威量），明细 sheet 同名条目不重复计入
    if matrix_sname and any("检测管" in _norm_text(it.get("fenxiang") or "")
                            and it.get("sheet") == matrix_sname for it in items):
        items = [it for it in items
                 if "检测管" not in _norm_text(it.get("fenxiang") or "")
                 or it.get("sheet") == matrix_sname]
    # 矩阵总表为权威工程量：同部位同材料的明细条目以矩阵值为准
    if matrix_items:
        covered = {(_norm_text(m["fenxiang"]), _norm_text(m["mat"]))
                   for m in matrix_items}
        detail_by_key = {}
        for it in items:
            if it.get("kind") == "材料":
                detail_by_key.setdefault((_norm_text(it["fenxiang"]),
                                          _norm_text(it["mat"])), it)
        merged = [it for it in items
                  if it.get("kind") != "材料"
                  or (_norm_text(it["fenxiang"]), _norm_text(it["mat"])) not in covered]
        # 桩基"单桩用量"明细条目（海泰数量列=1 的小钢种小值）：
        # 矩阵已按 HPB300/HRB400 给出该部位钢筋笼总量 → 删除，避免重复/口径混乱
        pile_rebar_fx = {_norm_text(m["fenxiang"]) for m in matrix_items
                         if "桩基" in _norm_text(m["fenxiang"])
                         and re.search(r"HPB|HRB", _norm_text(m["mat"]))}
        merged = [it for it in merged
                  if not (it.get("kind") == "材料"
                          and _norm_text(it["fenxiang"]) in pile_rebar_fx
                          and re.search(r"HPB235|HRB335", _norm_text(it["mat"]))
                          and (_to_float(it.get("value")) or 0) <= 100)]
        for m in matrix_items:
            mu = m["mat_unit"]
            if not mu:  # 矩阵钢筋行单位列为空 → 先借用明细条目单位
                d = detail_by_key.get((_norm_text(m["fenxiang"]), _norm_text(m["mat"])))
                mu = (d or {}).get("mat_unit") or ""
            if not mu:  # 仍为空 → 按材料推断（钢材类 t、混凝土/沥青/砂浆类 m3、涂料类 m2），备注注明
                mn = _norm_text(m["mat"])
                if re.search(r"HPB|HRB|CRB|冷轧|不锈钢|铸铁|钢棒|钢纤维|钢绞线|^Q\d", mn):
                    mu = "t"
                elif re.search(r"涂料|漆|涂装", mn):
                    mu = "m2"
                elif re.search(r"防水|黏层|透层|封层", mn) and not re.search(r"混凝土|砂浆", mn):
                    mu = "m2"
                elif re.search(r"混凝土|砂浆|沥青|碎石|垫层|砂|土|^C[aA]?\d", mn):
                    mu = "m3"
            merged.append({"sheet": matrix_sname, "kind": "材料", "fenxiang": m["fenxiang"],
                           "pos": "", "remark": "单位缺省，按材料推断" if not m["mat_unit"] and mu else "",
                           "count": None,
                           "mat": m["mat"], "mat_unit": mu,
                           "value": m["value"], "row": 0})
        items = merged
    if not items:
        raise BridgeMatchError(
            "未解析到任何分项材料条目。请上传海泰数量总表式文件"
            "（sheet 内含「分项名/适用位置/备注/数量/材料列」块，"
            "或「工程材料×部位」矩阵式总表）。")
    return items


def parse_text_lines(text: str) -> tuple[list[dict], list[str]]:
    """粘贴文本 → 条目列表 + 解析警告。每行「名称 数量 单位」。"""
    items, warns = [], []
    for ln, line in enumerate(str(text or "").splitlines(), 1):
        line = line.strip()
        if not line:
            continue
        if any(k in line for k in ("项目名称", "工程量", "序号")):
            continue
        m = re.match(r"^(.*?)[\s\t,，]+(\d+(?:\.\d+)?)[\s\t]*(\S{1,6})?$", line)
        if not m:
            warns.append(f"第 {ln} 行无法解析（缺少数量）：{line[:30]}")
            continue
        items.append({"sheet": "粘贴文本", "kind": "材料", "fenxiang": m.group(1).strip(),
                      "pos": "", "remark": "", "count": None,
                      "mat": "", "mat_unit": m.group(3) or "", "value": float(m.group(2)),
                      "row": ln})
    return items, warns


# ---------------------------------------------------------------------------
# 特征提取与打分
# ---------------------------------------------------------------------------

COMPONENT_GROUPS = [
    ("桩基础", ["钻孔灌注桩", "灌注桩", "钢护筒", "护筒", "桩基", "钢筋笼", "声测管",
                "超声波检测管", "检测管", "凿桩头", "桩头", "挖孔桩", "钢管桩",
                "预应力管桩", "预制桩"]),
    ("承台系", ["承台", "围堰", "垫层", "基坑", "垫石"]),
    ("墩台系", ["墩身", "墩柱", "桥台", "台身", "台帽", "盖梁", "耳墙", "背墙",
                "柱式墩", "墩台", "支撑梁", "横梁", "系梁", "墩"]),
    ("梁体", ["箱梁", "板梁", "T梁", "空心板", "实心板", "盖板", "预制梁", "现浇梁",
              "桥面板", "横隔板", "湿接缝", "梁板", "梁"]),
    ("预应力系", ["钢绞线", "预应力", "锚具", "波纹管"]),
    ("桥面系", ["桥面铺装", "铺装", "防水层", "防水混凝土", "伸缩缝", "搭板", "枕梁",
                "桥面连续", "沥青"]),
    ("护栏栏杆", ["防撞护栏", "护栏", "防撞墙", "栏杆", "人行道", "缘石",
                  "隔声屏障", "声屏障"]),
    ("排水", ["泄水孔", "泄水管", "透水管", "排水管", "盲沟", "排水"]),
    ("支座", ["支座", "板式支座", "盆式支座", "球形支座"]),
    ("临时工程", ["支架", "脚手架", "模板", "操作平台", "便道"]),
    ("土方", ["土方", "开挖", "基坑", "回填", "填方", "挖方", "开槽", "运土", "废料",
              "护坡", "锥坡"]),
    ("钢筋", ["钢筋", "HPB", "HRB", "CRB"]),
    ("混凝土", ["混凝土"]),
]

METHODS = ["钻孔", "成孔", "灌注", "埋设", "沉入", "打入", "锤击", "振动", "预制", "现浇",
           "安装", "制作", "绑扎", "张拉", "先张", "后张", "压浆", "运输", "吊装", "架设",
           "浇筑", "摊铺", "碾压", "拆除", "凿除", "抹面", "防水", "涂装", "焊接", "切割",
           "顶推", "悬浇"]
SOIL = ["砂土", "黏土", "砂砾", "砂卵石", "岩石", "风化岩", "淤泥"]
GRADES = ["细粒式", "中粒式", "粗粒式", "砂粒式"]
# 塑料类管材材质关键词（用户 2026-10-27 口径）：PVC/PE/HDPE/PPR/ABS/CPVC/PEX/PB/FRP
# 玻璃钢/铝塑/玻纤 等，名称含这些的管材确定为塑料类 → 排除钢管/铸铁管定额候选
_PLASTIC_PIPE_RE = re.compile(
    r"(?:pvc|upvc|cpvc|hdpe|ldpe|pe-rt|pp-r|pp-h|pp-b|ppr|pex|pe-x|xpe|mpp|frp|grp"
    r"|(?<![a-z0-9])pe(?![a-z0-9])|(?<![a-z0-9])pp(?![a-z0-9])"
    r"|(?<![a-z0-9])pb(?![a-z0-9])|(?<![a-z0-9])abs(?![a-z0-9])"
    r"|玻璃钢|玻纤|玻璃纤维|铝塑|塑料)", re.I)
# 塑料管具体材质族（用户 2026-10-27：PVC直通管 → 3-10-30 UPVC泄水管 唯一命中，
# 同族（PVC/PE/PP…）加分，泛「塑料」候选不加）
_PLASTIC_FAMILIES = (
    (re.compile(r"(?:pvc|upvc|cpvc)", re.I), "PVC"),
    (re.compile(r"(?:hdpe|ldpe|pe-rt|pex|pe-x|xpe|(?<![a-z0-9])pe(?![a-z0-9]))", re.I), "PE"),
    (re.compile(r"(?:ppr|pp-r|pp-h|pp-b|mpp|(?<![a-z0-9])pp(?![a-z0-9]))", re.I), "PP"),
    (re.compile(r"(?<![a-z0-9])abs(?![a-z0-9])", re.I), "ABS"),
    (re.compile(r"(?<![a-z0-9])pb(?![a-z0-9])", re.I), "PB"),
    (re.compile(r"frp|grp|玻璃钢|玻纤|玻璃纤维", re.I), "玻璃钢/玻纤"),
    (re.compile(r"铝塑", re.I), "铝塑"),
)


def _extract_features(text):
    """提取构件组/工艺/材料/规格/条件。"""
    t = _norm_text(text)
    comps, methods, materials, conds = set(), set(), set(), set()
    for group, kws in COMPONENT_GROUPS:
        for kw in sorted(kws, key=len, reverse=True):
            if kw in t:
                comps.add(group)
                break
    for kw in METHODS:
        if kw in t:
            methods.add(kw)
    if re.search(r"C\d{2}", t):
        materials.add(re.search(r"C\d{2}", t).group())
    for kw in ("混凝土", "钢筋", "钢绞线", "沥青", "浆砌", "片石", "碎石", "砂浆",
               "钢管", "铸铁管", "塑料管", "螺纹钢", "3号钢"):
        if kw in t:
            materials.add(kw)
    if _PLASTIC_PIPE_RE.search(t):
        materials.add("塑料管")
    for kw in SOIL:
        if kw in t:
            conds.add(kw)
    for kw in GRADES:
        if kw in t:
            conds.add(kw)
    specs = {}
    m = re.search(r"[φΦ∅⌀]\s*≤\s*(\d+(?:\.\d+)?)", t)
    if m:
        specs["dia_le"] = float(m.group(1))
    else:
        # 注意 (?!\d*…)：防止 \d+ 回溯到 φ10以内 的 "1"
        m = re.search(r"[φΦ∅⌀]\s*(\d+(?:\.\d+)?)(?!\d*\s*(以内|以外))", t)
        if m:
            specs["dia"] = float(m.group(1))
    m = re.search(r"[hH]\s*≤\s*(\d+(?:\.\d+)?)\s*m?", t)
    if m:
        specs["depth_le"] = float(m.group(1))
    m = re.search(r"桩长\s*≤?\s*(\d+(?:\.\d+)?)", t)
    if m and "depth_le" not in specs:
        specs["depth_le"] = float(m.group(1))
    m = re.search(r"(\d+)\s*孔", t)
    if m:
        specs["holes"] = int(m.group(1))
    m = re.search(r"(\d+)\s*号钢", t)
    if m:
        specs["steel_no"] = m.group(1)
    # 钢种：HRB/CRB→螺纹钢，HPB/Q235→3号钢
    if re.search(r"HRB|CRB", t):
        specs["steel_type"] = "螺纹钢"
        materials.add("螺纹钢")
    elif re.search(r"HPB|Q235", t):
        specs["steel_type"] = "3号钢"
        materials.add("3号钢")
    # 伸缩缝型号数字（KF-80/TS160 等 → 80/160）
    m = re.search(r"(?:KF|TS|GQF|CD|F)[-]?(\d{2,4})", t)
    if m:
        specs["type_no"] = int(m.group(1))
    # 钢筋直径边界（φ10以内/φ14以外）
    m = re.search(r"[φΦ∅]\s*(\d+(?:\.\d+)?)\s*(以内|以外)", t)
    if m:
        specs["dia_bound"] = (float(m.group(1)), m.group(2))
    return {"comps": comps, "methods": methods, "materials": materials,
            "conds": conds, "specs": specs, "text": t}


# 分部 → 构件类别（用于「构件类别与分部一致」+5）
_GROUP_FAMILY = {
    "桩基础": "桩基础", "承台": "承台系", "支撑梁": "墩台系", "系梁": "墩台系",
    "台帽": "墩台系", "墩帽": "墩台系", "桥台": "墩台系", "桥墩": "墩台系",
    "墩盖梁": "墩台系", "台盖梁": "墩台系", "预应力砼现浇箱梁": "梁体",
    "预应力混凝土小箱梁": "梁体", "预应力板梁": "梁体", "钢箱梁": "梁体",
    "桥面铺装": "桥面系", "沥青铺装": "桥面系", "防水铺装": "桥面系",
    "防水层": "桥面系", "搭板": "桥面系", "伸缩缝": "桥面系",
    "人行道板": "护栏栏杆", "防撞护栏": "护栏栏杆", "隔声屏障": "护栏栏杆",
    "支座": "支座", "排水": "排水", "土方": "土方", "护坡": "土方",
}
# 分部 → 全库补强检索的名称关键词（用于同分部变体定额）
_GROUP_FAMILY_KW = {
    "桩基础": "灌注桩|钻孔桩", "承台": "承台", "支撑梁": "支撑梁|横梁",
    "系梁": "系梁|横梁", "台帽": "台帽", "墩帽": "墩帽",
    "桥台": "台身|桥台|台帽", "桥墩": "墩|柱", "墩盖梁": "盖梁", "台盖梁": "盖梁",
    "预应力砼现浇箱梁": "箱梁", "预应力混凝土小箱梁": "小箱梁",
    "预应力板梁": "板梁|空心板", "钢箱梁": "钢箱梁|箱梁",
    "桥面铺装": "桥面|铺装", "沥青铺装": "沥青", "防水铺装": "桥面|防水",
    "防水层": "防水层", "搭板": "搭板",
    "防撞护栏": "护栏|防撞", "支座": "支座", "伸缩缝": "伸缩缝",
    "排水": "泄水|透水|盲沟", "人行道板": "人行道",
    "隔声屏障": "隔声屏障|声屏障", "土方": "土方|开挖|回填", "护坡": "护坡|锥形坡",
}


def _bigrams(t):
    return {t[i:i + 2] for i in range(len(t) - 1)} if len(t) >= 2 else set()


def _score(item_feats, entry_feats, item_unit, entry_unit, item_group="", item_mat=""):
    """规则打分 → (分数, 命中理由列表, 单位是否兼容)。"""
    score, reasons, unit_ok = 0, [], True
    if item_feats["text"] and item_feats["text"] == entry_feats["text"]:
        return 100, ["名称完全一致"], True
    # 构件
    hit = item_feats["comps"] & entry_feats["comps"]
    comp_pts = min(len(hit) * 12, 24)
    score += comp_pts
    for c in hit:
        reasons.append(f"构件一致：{c}")
    # 工艺
    hit = item_feats["methods"] & entry_feats["methods"]
    meth_pts = min(len(hit) * 6, 12)
    score += meth_pts
    if hit:
        reasons.append("工艺一致：" + "、".join(sorted(hit)))
    # 材料（C 等级 + 材料词）
    hit = item_feats["materials"] & entry_feats["materials"]
    mat_pts = min(len(hit) * 5, 10)
    score += mat_pts
    if hit:
        reasons.append("材料一致：" + "、".join(sorted(hit)))
    # 规格
    isp, esp = item_feats["specs"], entry_feats["specs"]
    if "dia" in isp and "dia" in esp and isp["dia"] == esp["dia"]:
        score += 8
        reasons.append(f"桩径一致 φ{isp['dia']:g}")
    if "dia" in isp and "dia_le" in esp and isp["dia"] <= esp["dia_le"]:
        score += 6
        reasons.append(f"桩径 φ{isp['dia']:g} 在定额 φ≤{esp['dia_le']:g} 范围内")
    if "dia" in isp and "dia" in esp and isp["dia"] != esp["dia"]:
        score -= 6
        reasons.append(f"直径不同：{isp['dia']:g} vs {esp['dia']:g}")
    # 钢筋直径边界（φ10以内/φ14以外）
    if "dia" in isp and "dia_bound" in esp:
        v, mode = esp["dia_bound"]
        if (mode == "以内" and isp["dia"] <= v) or (mode == "以外" and isp["dia"] > v):
            score += 6
            reasons.append(f"直径 ∅{isp['dia']:g} 在定额 φ{v:g}{mode}范围内")
        else:
            score -= 6
            reasons.append(f"直径 ∅{isp['dia']:g} 不在定额 φ{v:g}{mode}范围")
    if "dia_bound" in isp and "dia" in esp:
        v, mode = isp["dia_bound"]
        if (mode == "以内" and esp["dia"] <= v) or (mode == "以外" and esp["dia"] > v):
            score += 6
            reasons.append(f"定额直径 φ{esp['dia']:g} 在清单 φ{v:g}{mode}范围内")
        else:
            score -= 6
            reasons.append(f"定额直径 φ{esp['dia']:g} 不在清单 φ{v:g}{mode}范围")
    if "depth_le" in isp and "depth_le" in esp and isp["depth_le"] == esp["depth_le"]:
        score += 6
        reasons.append(f"深度一致 h≤{isp['depth_le']:g}m")
    if "depth_le" in isp and "depth_le" in esp and isp["depth_le"] < esp["depth_le"]:
        score += 3
        reasons.append(f"深度 h≤{isp['depth_le']:g}m 在定额 h≤{esp['depth_le']:g}m 范围内")
    c_item = [m for m in isp_materials(item_feats) if re.match(r"C\d{2}$", m)]
    c_ent = [m for m in esp_materials(entry_feats) if re.match(r"C\d{2}$", m)]
    if c_item and c_ent:
        if c_item[0] == c_ent[0]:
            score += 6
            reasons.append(f"混凝土强度一致 {c_item[0]}")
        else:
            score -= 6
            reasons.append(f"混凝土强度不同：清单 {c_item[0]} vs 定额 {c_ent[0]}")
    for k in ("holes", "steel_no"):
        if k in isp and k in esp:
            if isp[k] == esp[k]:
                score += 5
                reasons.append(f"{k} 一致")
            else:
                score -= 6
                reasons.append(f"{k} 不同：{isp[k]} vs {esp[k]}")
    # 钢种（螺纹钢/3号钢）
    if "steel_type" in isp and "steel_type" in esp:
        if isp["steel_type"] == esp["steel_type"]:
            score += 6
            reasons.append(f"钢种一致 {isp['steel_type']}")
        else:
            score -= 6
            reasons.append(f"钢种不同：{isp['steel_type']} vs {esp['steel_type']}")
    # 伸缩缝型号（KF-80 ↔ TS80 伸缩量一致）
    if "type_no" in isp and "type_no" in esp:
        if isp["type_no"] == esp["type_no"]:
            score += 8
            reasons.append(f"伸缩量规格一致 {isp['type_no']}")
        else:
            score -= 6
            reasons.append(f"伸缩量不同：{isp['type_no']} vs {esp['type_no']}")
    # 地层
    hit = item_feats["conds"] & entry_feats["conds"]
    if hit:
        score += 4
        reasons.append("条件一致：" + "、".join(sorted(hit)))
    # 构件类别与分部一致
    if item_group:
        fam = _GROUP_FAMILY.get(_map_group(item_group))
        if fam and fam in entry_feats["comps"]:
            score += 5
            reasons.append(f"构件类别与分部一致（{fam}）")
    # 塑料管具体材质族一致（用户 2026-10-27：PVC直通管 → 3-10-30 UPVC泄水管 唯一命中）
    if "塑料管" in item_feats["materials"] and "塑料管" in entry_feats["materials"]:
        for fam_re, fam_label in _PLASTIC_FAMILIES:
            if fam_re.search(item_feats["text"]) and fam_re.search(entry_feats["text"]):
                score += 8
                reasons.append(f"材质族一致：{fam_label}（塑料管）")
                break
    # bigram
    score += min(len(_bigrams(item_feats["text"]) & _bigrams(entry_feats["text"])), 15)
    # 名称包含
    a, b = item_feats["text"], entry_feats["text"]
    if len(a) >= 4 and len(b) >= 4 and (a in b or b in a):
        score += 8
        reasons.append("名称相互包含")
    # 材料名完整包含（用户 2026-10-27：泄水孔 碎石 → 1-6-13 碎石垫层、
    # 浆砌片石(M15砂浆) → 1-1-7 浆砌片石 护底 唯一命中）
    mat_core = re.sub(r"[（(][^）)]*[）)]", "", _norm_text(item_mat or ""))
    if len(mat_core) >= 2 and mat_core in entry_feats["text"]:
        score += 8
        reasons.append(f"材料名完整一致：{mat_core}")
    # 单位
    ib, im = _unit_parts(item_unit) if item_unit else ("", 1.0)
    eb, em = _unit_parts(entry_unit)
    if item_unit and ib and eb:
        if ib == eb and im == em:
            score += 4
            reasons.append(f"单位一致 {item_unit}")
        elif ib == eb:
            score += 3
            reasons.append(f"单位同基类（{item_unit} ↔ {entry_unit}）")
        elif ib in _COUNT_UNITS and eb in _COUNT_UNITS:
            score += 3
            reasons.append("同为计数单位")
        elif {ib, eb} <= {"kg", "t"}:
            score += 3
            reasons.append(f"单位 {item_unit} ↔ {entry_unit}（kg/t）")
        elif "沥青" in item_feats["materials"] and {ib, eb} <= {"m3", "t"}:
            # 沥青混凝土清单 m3 ↔ 定额 t 为合法换算（定额说明第8条密度换算）
            score += 3
            reasons.append(f"沥青混凝土 {item_unit} ↔ {entry_unit}（定额说明第8条密度换算）")
        else:
            score -= 10
            unit_ok = False
            reasons.append(f"单位基类不符（{item_unit} vs {entry_unit}）")
    return score, reasons, unit_ok


def isp_materials(feats):
    return feats["materials"]


def esp_materials(feats):
    return feats["materials"]


# ---------------------------------------------------------------------------
# 匹配主流程
# ---------------------------------------------------------------------------

TH_SINGLE = 16
MARGIN = 8
MIN_CANDIDATE = 10
MAX_CANDIDATES = 5

_GROUP_MAP = [
    ("桩基础", ["桩基", "基桩", "灌注桩", "凿桩头", "检测管", "声测管"]),
    ("承台", ["承台"]),
    ("支撑梁", ["支撑梁", "横梁"]),
    ("系梁", ["系梁"]),
    ("台帽", ["台帽"]),
    ("墩帽", ["墩帽"]),
    ("桥台", ["台身", "桥台"]),
    ("桥墩", ["墩柱", "墩身", "桥墩"]),
    ("钢箱梁", ["钢箱梁"]),
    ("预应力混凝土小箱梁", ["小箱梁"]),
    ("预应力砼现浇箱梁", ["现浇箱梁", "箱梁"]),
    ("预应力板梁", ["板梁", "空心板"]),
    ("桥面铺装", ["桥面铺装"]),
    ("沥青铺装", ["沥青混凝土铺装", "沥青铺装"]),
    ("防水铺装", ["防水混凝土铺装", "防水铺装"]),
    ("防水层", ["防水层"]),
    ("搭板", ["搭板"]),
    ("防撞护栏", ["防撞护栏", "护栏"]),
    ("支座", ["支座"]),
    ("伸缩缝", ["伸缩缝"]),
    ("排水", ["泄水孔", "泄水管", "透水管", "排水"]),
    ("人行道板", ["人行道"]),
    ("隔声屏障", ["隔声屏障", "声屏障"]),
    ("土方", ["土方", "开槽", "回填", "挖方", "填方"]),
    ("护坡", ["护坡", "锥坡"]),
    # 梯道放最后：「梯道下部 桥墩/承台」等复合分项名先命中前面的 桥墩/承台
    ("梯道", ["梯道"]),
]


def _steel_mat(mat):
    """非钢筋钢材判定：Q 牌号钢（Q235/Q345…）或钢板/型钢类材料 → 按预埋铁件口径。"""
    if re.match(r"^Q\d{3}", mat):
        return True
    return any(k in mat for k in ("钢板", "型钢", "角钢", "槽钢", "工字钢", "扁钢"))


def _pipe_kg_per_m(mat):
    """钢管 φD×t → 每米理论重量 kg/m（碳钢 0.02466 系数，用户 2026-10-26 提供公式
    W = 0.02466 × (D − t) × t，D/t 单位 mm；如 φ57×3.5 → 4.618 kg/m）。"""
    m = re.search(r"[φd](\d+(?:\.\d+)?)\s*[x×]\s*(\d+(?:\.\d+)?)", _norm_text(mat))
    if not m:
        return None
    d, t = float(m.group(1)), float(m.group(2))
    return 0.02466 * (d - t) * t


def _map_group(fenxiang):
    t = _norm_text(fenxiang)
    if "台盖梁" in t:
        return "台盖梁"
    if "盖梁" in t:
        return "台盖梁" if "桥台" in t else "墩盖梁"
    # 矩阵式数量表「小箱梁现浇部分」= 现浇施工 → 套现浇箱梁规则（用户 2026-10-27 口径）
    if "小箱梁" in t and "现浇" in t:
        return "预应力砼现浇箱梁"
    # 墩柱及柱系梁合并量：整体套桥墩墩台身定额（用户 2026-10-26 口径）；
    # 量表中墩柱/系梁分开成行时，按模板分套（墩柱→桥墩、系梁→系梁）
    if "墩柱" in t and "系梁" in t:
        return "桥墩"
    for g, kws in _GROUP_MAP:
        for kw in sorted(kws, key=len, reverse=True):
            if kw in t:
                return g
    return None


def _group_of(item, rules):
    """分项名 → 规则组名。桥面铺装等复合分项按材料细分：
    沥青 → 沥青铺装、防水混凝土 → 防水铺装、防水层 → 防水层（仅当模板含该组）。"""
    group = _map_group(item.get("fenxiang") or "")
    mat = _norm_text(item.get("mat") or "")
    if "沥青" in mat and "沥青铺装" in rules:
        group = "沥青铺装"
    elif "防水" in mat and "混凝土" in mat and "防水铺装" in rules:
        group = "防水铺装"
    elif "防水" in mat and "层" in mat and "防水层" in rules:
        group = "防水层"
    return group


# 混凝土→模板定额映射：不同部位混凝土随同套用相应模板定额，
# 模板量=混凝土量×系数。系数默认取参考表样例比例（模板样例量÷混凝土样例量），
# 页面可自行修改；梯道/预制小箱梁/预制板梁样例表无比例，默认系数待定（页面可改）。
# 套法按单位工程.xlsx 模板：承台→3-7-26、支撑梁→3-7-25、系梁→3-7-24、
# 桥台/台帽→3-7-28、桥墩/墩帽→3-7-29、台盖梁→3-7-34、墩盖梁→3-7-33、
# 现浇箱梁→3-7-38、预制小箱梁→3-7-58、预制板梁→3-7-60、防撞护栏→3-7-47、搭板→3-7-42。
# (fenxiang 关键词, 模板编号, 系数键, 默认系数)
_TEMPLATE_MAP = [
    ("承台", "3-7-26", "承台", 1.207),          # 692.6973/573.9
    ("台盖梁", "3-7-34", "盖梁", 3.031),        # 样例 0/0，按墩盖梁系数
    ("盖梁", "3-7-33", "盖梁", 3.031),          # 700.161/231
    ("台帽", "3-7-28", "台帽", 1.499),          # 模板同 3-7-28，系数类推台身
    ("墩帽", "3-7-29", "墩帽", 4.295),          # 模板同 3-7-29，系数类推墩柱
    ("台身", "3-7-28", "桥台", 1.499),          # 299.0505/199.5
    ("桥台", "3-7-28", "桥台", 1.499),          # 分项名直接为「桥台」
    ("支撑梁", "3-7-25", "支撑梁", 1.207),      # 模板 3-7-25，系数类推承台
    ("系梁", "3-7-24", "系梁", 1.207),          # 模板 3-7-24，系数类推承台
    ("墩", "3-7-29", "墩柱", 4.295),            # 12126.9325/2823.5（含柱系梁，样例整体口径）
    ("预制小箱梁", "3-7-58", "预制小箱梁", 1.0),  # 模板 3-7-58，系数待定（页面可改）
    ("预制板梁", "3-7-60", "预制板梁", 1.0),      # 模板 3-7-60，系数待定（页面可改）
    ("箱梁", "3-7-38", "现浇箱梁", 5.387),      # 51725.974/9602
    ("梯道", "3-7-41", "梯道", 5.387),          # 样例无梯道，用户指定默认 5.387
    ("搭板", "3-7-42", "搭板", 3.209),          # 346.572/108
    ("护栏", "3-7-47", "护栏", 4.810),          # 3261.18/678
    ("挡", "3-7-49", "挡墙", 1.61),             # 用户指定默认 1.61（挡土墙）
]
# 页面输入框顺序（键去重）
TPL_RATIOS = []
for _kw, _code, _key, _def in _TEMPLATE_MAP:
    if not any(k == _key for k, _ in TPL_RATIOS):
        TPL_RATIOS.append((_key, _def))


def _template_for(item) -> dict | None:
    """混凝土条目对应模板定额（模板量=混凝土量×系数，系数页面可改）。
    桩基（灌注桩）、桥面铺装、沥青混凝土不套模板。"""
    fx = _norm_text(item.get("fenxiang") or "")
    mat = _norm_text(item.get("mat") or "")
    if "桩基" in fx or "沥青" in mat:
        return None
    if any(k in fx for k in ("拆除", "切除", "改造", "拆桥")):
        return None  # 拆除/改造类条目不套模板
    if not (re.search(r"C\d{2}", mat) or "混凝土" in mat):
        return None
    # 墩柱及柱系梁合并量：整体按桥墩口径套墩台身模板（用户 2026-10-26）
    if "墩柱" in fx and "系梁" in fx:
        return {"code": "3-7-29", "key": "墩柱", "default": 4.295}
    for kw, code, key, default in _TEMPLATE_MAP:
        if kw in fx:
            return {"code": code, "key": key, "default": default}
    return None


def _item_display(item):
    if item["kind"] == "桩基参数":
        return (f"{item['fenxiang']}（{item['count']:g} 根 φ{item['dia']:g}m "
                f"桩长 {item['depth']:g}m）")
    mat = item.get("mat") or ""
    return f"{item['fenxiang']} {mat}".strip()


def _classify(cands):
    """候选分类 → (status, picked, candidates, note)。"""
    cands = [c for c in cands if c["score"] >= MIN_CANDIDATE]
    if not cands:
        return "none", None, [], ""
    best = cands[0]
    if best["score"] >= 100:
        return "single", best, cands[:1], ""
    ok = [c for c in cands if c["unit_ok"]]
    best_ok = ok[0] if ok else None
    second = cands[1]["score"] if len(cands) > 1 else -999
    if best_ok is not None and best["score"] >= TH_SINGLE and best["score"] - second >= MARGIN:
        return "single", best_ok, cands[:MAX_CANDIDATES], ""
    if best["score"] >= TH_SINGLE and best["score"] - second < MARGIN:
        return "multi", None, cands[:MAX_CANDIDATES], "多个候选定额得分接近，请结合差异说明人工选择。"
    return "none", None, cands[:3], "得分最高的候选未达到可信阈值，需补充特征后重新匹配。"


def _checklist(item, feats=None, soil=""):
    """缺失特征 → 补充信息清单（分项类型感知）。
    soil：页面已选地层条件时不再提示补充地层。"""
    feats = feats or _extract_features(_item_display(item))
    t = feats["text"]
    fx = _norm_text(item["fenxiang"])
    mat = _norm_text(item.get("mat") or "")
    missing = []
    is_pile_param = item.get("kind") == "桩基参数"
    is_pile_material = "桩基础" == _map_group(item["fenxiang"]) and item.get("kind") == "材料"
    if not feats["comps"]:
        missing.append("构件类型（如 承台/墩柱/盖梁/护栏/支座…）")
    if (is_pile_param or ("桩" in fx and not is_pile_material)) and "dia" not in feats["specs"] and "dia_le" not in feats["specs"]:
        missing.append("桩径 φ（如 φ1200）")
    if (is_pile_param or ("桩" in fx and not is_pile_material)) and "depth_le" not in feats["specs"]:
        missing.append("桩长/钻孔深度 h（如 h≤40m）")
    if is_pile_param and not (feats["conds"] & set(SOIL)) and not soil:
        missing.append("地层条件（如 砂土/黏土/砂砾/岩石）")
    if ("混凝土" in feats["materials"] and "沥青" not in feats["materials"]
            and not any(k in mat for k in ("涂装", "涂料", "漆"))
            and not any(k in fx for k in ("防腐", "拆除", "切除", "改造", "拆桥"))
            and not any(re.match(r"C\d{2}$", m) for m in feats["materials"])):
        missing.append("混凝土强度等级（如 C30）")
    if re.search(r"HPB|HRB|CRB|钢筋", mat + fx):
        missing.append("钢筋直径 φ10/φ14 以内或以外")
    if "预应力系" in feats["comps"]:
        missing.append("钢绞线束长/孔数")
    if "支座" in feats["comps"]:
        missing.append("支座类型（板式/四氟板式/盆式…）")
    if "护栏栏杆" in feats["comps"]:
        missing.append("栏杆材质类型（钢管/不锈钢/铸铁…）")
    if "排水" in feats["comps"] and "管" in mat:
        if _unit_parts(item.get("mat_unit") or "")[0] == "kg":
            missing.append("泄水管长度（m，定额按 10m 计，kg 无法换算）")
        else:
            if not _PLASTIC_PIPE_RE.search(mat):
                missing.append("泄水管材质（钢管/铸铁管/塑料管）")
    if item.get("kind") != "桩基参数" and not item.get("mat_unit") and not item.get("unit"):
        missing.append("单位（m/m2/m3/t/套）")
    return missing


def _asphalt_width_bands(entries):
    """库内沥青摊铺宽度档（升序），如 [8.5, 14]。"""
    out = set()
    for e in entries:
        m = re.search(r"摊铺宽度\s*([\d.]+)\s*m", e["name"])
        if m and "沥青" in e["name"]:
            out.add(float(m.group(1)))
    return sorted(out)


def _wd_is(name, tw):
    """条目名称的摊铺宽度档是否等于 tw。"""
    m = re.search(r"摊铺宽度\s*([\d.]+)\s*m", name)
    return m is not None and abs(float(m.group(1)) - tw) < 1e-9


def _strip_metal_pipe_cands(item, cands):
    """塑料类管材条目 → 排除候选里名称含 钢管/铸铁管 的定额
    （用户 2026-10-27 口径：量表写 PVC/PE/PPR 等材质即确定为塑料管）。
    返回 (cands, 是否发生过排除)。"""
    if _PLASTIC_PIPE_RE.search(_norm_text(item.get("mat") or "")):
        kept = [c for c in cands
                if not re.search(r"钢管|铸铁管",
                                 _norm_text((c["entry"] or {}).get("name") or ""))]
        if len(kept) < len(cands):
            return kept, True
    return cands, False


def _match_text(item, entries, bridge_width=8.5):
    """全库文本打分（未映射/回退路径）。"""
    text = _item_display(item)
    feats = _extract_features(text)
    cands = []
    # 沥青混凝土摊铺宽度档（用户 2026-10-27 口径）：取 ≥桥梁宽度 的最小档，默认 8.5m 以内；
    # 改性：量表 mat 含 改性/SBS → 套改性沥青混凝土，否则套普通沥青混凝土（不符者降分）
    _mat_n = _norm_text(item.get("mat") or "")
    _mod = bool(re.search(r"改性|sbs", _mat_n, re.I))
    _bands = _asphalt_width_bands(entries)
    _w = bridge_width or 8.5
    _tw = next((b for b in _bands if b >= _w), None)
    _w_over = _tw is None and bool(_bands)
    if _tw is None:
        _tw = _bands[-1] if _bands else None
    _wtag = (f"摊铺宽度{_tw:g}m以内优先（桥梁宽度 {_w:g}m）"
             + ("；宽度超库内最大档，按最大档套用" if _w_over else "")
             if _tw is not None else "")
    for e in entries:
        ef = _extract_features(e["name"])
        s, reasons, ok = _score(feats, ef, item.get("mat_unit") or item.get("unit"),
                                e["unit"], item_group=item.get("fenxiang", ""),
                                item_mat=item.get("mat") or "")
        # 沥青混凝土：机械摊铺子目优先（定额单位 t）、粒式吻合再加分，
        # 人工摊铺/混凝土类子目不套；粘层等非混凝土沥青按文本打分不调整
        if ("沥青" in _mat_n and "沥青" in e["name"]
                and ("混凝土" in _mat_n or "粒式" in _mat_n)):
            _grain = next((k for k in ("粗粒", "中粒", "细粒") if k in _mat_n), None)
            if "机械摊铺" in e["name"]:
                s += 25 if _grain and _grain in e["name"] else 15
                if _tw is not None and _wd_is(e["name"], _tw):
                    s += 10
                    reasons.insert(0, f"沥青混凝土套机械摊铺（用户口径），{_wtag}，"
                                    "定额单位 t（说明第8条）")
                else:
                    reasons.insert(0, "沥青混凝土套机械摊铺（用户口径），定额单位 t（说明第8条）")
            elif "摊铺" in e["name"]:
                s -= 15
                reasons.insert(0, "沥青混凝土套机械摊铺，人工摊铺不套（用户口径）")
            elif "混凝土" in e["name"]:
                s -= 15
                reasons.insert(0, "沥青混凝土套机械摊铺，混凝土类子目不套（用户口径）")
            if ("改性" in e["name"]) != _mod:
                s -= 15
                reasons.insert(0, ("量表含 改性/SBS → 套改性沥青混凝土，不套普通沥青（用户口径）"
                                   if _mod
                                   else "量表无 改性/SBS → 套普通沥青混凝土，不套改性沥青（用户口径）"))
        cands.append({"entry": e, "score": s, "reasons": reasons, "unit_ok": ok})
    cands.sort(key=lambda c: -c["score"])
    return cands


def _find_candidates(item, entries, rules, bridge_width=8.5):
    """按材料语义在规则组（优先）与全库中找候选。
    bridge_width：桥梁宽度（m），沥青摊铺宽度档取 ≥W 的最小档（默认 8.5m 以内）。

    返回 [{"entry": 定额条目副本, "boost": 规则加分, "boost_reason": str}, ...]
    """
    mat = _norm_text(item.get("mat") or "")
    fx = _norm_text(item["fenxiang"])
    group = _group_of(item, rules)
    rule_rows = rules.get(group) or []
    fam_kw = _GROUP_FAMILY_KW.get(group, "")
    pool = []

    def add_entry(e, boost=0, reason=""):
        pool.append({"entry": dict(e), "boost": boost, "boost_reason": reason})

    def add_rule_rows(class_token, boost_reason, boost_pred=None):
        """规则样例行：名称含类别词（混凝土/螺纹钢…）的给予规则加分。

        boost_pred(name) 可额外约束（如伸缩缝型号一致才加分）。
        """
        for rr in rule_rows:
            e = next((x for x in entries if x["code"] == rr["code"]), None)
            if e is None:
                e = {"code": rr["code"], "name": rr["name"], "unit": rr["unit"],
                     "price": None, "row": 0}
            nm = _norm_text(rr["name"])
            boost = 8 if class_token and class_token in nm and (boost_pred is None or boost_pred(nm)) else 0
            add_entry(e, boost, boost_reason if boost else "")

    def add_library(pred, boost=0, reason=""):
        for e in entries:
            if pred(e):
                add_entry(e, boost, reason)

    m_c = re.search(r"C\d{2}", mat)
    demo = (re.search(r"拆除|铣刨|凿除|切除", mat + fx)
            and not re.search(r"重浇|浇筑|浇注|灌注|现浇|预制", mat))
    if demo:
        # 拆除/铣刨/凿除类条目：优先拆除/铣刨定额（如 1-3-47/48 拆除混凝土结构、
        # 1-3-8/9 铣刨沥青混凝土路面），浇筑/摊铺类子目不套（重浇除外）
        add_library(lambda e: re.search(r"拆除|铣刨|凿除", e["name"]), 14,
                    "拆除类条目：优先拆除/铣刨定额")
        add_library(lambda e: re.search(r"现浇|预制|浇筑|灌注|摊铺", e["name"]), -14,
                    "拆除类条目：不套浇筑/摊铺类定额")
    elif "防水" in mat:
        # 防水混凝土/防水层 → 防水相关子目
        add_rule_rows("防水", f"样例规则指向（{group}）")
        add_library(lambda e: "防水" in e["name"], 0, "")
    elif m_c:
        # Cxx 混凝土材料 → 现浇混凝土定额（强度精确匹配优先）
        # 规则加分仅给浇筑类子目（灌注/现浇/预制/浇筑），模板/凿除等配套子目不抬分
        add_rule_rows("混凝土", f"样例规则指向（{group}）",
                      boost_pred=lambda nm: not re.search(r"模板|凿除|拆除|凿除桩顶", nm)
                      and re.search(r"灌注|现浇|预制|浇筑", nm) is not None)
        if group == "伸缩缝":
            # 伸缩缝后浇混凝土（如 C50混凝土 或 C50钢纤维混凝土，秦滨矩阵式分两行）：
            # 库内无伸缩缝混凝土定额，引钢纤维混凝土子目（如 3-3-55 桥面铺装）作参考候选
            add_library(lambda e: "钢纤维混凝土" in e["name"], 6,
                        "伸缩缝后浇混凝土：库内无伸缩缝混凝土定额，引钢纤维混凝土子目参考（强度不符注明）")
        # 全库补强：同强度 + 工艺（现浇/预制/灌注）+ 同构件类别
        def pred(e, m=m_c):
            n = e["name"]
            return ("混凝土" in n and m.group() in n
                    and re.search(r"现浇|预制|灌注", n)
                    and (not fam_kw or re.search(fam_kw, n)))
        add_library(pred, 0, "")
    elif "沥青" in mat:
        add_rule_rows("沥青", f"样例规则指向（{group}）")
        if "混凝土" in mat or "粒式" in mat:
            # 沥青混凝土（用户口径）：机械摊铺子目优先、粒式吻合再加分、
            # 摊铺宽度档按桥梁宽度取 ≥W 的最小档（默认 8.5m 以内）；
            # 改性（量表含 改性/SBS）→ 套改性沥青混凝土，否则套普通沥青混凝土，
            # 不符者降分；人工摊铺/混凝土类子目不套；定额单位 t（说明第8条）
            _grain = next((k for k in ("粗粒", "中粒", "细粒") if k in mat), None)
            _mod = bool(re.search(r"改性|sbs", mat, re.I))
            _bands = _asphalt_width_bands(entries)
            _w = bridge_width or 8.5
            _tw = next((b for b in _bands if b >= _w), None)
            _w_over = _tw is None and bool(_bands)
            if _tw is None:
                _tw = _bands[-1] if _bands else None
            _wtag = (f"摊铺宽度{_tw:g}m以内优先（桥梁宽度 {_w:g}m）"
                     + ("；宽度超库内最大档，按最大档套用" if _w_over else "")
                     if _tw is not None else "")
            _e_mod = lambda e: "改性" in e["name"]
            _e_wd = lambda e: _tw is not None and _wd_is(e["name"], _tw)
            add_library(lambda e: "沥青" in e["name"] and "机械摊铺" in e["name"]
                        and _e_mod(e) == _mod and _e_wd(e)
                        and (_grain is None or _grain in e["name"]), 35,
                        "沥青混凝土套机械摊铺（用户口径），" + _wtag
                        + "，定额单位 t（说明第8条）")
            add_library(lambda e: "沥青" in e["name"] and "机械摊铺" in e["name"]
                        and _e_mod(e) == _mod and not _e_wd(e)
                        and (_grain is None or _grain in e["name"]), 25,
                        "沥青混凝土套机械摊铺（用户口径），定额单位 t（说明第8条）")
            add_library(lambda e: "沥青" in e["name"] and "机械摊铺" in e["name"]
                        and _e_mod(e) == _mod
                        and _grain is not None and _grain not in e["name"], 15,
                        "沥青混凝土套机械摊铺（用户口径），定额单位 t（说明第8条）")
            # 改性不符（含机械/人工摊铺等）：-15
            add_library(lambda e: "沥青" in e["name"] and "摊铺" in e["name"]
                        and _e_mod(e) != _mod, -15,
                        "量表含 改性/SBS → 套改性沥青混凝土（用户口径）" if _mod
                        else "量表无 改性/SBS → 套普通沥青混凝土（用户口径）")
            add_library(lambda e: "沥青" in e["name"] and "摊铺" in e["name"]
                        and "机械摊铺" not in e["name"] and _e_mod(e) == _mod, -15,
                        "沥青混凝土套机械摊铺，人工摊铺不套（用户口径）")
            add_library(lambda e: "沥青" in e["name"] and "摊铺" not in e["name"]
                        and "混凝土" in e["name"], -15,
                        "沥青混凝土套机械摊铺，混凝土类子目不套（用户口径）")
            add_library(lambda e: "沥青" in e["name"] and "摊铺" not in e["name"]
                        and "混凝土" not in e["name"], 0, "")
        else:
            # 粘层/透层/沥青砂等：不抬机械摊铺分，按文本打分
            add_library(lambda e: "沥青" in e["name"], 0, "")
    elif any(k in mat for k in ("涂装", "涂料", "漆")):
        # 涂装/涂料/漆类材料：引 3-8 抹灰油/乳胶漆/水质涂料 与 3-9 钢结构重防腐子目
        add_library(lambda e: "涂料" in e["name"] or "漆" in e["name"], 20,
                    "涂装类材料引涂装/油漆子目")
    elif re.search(r"HPB|HRB|CRB|Q235|钢筋", mat):
        feats = _extract_features(mat)
        st = feats["specs"].get("steel_type", "")
        # 规则加分仅给钢种相符的样例行（HRB400 只抬 螺纹钢 行）
        add_rule_rows("钢筋", f"样例规则指向（{group}）",
                      boost_pred=lambda nm, st=st: bool(st) and st in nm)
        if st:
            add_library(lambda e, st=st: "钢筋" in e["name"] and st in e["name"], 0, "")
        else:
            add_library(lambda e: "钢筋" in e["name"]
                        and ("螺纹钢" in e["name"] or "3号钢" in e["name"]), 0, "")
        if group == "桩基础" and st in ("3号钢", "螺纹钢"):
            # 桩基钢筋按设计要求锚定：HPB300/Q235→3-2-6（3号钢 φ10以外）、
            # HRB400/CRB550→3-2-7（螺纹钢 φ14以外）
            want = "3-2-6" if st == "3号钢" else "3-2-7"
            dia_w = "φ10以外" if st == "3号钢" else "φ14以外"
            for p in pool:
                if str(p["entry"]["code"]).split("*")[0].strip() == want:
                    p["boost"] = 30
                    p["boost_reason"] = (f"桩基钢筋：{mat} → {want}"
                                         f"（{st} {dia_w}，设计要求）")
    elif "钢绞线" in mat:
        add_rule_rows("钢绞线", f"样例规则指向（{group}）")
        add_library(lambda e: "钢绞线" in e["name"], 0, "")
    elif "支座" in mat or "支座" in fx:
        add_rule_rows("支座", f"样例规则指向（{group}）")
        add_library(lambda e: "支座" in e["name"], 0, "")
    elif "伸缩" in mat or "伸缩" in fx:
        # 型号一致（KF-80 ↔ TS80）才给样例规则加分，避免型号不符的子目被抬分
        tn = _extract_features(mat)["specs"].get("type_no")
        add_rule_rows("伸缩缝", f"样例规则指向（{group}）",
                      boost_pred=lambda nm, tn=tn: tn is None
                      or _extract_features(nm)["specs"].get("type_no") == tn)
        add_library(lambda e: "伸缩缝" in e["name"]
                    and "沥青" not in e["name"] and "麻丝" not in e["name"], 0, "")
    elif group == "土方":
        # 挖方/填方：按语义锚定样例规则（挖方→开槽挖土、填方→填压土方）
        kw = ("挖" if "挖方" in fx or "挖方" in mat
              else "填" if "填方" in fx or "填方" in mat else None)
        for rr in rule_rows:
            e = next((x for x in entries if x["code"] == rr["code"]), None)
            if e is None:
                e = {"code": rr["code"], "name": rr["name"], "unit": rr["unit"],
                     "price": None, "row": 0}
            nm = _norm_text(rr["name"])
            boost = 25 if kw and kw in nm else 0
            add_entry(e, boost, f"样例规则指向（{group}）" if boost else "")
    else:
        add_rule_rows(None, "")
    # 去重（编号+名称+单位）：同一条目重复入池时保留加分更高的一条
    # （如 2-2-91/93 既命中规则行又命中机械摊铺库检索，需保留 15 而非 8）
    seen, out = {}, []
    for p in pool:
        e = p["entry"]
        k = (e["code"], e["name"], e["unit"])
        prev = seen.get(k)
        if prev is None:
            seen[k] = len(out)
            out.append(p)
        elif p["boost"] > out[prev]["boost"]:
            out[prev] = p
    return out


def _match_pile_params(item, entries, rules, soil=""):
    """桩基参数（根数/桩径/桩长）→ 钻孔定额 + 样例组合。
    soil：页面选择的地层条件（如 砂土/黏土、砂砾、岩石，默认 砂土/黏土），
    名称含所选地层的钻孔子目加分（2026-10-27 用户口径）。"""
    dia_mm = item["dia"] * 1000 if item["dia"] < 10 else item["dia"]
    depth = item["depth"]
    total_m = item["count"] * depth
    cands = []
    _soil_kws = tuple(k for k in (soil or "").split("/") if k)

    def _add_cand(e):
        ef = _extract_features(e["name"])
        specs = ef["specs"]
        if "钻孔" not in ef["methods"] or "灌注" in ef["methods"]:
            return
        if "dia_le" not in specs or "depth_le" not in specs:
            return
        if specs["dia_le"] < dia_mm or specs["depth_le"] < depth:
            return
        score, reasons = 10, [f"桩径 φ{dia_mm:g} ≤ φ≤{specs['dia_le']:g}"]
        if specs["dia_le"] == dia_mm:
            score += 8
            reasons[0] = f"桩径一致 φ≤{specs['dia_le']:g}"
        if specs["depth_le"] == depth:
            score += 6
            reasons.append(f"深度一致 h≤{specs['depth_le']:g}m")
        else:
            score += 3
            reasons.append(f"深度 h≤{depth:g}m 在定额 h≤{specs['depth_le']:g}m 范围内")
        cso = ef["conds"] & set(SOIL)
        # 已选地层时，名称所含地层与所选完全无交集的钻孔子目不进候选
        # （用户 2026-10-27 口径）：选 砂土/黏土 时「砂土、黏土」合并名含所选地层 → 入选，
        # 仅 砂砾/岩石 等不含所选地层的变体排除
        if _soil_kws and cso and not (cso & set(_soil_kws)):
            return
        if cso:
            reasons.append("定额适用地层：" + "、".join(sorted(cso)))
        if _soil_kws and (cso & set(_soil_kws)):
            score += 8
            reasons.append(f"地层一致（页面选择：{soil}）")
        score += min(len(_bigrams(_norm_text(item["fenxiang"])) & _bigrams(ef["text"])), 15)
        cands.append({"entry": e, "score": score, "reasons": reasons, "unit_ok": True})

    # 库中缺少小桩径钻孔变体（如 φ≤1800 的 3-1-32*0.87 换乘系数子目）时，
    # 从样例规则补入：单位借同号基础子目，单价按换乘系数折算
    extended = list(entries)
    for rr in rules.get("桩基础") or []:
        if "钻孔" not in _norm_text(rr["name"]):
            continue
        if any(str(x["code"]) == rr["code"] for x in entries):
            continue
        base = str(rr["code"]).split("*")[0].strip()
        be = next((x for x in entries if str(x["code"]) == base), None)
        price = (round(be["price"] * 0.87, 2)
                 if be and be["price"] is not None and "0.87" in rr["name"]
                 else be["price"] if be else None)
        extended.append({"code": rr["code"], "name": rr["name"],
                         "unit": be["unit"] if be else rr["unit"],
                         "price": price, "row": 0})
    for e in extended:
        _add_cand(e)
    if not cands:
        # 桩长超出库内所有深度档、或库内无该桩径档（如 φ1800 h64：库内孔径档到 φ≤2000，
        # 无 φ≤1800 档，且 64m 超出 h≤60m）：取能覆盖该桩径的最小孔径档中深度最大者套用并标注，
        # 工程量仍按实际桩长计算
        best = None
        best_dia = None
        for e in extended:
            ef = _extract_features(e["name"])
            specs = ef["specs"]
            if ("钻孔" not in ef["methods"] or "灌注" in ef["methods"]):
                continue
            if "depth_le" not in specs or specs["dia_le"] < dia_mm:
                continue
            cso = ef["conds"] & set(SOIL)
            if _soil_kws and cso and not (cso & set(_soil_kws)):
                continue  # 兜底取最大档时同样排除与所选地层无交集的变体
            if (best is None or specs["dia_le"] < best_dia
                    or (specs["dia_le"] == best_dia and specs["depth_le"] > best[0])):
                best = (specs["depth_le"], e)
                best_dia = specs["dia_le"]
        if best:
            e = best[1]
            specs = _extract_features(e["name"])["specs"]
            reasons = [f"桩径 φ{dia_mm:g} 在定额 φ≤{specs['dia_le']:g} 范围内"
                       + ("（库内无 φ≤1800 档，借用 φ≤2000 档）" if dia_mm == 1800 else "")]
            reasons.append(f"桩长 {depth:g}m 超出库内该孔径钻孔定额最大深度"
                           f" h≤{specs['depth_le']:g}m，按最大档套用（工程量按实际桩长）")
            cands.append({"entry": e, "score": 20, "reasons": reasons,
                          "unit_ok": True})
    # 孔径档取紧（用户 2026-10-27 口径）：孔径档为桩径区间上界（φ≤1200 即 1000＜φ≤1200），
    # 只保留能覆盖桩径的最小孔径档——φ1m 桩只套 φ≤1000，φ≤1200/1500/2000 等更大档不适用
    if cands:
        _mdia = min(_extract_features(c["entry"]["name"])["specs"]["dia_le"]
                    for c in cands)
        _tight = [c for c in cands
                  if _extract_features(c["entry"]["name"])["specs"]["dia_le"] == _mdia]
        if len(_tight) < len(cands):
            for c in _tight:
                c["reasons"].append(
                    f"孔径档取紧：桩径 φ{dia_mm / 1000:g}m 套最小可用孔径档 φ≤{_mdia:g}"
                    f"（孔径档为桩径区间上界，更大档不适用）")
        cands = _tight
    # 深度档取紧（用户 2026-10-27 口径）：桩长=桩总长度÷桩数（工程数量表汇总
    # 解析时已算出 depth），只保留能覆盖桩长的最小深度档——≤40m 套 h≤40、
    # 40<L≤60m 套 h≤60；L>60m 走上方「无候选」兜底取最大档 h≤60 并标注
    if cands:
        _md = min(_extract_features(c["entry"]["name"])["specs"]["depth_le"]
                  for c in cands)
        _tight = [c for c in cands
                  if _extract_features(c["entry"]["name"])["specs"]["depth_le"] == _md]
        if len(_tight) < len(cands):
            for c in _tight:
                c["reasons"].append(
                    f"深度档取紧：桩长 {depth:g}m（桩总长度÷桩数）套最小可用档 h≤{_md:g}m")
        cands = _tight
    cands.sort(key=lambda c: -c["score"])
    # 钻孔工程量 = 根数 × 桩长（自然米）
    for c in cands:
        c["qty_natural"] = total_m
        c["qty_de"] = total_m / _unit_parts(c["entry"]["unit"])[1]
        c["amount"] = round(c["qty_de"] * c["entry"]["price"], 2) if c["entry"]["price"] is not None else None
    # 样例组合
    combo = []
    for rr in rules.get("桩基础") or []:
        nm = _norm_text(rr["name"])
        disp_code = rr["code"] + (" 换" if rr.get("huan") else "")
        st, need = "需补充", ""
        if "钢护筒" in nm:
            m = re.search(r"φ(\d+)", nm)
            if m and float(m.group(1)) == dia_mm and "陆地" in rr["name"]:
                st = "已匹配"
                need = f"工程量=桩数×2m（每根桩 2m 护筒，{item['count']:g}根×2m）"
            else:
                st = "不适用"
                need = (f"桩径不同（定额 φ{m.group(1)}）" if m
                        else "水中埋设，陆地不套")
        elif "回旋钻机钻孔" in nm:
            st = "已匹配" if any(c["entry"]["code"] == rr["code"] for c in cands[:MAX_CANDIDATES]) else "不适用"
        elif "泥浆" in nm or "灌注桩混凝土" in nm:
            need = "工程量=桩基混凝土量（随桩基混凝土条目套用）"
        elif "黏土" in nm:
            need = "随泥浆制作套用（材料行）"
        elif "声测管" in nm:
            need = "声测管长度（m）；库内无此子目，按模板借 2-33"
        elif "凿除" in nm:
            st = "已匹配"
            need = "工程量=桩数×0.5m×截面积（分部内所有桩合计，套 3-10-44 一次）"
        elif "钢筋" in nm:
            need = "桩基钢筋用量及直径 φ10/φ14 内外"
        elif "运" in nm or "废料" in nm:
            st = "已匹配"
            need = ("凿桩头弃运：1-6-71=凿桩头量×2（建筑垃圾按虚方），"
                    "1-6-72=×2×(运距-1)，运距在页面设置")
        combo.append({"code": disp_code, "name": rr["name"], "unit": rr["unit"],
                      "status": st, "need": need})
    # 陆地埋设钢护筒：每根桩 2m，工程量 = 桩数×2（定额单位 10m）
    liner = None
    for e in entries:
        nm = _norm_text(e["name"])
        if "陆地埋设钢护筒" in nm and f"φ{dia_mm:g}" in nm:
            liner = dict(e)
            liner["qty_natural"] = item["count"] * 2
            liner["qty_de"] = liner["qty_natural"] / _unit_parts(e["unit"])[1]
            liner["amount"] = (round(liner["qty_de"] * e["price"], 2)
                               if e["price"] is not None else None)
            break
    # 水中变体仅作提示（如实际水中施工改套）
    supplement = []
    for e in entries:
        if "水中埋设钢护筒" in e["name"] and f"φ{dia_mm:g}" in _norm_text(e["name"]):
            supplement.append(e)
    return cands, combo, supplement, liner


def _asphalt_thickness(item, grain):
    """沥青混凝土设计厚度（m）。用户 2026-10-27 口径：从数量表里读，找不到回退
    细粒0.04/中粒0.06。表内来源：备注/分项名/材料名里的厚度字样（4cm、0.04m厚等）。
    返回 (厚度 或 None, 来源说明)。"""
    txt = _norm_text((item.get("remark") or "") + (item.get("fenxiang") or "")
                     + (item.get("mat") or ""))
    # 带「厚」字样的 cm/米 显式厚度（如 厚4cm / 4cm厚 / 厚0.04m / 0.04m厚）
    m = re.search(r"厚(?:度)?\s*(\d+(?:\.\d+)?)\s*(?:cm|厘米)"
                  r"|(\d+(?:\.\d+)?)\s*(?:cm|厘米)\s*厚"
                  r"|厚(?:度)?\s*(\d+(?:\.\d+)?)\s*[m米](?![0-9²³])"
                  r"|(\d+(?:\.\d+)?)\s*[m米](?![0-9²³])\s*厚", txt)
    if m:
        v = float(next(g for g in m.groups() if g is not None))
        if m.group(1) or m.group(2):
            return v / 100.0, f"表内厚度 {v:g}cm"
        return v, f"表内厚度 {v:g}m"
    # 无「厚」字样的裸 cm 值（如备注仅写 4cm），限定合理范围防误读
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:cm|厘米)", txt)
    if m and 0 < float(m.group(1)) <= 30:
        v = float(m.group(1))
        return v / 100.0, f"表内厚度 {v:g}cm"
    return ({"细粒": 0.04, "中粒": 0.06}.get(grain),
            {"细粒": "细粒式默认4cm", "中粒": "中粒式默认6cm"}.get(grain, ""))


def _match_one(idx, item, entries, code_index, rules, all_items, span_len=30.0,
                soil="", bridge_width=8.5):
    status, matches, combo, supplement, checklist, notes = "none", [], [], [], [], []
    mud, cage = [], None
    feats = _extract_features(_item_display(item))
    if item["kind"] == "桩基参数":
        cands, combo, supplement, liner = _match_pile_params(item, entries, rules, soil)
        status, picked, top, note = _classify(cands)
        matches = top
        checklist = _checklist(item, soil=soil)
        if note:
            notes.append(note)
        if not cands:
            notes.append(f"桩径 φ{item['dia']:g}/桩长 {item['depth']:g}m 超出库内回旋钻机钻孔定额范围，"
                         "无法匹配，请核对桩参数"
                         + ("。" if soil else "或补充地层条件。"))
        else:
            notes.append(f"钻孔工程量 = Σ根数×桩长 = {item['count']:g}×{item['depth']:g} = "
                         f"{item['count'] * item['depth']:g}m；陆地埋设钢护筒按每根桩 2m 套用，"
                         f"工程量 = {item['count']:g}根×2m = {item['count'] * 2:g}m；"
                         "泥浆/灌注桩混凝土随桩基混凝土条目套用；"
                         "凿桩头按 0.5m/根计，分部内所有桩合计套 3-10-44 一次；"
                         "凿桩头弃运套 1-6-71/1-6-72（建筑垃圾按虚方×2，"
                         "1-6-72 次数=页面运距-1）。")
            if status in ("single", "multi") and supplement:
                for s in supplement:
                    notes.append(f"如水中施工改套：{s['code']} {s['name']}（{s['unit']}）")
        return {"idx": idx, "item": item, "status": status, "matches": matches,
                "rule_combo": combo, "checklist": checklist, "notes": notes,
                "pile_liner": liner}

    mat = _norm_text(item.get("mat") or "")
    if "面积" in mat and not re.search(r"铣刨|拆除|切除|凿除", mat):
        # 面积类列（桥梁面积等）为指标数据，非定额计量子目
        return {"idx": idx, "item": item, "status": "none", "matches": [],
                "rule_combo": [], "checklist": [],
                "notes": ["「面积」列为指标数据，非定额计量子目，不参与定额匹配。"]}
    fx_s = _norm_text(item.get("fenxiang") or "")
    # 泄水管（泄水孔）内按重量计的钢材类（钢管/栅盖/扣件/补强钢筋等）→
    # 合并一套预埋铁件 3-2-8（用户 2026-10-26：加在一起一律套预埋铁件；
    # 合并求和写入在 build_result_excel 的 排水 分项）；玻纤/塑料等非钢材不在此列
    if (("泄水" in fx_s or "排水" in fx_s)
            and _unit_parts(item.get("mat_unit") or "")[0] == "kg"
            and not _PLASTIC_PIPE_RE.search(mat)
            and any(k in mat for k in ("钢", "铁", "扣件", "栅盖", "法兰", "管卡"))):
        e = next((x for x in entries
                  if str(x["code"]).split("*")[0].strip() == "3-2-8"), None)
        if e is not None:
            status, matches = "single", [{"entry": e, "score": 40, "unit_ok": True,
                                          "reasons": [f"泄水管内钢材类（{item.get('mat')}）按重量计，"
                                                      "按用户口径套预埋铁件（同类合并计价）"]}]
            notes.append(f"泄水管钢材类（{item.get('mat')}）按用户口径套预埋铁件 3-2-8。")
        else:
            status, matches = "none", []
            notes.append("泄水管钢材类应套预埋铁件，但库内无 3-2-8 子目，需补充。")
        for c in matches:
            mult = _unit_parts(c["entry"]["unit"])[1]
            val = item.get("value") if item.get("value") is not None else 0
            c["qty_natural"] = val
            c["qty_de"] = val / 1000 / mult
            c["note_conv"] = (f"已按 kg→t 换算（{val:g} kg = {c['qty_de']:g} "
                              f"{c['entry']['unit']}）")
            c["amount"] = (round(c["qty_de"] * c["entry"]["price"], 2)
                           if c["entry"]["price"] is not None else None)
        return {"idx": idx, "item": item, "status": status, "matches": matches,
                "rule_combo": [], "checklist": _checklist(item, soil=soil), "notes": notes,
                "mud": [], "cage": None, "tpl": None}
    # 非钢筋钢材（Q235/Q345 等牌号、钢板/型钢类）→ 预埋铁件 3-2-8
    # （用户 2026-10-26 口径：除钢筋外的钢材套预埋铁件，不套钢筋定额）；
    # 钢管φD×t 为声测管材料，走下方借 2-33 分支
    # 钢绞线：束长按一跨桥长（页面输入，默认 30m → 40m 以内档）；
    # 量表有锚具 → 后张法（孔数与锚具一致），无锚具 → 先张法（用户 2026-10-26 口径）
    if "钢绞线" in mat:
        anchors = [x for x in all_items
                   if ("锚" in _norm_text(x.get("mat") or "")
                       or "锚" in _norm_text(x.get("fenxiang") or ""))
                   and "钢绞线" not in _norm_text(x.get("mat") or "")]
        holes = []
        for x in anchors:
            mh = re.search(r"15-(\d+)", _norm_text(x.get("mat") or "")) \
                or re.search(r"(\d+)孔", _norm_text(x.get("mat") or ""))
            if mh:
                holes.append(int(mh.group(1)))
        max_holes = max(holes) if holes else 0
        if anchors:
            band = "20m以内" if span_len <= 20 else "40m以内"
            if max_holes:
                if band == "20m以内":
                    code = "3-5-7" if max_holes <= 3 else "3-5-8" if max_holes <= 7 else "3-5-9"
                    hole_label = ("3孔" if max_holes <= 3
                                  else "7孔以内" if max_holes <= 7 else "12孔以内")
                else:
                    code = "3-5-10" if max_holes <= 7 else "3-5-11" if max_holes <= 12 else "3-5-12"
                    hole_label = ("7孔以内" if max_holes <= 7
                                  else "12孔以内" if max_holes <= 12 else "19孔以内")
                e = next((x for x in entries
                          if str(x["code"]).split("*")[0].strip() == code), None)
                if e is not None:
                    status, matches = "single", [{"entry": e, "score": 45, "unit_ok": True,
                                                  "reasons": [f"量表含锚具（{max_holes}孔），按后张法",
                                                              f"束长按一跨桥长 {span_len:g}m → {band}档",
                                                              f"孔数与锚具一致 → {hole_label}"]}]
                    notes.append(f"钢绞线按后张法套 {e['code']}（锚具 {max_holes} 孔、束长 {span_len:g}m）。")
                else:
                    status, matches = "none", []
                    checklist.append(f"后张法钢绞线应套 {code}，但库内无此子目，需补充")
                    notes.append(f"后张法应套 {code}，但库内无此子目，需补充。")
            else:
                codes = (("3-5-7", "3-5-8", "3-5-9") if band == "20m以内"
                         else ("3-5-10", "3-5-11", "3-5-12"))
                cands = [x for x in entries
                         if str(x["code"]).split("*")[0].strip() in codes]
                if cands:
                    status, matches = "multi", [
                        {"entry": x, "score": 40, "unit_ok": True,
                         "reasons": [f"量表含锚具，按后张法；束长 {span_len:g}m → {band}档"]}
                        for x in cands]
                else:
                    status, matches = "none", []
                checklist.append("钢绞线孔数（量表锚具未标注孔数）")
                notes.append("后张法钢绞线孔数需按锚具孔数补充。")
            if span_len > 40:
                notes.append("束长超库内最大档（40m 以内），按 40m 以内档套用。")
        else:
            e = next((x for x in entries
                      if str(x["code"]).split("*")[0].strip() == "3-5-2"), None)
            if e is not None:
                status, matches = "single", [{"entry": e, "score": 45, "unit_ok": True,
                                              "reasons": ["量表无锚具，按先张法（用户口径）",
                                                          f"束长按一跨桥长 {span_len:g}m（页面可改，默认 30m）"]}]
                notes.append(f"钢绞线按先张法套 3-5-2（量表无锚具；束长 {span_len:g}m）。")
            else:
                status, matches = "none", []
                checklist.append("先张法钢绞线应套 3-5-2，但库内无此子目，需补充")
                notes.append("先张法钢绞线应套 3-5-2，但库内无此子目，需补充。")
        for c in matches:
            mult = _unit_parts(c["entry"]["unit"])[1]
            val = item.get("value") if item.get("value") is not None else 0
            c["qty_natural"] = val
            if _unit_parts(item.get("mat_unit") or "")[0] == "t":
                c["qty_de"] = val / mult  # 量表单位已是 t，不换算
            else:
                c["qty_de"] = val / 1000 / mult
                c["note_conv"] = (f"已按 kg→t 换算（{val:g} kg = {c['qty_de']:g} "
                                  f"{c['entry']['unit']}）")
            c["amount"] = (round(c["qty_de"] * c["entry"]["price"], 2)
                           if c["entry"]["price"] is not None else None)
        return {"idx": idx, "item": item, "status": status, "matches": matches,
                "rule_combo": [], "checklist": checklist, "notes": notes,
                "mud": [], "cage": None, "tpl": None}
    if re.match(r"^Q\d", mat) and "检测管" not in fx_s and "声测管" not in fx_s:
        # 用户 2026-10-26 补充：钢材部分中的 Q235 钢不套定额
        # （声测管块内 Q235B 除外，仍套预埋铁件 3-2-8）
        return {"idx": idx, "item": item, "status": "none", "matches": [],
                "rule_combo": [], "checklist": [f"{item.get('mat')} 钢按用户口径不套定额（声测管除外）"],
                "notes": [f"钢材部分 {item.get('mat')} 按用户口径不套定额。"]}
    if _steel_mat(mat):
        e = next((x for x in entries
                  if str(x["code"]).split("*")[0].strip() == "3-2-8"), None)
        if e is not None:
            status, matches = "single", [{"entry": e, "score": 40, "unit_ok": True,
                                          "reasons": [f"清单材料 {item.get('mat')} 为非钢筋钢材，"
                                                      "按用户口径套预埋铁件定额（不套钢筋定额）"]}]
            notes.append(f"非钢筋钢材（{item.get('mat')}）按用户口径套预埋铁件 3-2-8。")
        else:
            status, matches = "none", []
            notes.append("非钢筋钢材应套预埋铁件，但库内无 3-2-8 子目，需补充。")
        for c in matches:
            mult = _unit_parts(c["entry"]["unit"])[1]
            val = item.get("value") if item.get("value") is not None else 0
            ib = _unit_parts(item.get("mat_unit") or "")[0]
            c["qty_natural"] = val
            if ib == "kg":
                c["qty_de"] = val / 1000 / mult
                c["note_conv"] = (f"已按 kg→t 换算（{val:g} kg = {c['qty_de']:g} "
                                  f"{c['entry']['unit']}）")
            else:
                c["qty_de"] = val / mult
            c["amount"] = (round(c["qty_de"] * c["entry"]["price"], 2)
                           if c["entry"]["price"] is not None else None)
        return {"idx": idx, "item": item, "status": status, "matches": matches,
                "rule_combo": [], "checklist": _checklist(item, soil=soil), "notes": notes,
                "mud": [], "cage": None, "tpl": None}
    # 声测管（钢管φD×t）：按模板借 2-33 灌注桩声测管埋设（定额单位 m；库内无此子目，
    # 单价需补充）。清单为重量 kg 时按用户公式 W=0.02466×(D-t)×t kg/m 换算长度
    if re.match(r"^钢管φ", mat) or (
            "声测管" in mat and re.search(r"[dφ]\d+(?:\.\d+)?\s*[x×]\s*\d+(?:\.\d+)?", mat)):
        w_kgm = _pipe_kg_per_m(mat)
        e = next((x for x in entries if str(x["code"]).split("*")[0].strip() == "2-33"), None)
        if e is None:
            e = {"code": "借2-33", "name": "灌注桩声测管埋设 钢管", "unit": "m",
                 "price": None, "row": 0}
        val = item.get("value") if item.get("value") is not None else 0
        ib = _unit_parts(item.get("mat_unit") or "")[0]
        mult = _unit_parts(e["unit"])[1]
        c = {"entry": e, "score": 40, "unit_ok": True, "qty_natural": val,
             "reasons": ["声测管按模板借 2-33（库内无此子目，单价需补充）"]}
        if ib == "kg" and w_kgm:
            len_m = val / w_kgm
            c["qty_de"] = len_m / mult
            c["conv_pipe"] = {"w_kgm": w_kgm, "val": val}
            c["reasons"].append(f"按公式 W=0.02466×(D-t)×t={w_kgm:g} kg/m，"
                                f"长度={val:g}÷{w_kgm:g}={len_m:g} m")
            notes.append(f"声测管 {mat} 按重量换算长度：{val:g} kg ÷ {w_kgm:g} kg/m"
                         f" = {len_m:g} m（W=0.02466×(D-t)×t）。")
        elif ib == "kg" and not w_kgm:
            c["qty_de"] = val
            c["reasons"].append("未识别钢管规格 φD×t，无法换算长度，工程量按原值")
            notes.append("声测管重量需按规格 φD×t 换算长度（W=0.02466×(D-t)×t kg/m）。")
        else:
            c["qty_de"] = val / mult
        c["amount"] = None
        notes.append(f"声测管 {mat} 按模板借 2-33（m）套用，"
                     f"工程量 {_fmt_num(val)} m；库内无此子目，单价需补充。")
        combo = [{"code": "借2-33", "name": "灌注桩声测管埋设 钢管", "unit": "m",
                  "status": "已匹配", "need": "声测管长度（m），按清单值直接套用"}]
        return {"idx": idx, "item": item, "status": "single", "matches": [c],
                "rule_combo": combo, "checklist": _checklist(item, soil=soil), "notes": notes,
                "mud": [], "cage": None, "tpl": None}
    group = _group_of(item, rules)
    # 拆除/铣刨类条目（分项不在规则组内，如 主梁悬臂改造/旧桥铺装拆除）也走候选池路径：
    # _find_candidates 内 demo 分支优先拆除/铣刨定额、排除浇筑类子目（重浇除外）
    demo = (re.search(r"拆除|铣刨|凿除|切除", mat + fx_s)
            and not re.search(r"重浇|浇筑|浇注|灌注|现浇|预制", mat))
    # 混凝土条目随同套用相应模板定额（模板量=混凝土量×系数，系数默认参考表比例、页面可改）
    tpl = None
    tfi = _template_for(item)
    if tfi:
        e = next((x for x in entries
                  if str(x["code"]).split("*")[0].strip() == tfi["code"]), None)
        if e is not None:
            te = dict(e)
            te["tpl_key"] = tfi["key"]
            te["tpl_default"] = tfi["default"]
            te["conc_val"] = item.get("value")
            tpl = te
    if (group and group in rules) or demo:
        pool = _find_candidates(item, entries, rules, bridge_width)
        cands = []
        for p in pool:
            e = p["entry"]
            ef = _extract_features(e["name"])
            s, reasons, ok = _score(feats, ef, item.get("mat_unit"), e["unit"],
                                    item_group=item.get("fenxiang", ""),
                                    item_mat=item.get("mat") or "")
            if p["boost"]:
                s += p["boost"]
                reasons.insert(0, p["boost_reason"])
            cands.append({"entry": e, "score": s, "reasons": reasons, "unit_ok": ok})
        cands.sort(key=lambda c: -c["score"])
        if not cands:
            cands = _match_text(item, entries, bridge_width)
        cands, _strip_metal = _strip_metal_pipe_cands(item, cands)
        if _strip_metal:
            notes.append(f"材质为塑料类（{item.get('mat')}），已排除钢管/铸铁管定额候选（用户口径）。")
        status, picked, top, note = _classify(cands)
        matches = top
        if note:
            notes.append(note)
        # 桩基混凝土条目（top 为 3-1-60 灌注桩混凝土）：泥浆制作/运输随同套用
        is_pile_conc = (group == "桩基础" and re.search(r"C\d{2}", mat)
                        and matches
                        and str(matches[0]["entry"]["code"]).split("*")[0].strip() == "3-1-60")
        # 钻孔桩钢筋笼入孔 3-2-10：工程量=同部位 HPB300+HRB400 钢筋量之和（kg→t），
        # 由该部位 HPB300 条目携带套用（避免每行重复）
        cage = None
        if group == "桩基础" and mat == "HPB300":
            fx = _norm_text(item["fenxiang"])
            hpb_kg = hrb_kg = 0.0
            for o in all_items:
                if _norm_text(o.get("fenxiang") or "") != fx:
                    continue
                m = _norm_text(o.get("mat") or "")
                v = _to_float(o.get("value")) or 0
                if m == "HPB300":
                    hpb_kg += v
                elif m == "HRB400":
                    hrb_kg += v
            if hpb_kg or hrb_kg:
                e = next((x for x in entries if str(x["code"]) == "3-2-10"), None)
                if e is not None:
                    ce = dict(e)
                    ce["qty_natural"] = hpb_kg + hrb_kg
                    ce["qty_de"] = (hpb_kg + hrb_kg) / 1000 / _unit_parts(e["unit"])[1]
                    ce["amount"] = (round(ce["qty_de"] * e["price"], 2)
                                    if e["price"] is not None else None)
                    ce["cage_parts"] = (hpb_kg, hrb_kg)
                    cage = ce
        combo = []
        for rr in rules.get(group) or []:
            nm = _norm_text(rr["name"])
            disp_code = rr["code"] + (" 换" if rr.get("huan") else "")
            st, need = "需补充", ""
            if any(c["entry"]["code"] == rr["code"] for c in cands[:MAX_CANDIDATES]):
                st = "已匹配" if cands and cands[0]["entry"]["code"] == rr["code"] else "候选"
            if "模板" in nm:
                if tpl and str(rr["code"]).split("*")[0].strip() == tpl["code"]:
                    st = "已匹配"
                    need = (f"模板量=混凝土量×系数（默认 {tpl['tpl_default']:g}，页面可改；"
                            f"定额单位 {tpl['unit']}）")
                else:
                    need = "模板面积（m2）"
            elif "泥浆" in nm:
                if is_pile_conc:
                    st = "已匹配"
                    need = "工程量=桩基混凝土量（随本条目套用，÷10 换算定额单位）"
                else:
                    need = "工程量=桩基混凝土量（随桩基混凝土条目套用）"
            elif "黏土" in nm:
                need = "随泥浆制作套用（材料行）"
            elif "声测管" in nm:
                need = "声测管长度（m）；库内无此子目，按模板借 2-33"
            elif "支架" in nm:
                need = "支架面积（m2）按施工方案"
            elif "垫层" in nm:
                need = "垫层方量（m3）"
            elif "拆除" in nm:
                need = "拆除方量（m3）"
            elif "锚具" in nm:
                need = "锚具套数（按钢绞线孔数）"
            elif "灌浆" in nm:
                need = "管道压浆量（m3）"
            elif "凿" in nm:
                if "梯道" in _norm_text(item["fenxiang"]):
                    st = "需补充"
                    need = "梯道下部输入表未提供桩数（无桩基类型块），凿桩头量需补充"
                else:
                    st = "已匹配"
                    need = "工程量=桩数×0.5m×截面积（分部内所有桩合计，套 3-10-44 一次）"
            elif "钢筋笼入孔" in nm:
                if group == "桩基础" and mat in ("HPB300", "HRB400"):
                    st = "已匹配" if cage else "候选"
                    need = "工程量=HPB300+HRB400 钢筋量之和（kg 按 t 换算）"
                else:
                    need = "钢筋笼重量（kg 按 t 换算）"
            elif "钢筋" in nm:
                if group == "桩基础":
                    need = ("桩基钢筋按设计要求套 φ10/φ14 以外子目（kg 按 t 换算）"
                            if "骨架" in nm else "钢筋笼重量（kg 按 t 换算）")
                else:
                    need = "钢筋用量及直径 φ10/φ14 内外"
            elif "钢绞线" in nm:
                need = "束长/孔数"
            elif "波纹管" in nm:
                need = "管道直径与长度"
            elif "运" in nm or "废料" in nm:
                if group == "桩基础":
                    if "梯道" in _norm_text(item["fenxiang"]):
                        st = "需补充"
                        need = "梯道下部无桩数信息，凿桩头弃运量需补充"
                    else:
                        st = "已匹配"
                        need = ("凿桩头弃运：1-6-71=凿桩头量×2（建筑垃圾按虚方），"
                                "1-6-72=×2×(运距-1)，运距在页面设置")
                else:
                    need = "弃运方量（按样例规则）"
            combo.append({"code": disp_code, "name": rr["name"], "unit": rr["unit"],
                          "status": st, "need": need})
        checklist = _checklist(item, soil=soil)
        # 泥浆制作 3-1-58 / 泥浆运输 3-1-59：工程量=桩基混凝土量（定额单位 10m3）
        if is_pile_conc:
            val = item.get("value") or 0
            for code in ("3-1-58", "3-1-59"):
                e = next((x for x in entries if str(x["code"]) == code), None)
                if e is None:
                    continue
                me = dict(e)
                me["qty_natural"] = val
                me["qty_de"] = val / _unit_parts(e["unit"])[1]
                me["amount"] = (round(me["qty_de"] * e["price"], 2)
                                if e["price"] is not None else None)
                mud.append(me)
            notes.append("泥浆制作 3-1-58 / 泥浆运输 3-1-59 随灌注桩混凝土套用，"
                         f"工程量=桩基混凝土量 {_fmt_num(val)} m3。")
        if cage:
            hpb_kg, hrb_kg = cage["cage_parts"]
            notes.append("钻孔桩钢筋笼入孔 3-2-10 随本条目套用，工程量=HPB300+HRB400 钢筋量之和 "
                         f"{_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)} kg（÷1000 换 t）。")
    else:
        cands = _match_text(item, entries, bridge_width)
        cands, _strip_metal = _strip_metal_pipe_cands(item, cands)
        if _strip_metal:
            notes.append(f"材质为塑料类（{item.get('mat')}），已排除钢管/铸铁管定额候选（用户口径）。")
        status, picked, top, note = _classify(cands)
        matches = top
        if note:
            notes.append(note)
        checklist = _checklist(item, soil=soil)
        if group:
            notes.append(f"分项「{item['fenxiang']}」无套定额样例规则，已按全库文本匹配。")
        if tpl and not any(str(rr.get("code", "")).split("*")[0].strip() == tpl["code"]
                           for rr in combo):
            combo.append({"code": tpl["code"], "name": tpl["name"], "unit": tpl["unit"],
                          "status": "已匹配",
                          "need": (f"模板量=混凝土量×系数（默认 {tpl['tpl_default']:g}，"
                                   f"页面可改；定额单位 {tpl['unit']}）")})

    # 语义补充说明
    fx = _norm_text(item["fenxiang"])
    # 沥青混凝土：定额单位 t，按定额说明第8条 m3→t 密度换算、第9条厚度附加
    _ASPHALT_DENS = (("粗粒", 2.36), ("中粒", 2.35), ("细粒", 2.30))
    asphalt_conv = None  # {"dens","grain","extra","thick","thick_src"}：m3→t 换算参数
    if "沥青" in mat and matches:
        for k, d in _ASPHALT_DENS:
            if k in mat:
                asphalt_conv = {"dens": d, "grain": k, "extra": 0.0}
                break
        # 第9条 厚度附加：看同分项块内其他沥青层与基层、面积条目
        peers = [x for x in all_items
                 if (x.get("fenxiang") or "") == item.get("fenxiang") and x is not item]
        has_coarse = any("粗粒" in (x.get("mat") or "") for x in peers)
        has_mid = any("中粒" in (x.get("mat") or "") for x in peers)
        extra = 0.0
        extra_tag = ""
        if asphalt_conv and asphalt_conv["grain"] == "细粒" and (has_coarse or has_mid):
            extra = 0.003
            extra_tag = "铺筑于粗粒/中粒式沥青混凝土上，厚度另加0.3cm"
        elif asphalt_conv and asphalt_conv["grain"] == "中粒" and has_coarse:
            extra = 0.003
            extra_tag = "铺筑于粗粒式沥青混凝土上，厚度另加0.3cm"
        elif asphalt_conv and asphalt_conv["grain"] in ("粗粒", "中粒"):
            base_txt = _norm_text((item.get("fenxiang") or "")
                                  + (item.get("remark") or ""))
            if any(k in base_txt for k in ("石灰土", "多合土", "两渣")):
                extra = 0.005
                extra_tag = "铺筑于石灰土/多合土/两渣基层，厚度另加0.5cm"
            elif "渣石" in base_txt:
                extra = 0.01
                extra_tag = "铺筑于渣石基层，厚度另加1cm"
            elif "桥面铺装" in base_txt:
                extra = 0.005
                extra_tag = "桥面铺装底层沥青厚度另加0.5cm"
        if asphalt_conv:
            asphalt_conv["extra"] = extra
            # 设计厚度（用户 2026-10-27：从数量表里读，找不到回退细粒0.04/中粒0.06；
            # 粗粒式等无默认者用方量÷同分项 m2 面积条目推导）
            thick, thick_src = _asphalt_thickness(item, asphalt_conv["grain"])
            if thick is None:
                area_it = next((x for x in peers
                                if _unit_parts(x.get("mat_unit") or "")[0] == "m2"
                                and x.get("value")), None)
                if area_it:
                    thick = round((item.get("value") or 0) / area_it["value"], 4)
                    thick_src = f"方量÷同分项「{area_it.get('mat')}」面积推导"
            asphalt_conv["thick"] = thick
            asphalt_conv["thick_src"] = thick_src
            if thick:
                t2 = round(thick + extra, 4)
                notes.append(f"沥青混凝土工程量=方量/{thick:g}×{t2:g}"
                             f"×{asphalt_conv['dens']:g} t（设计厚度{thick:g}m，"
                             f"来源：{thick_src}"
                             + (f"；{extra_tag}" if extra_tag else "")
                             + "）。")
            else:
                notes.append("沥青混凝土设计厚度未在表内找到，按清单方量直接乘密度"
                             "（定额说明第9条厚度附加未计入）。")
        if "玄武岩" in mat:
            notes.append("玄武岩沥青密度不同时可按实际计算（定额说明第8条）。")
        if asphalt_conv is None and _unit_parts(item.get("mat_unit") or "")[0] == "m3":
            notes.append("沥青混凝土定额按 t 计，清单为 m3，需按密度换算"
                         "（粗粒2.36/中粒2.35/细粒2.30 t/m3，定额说明第8条）。")
    if re.search(r"KF|TS", mat) and status == "single" and matches:
        if "TS" in matches[0]["entry"]["name"] and "KF" in mat:
            notes.append(f"清单为 KF 型定型产品，定额 {matches[0]['entry']['code']} 为 TS 型制作、"
                         "安装子目，伸缩量规格一致，型式不同请人工确认。")
    if status == "single" and matches:
        m_c = re.search(r"C\d{2}", mat)
        m_e = re.search(r"C\d{2}", matches[0]["entry"]["name"])
        if m_c and m_e and m_c.group() != m_e.group():
            notes.append(f"清单混凝土 {m_c.group()} 与定额 {m_e.group()} 强度不同，"
                         "已按构件/工艺匹配，套用请人工确认或换算。")
    if status == "none" and matches:
        notes.append("以下为库内较接近的定额子目（仅供参考，不得直接套用），"
                     "请按补充清单完善特征后重新匹配。")

    # 定额工程量与合价
    for c in matches:
        mult = _unit_parts(c["entry"]["unit"])[1]
        val = item.get("value") if item.get("value") is not None else 0
        ib = _unit_parts(item.get("mat_unit") or "")[0]
        eb = _unit_parts(c["entry"]["unit"])[0]
        qty_de = val / mult
        c["qty_natural"] = val
        if ib and eb and {ib, eb} == {"kg", "t"} and ib == "kg":
            qty_de = val / 1000 / mult
            c["qty_de"] = qty_de
            c["note_conv"] = f"已按 kg→t 换算（{val:g} kg = {qty_de:g} {c['entry']['unit']}）"
        elif asphalt_conv and ib == "m3" and eb == "t" and "沥青" in c["entry"]["name"]:
            ac = asphalt_conv
            if ac.get("thick"):  # 方量/设计厚度×(设计厚度+附加)×密度（用户 2026-10-27 公式）
                qty_de = round(val / ac["thick"] * (ac["thick"] + ac["extra"])
                               * ac["dens"] / mult, 4)
            else:
                qty_de = round(val * ac["dens"] / mult, 4)
            c["qty_de"] = qty_de
            c["asphalt_conv"] = ac
            parts = [f"已按 m3→t 换算（{ac['grain']}式 ×{ac['dens']:g} t/m3，"
                     f"定额说明第8条）"]
            if ac.get("thick"):
                if ac.get("extra"):
                    parts.append(f"设计厚度{ac['thick']:g}m，另加{ac['extra'] * 100:g}cm（第9条）")
                else:
                    parts.append(f"设计厚度{ac['thick']:g}m")
            c["note_conv"] = "；".join(parts)
        else:
            c["qty_de"] = qty_de
        if "沥青" in mat and ib and eb and ib == "m3" and eb == "t" and not asphalt_conv:
            c["amount"] = None  # 无密度不计算合价
        else:
            c["amount"] = round(qty_de * c["entry"]["price"], 2) if c["entry"]["price"] is not None else None
    if tpl:
        notes.append(f"模板 {tpl['code']} {tpl['name']} 随本条目套用，"
                     f"模板量=混凝土量×系数（默认 {tpl['tpl_default']:g}，页面可改）。")
    return {"idx": idx, "item": item, "status": status, "matches": matches,
            "rule_combo": combo, "checklist": checklist, "notes": notes,
            "mud": mud, "cage": cage, "tpl": tpl}


def match_quotas(items, lib, span_len=30.0, soil="", bridge_width=8.5) -> list[dict]:
    """span_len：钢绞线束长/一跨桥长（m，页面输入，默认 30m → 套 40m 以内档）。
    bridge_width：桥梁宽度（m），沥青摊铺宽度档取 ≥W 的最小档（默认 8.5m 以内）。"""
    entries = lib["entries"]
    rules = lib.get("rules") or {}
    code_index = {}
    for e in entries:
        code_index.setdefault(e["code"], []).append(e)
    return [_match_one(i, it, entries, code_index, rules, items, span_len, soil,
                       bridge_width)
            for i, it in enumerate(items)]


# ---------------------------------------------------------------------------
# AI 说明（DeepSeek，仅解释）
# ---------------------------------------------------------------------------

_AI_SYSTEM_PROMPT = (
    "你是天津桥梁定额匹配的说明助手。你只能依据给定的候选定额文本撰写理由，"
    "禁止编造定额编号、单价或定额库中不存在的内容；不得改变候选列表，"
    "不得推荐候选之外的定额。输出一个 JSON 对象："
    '{"candidates":[{"index":1,"reason":"≤40字","diff":"≤40字"}],"overall":"≤60字"}。'
)

_client = None


def _get_client():
    global _client
    if _client is None:
        from openai import OpenAI

        from config import DEEPSEEK_API_KEY, DEEPSEEK_BASE_URL
        _client = OpenAI(api_key=DEEPSEEK_API_KEY, base_url=DEEPSEEK_BASE_URL)
    return _client


def _parse_json_loose(raw):
    raw = re.sub(r"```(?:json)?\s*(.*?)```", r"\1", raw, flags=re.S)
    s = raw[raw.find("{"):raw.rfind("}") + 1] if "{" in raw and "}" in raw else raw
    return json.loads(s)


def ai_explain(item_desc: str, candidates: list[dict]) -> dict | None:
    """多候选 AI 对比说明 → {"candidates":[...],"overall":...}；失败返回 None。"""
    try:
        from config import DEEPSEEK_MODEL
        cand_text = "\n".join(
            f"{i + 1}. {c['entry']['code']} {c['entry']['name']}（{c['entry']['unit']}，"
            f"单价 {c['entry']['price'] if c['entry']['price'] is not None else '—'}）"
            f" 规则命中：{'；'.join(c.get('reasons') or [])}"
            for i, c in enumerate(candidates))
        user = (f"## 清单条目\n{item_desc}\n"
                f"## 候选定额（规则检索结果）\n{cand_text}\n"
                "请输出 JSON 说明各候选的匹配理由与差异。")
        resp = _get_client().chat.completions.create(
            model=DEEPSEEK_MODEL,
            messages=[{"role": "system", "content": _AI_SYSTEM_PROMPT},
                      {"role": "user", "content": user}],
            temperature=0.2, max_tokens=800, timeout=20)
        raw = resp.choices[0].message.content or ""
        try:
            return _parse_json_loose(raw)
        except Exception:
            resp2 = _get_client().chat.completions.create(
                model=DEEPSEEK_MODEL,
                messages=[{"role": "system", "content": _AI_SYSTEM_PROMPT},
                          {"role": "assistant", "content": raw},
                          {"role": "user", "content": "你上次的输出格式不正确，请只输出一个 JSON 对象。"}],
                temperature=0.2, max_tokens=800, timeout=20)
            return _parse_json_loose(resp2.choices[0].message.content or "")
    except Exception:
        return None


# ---------------------------------------------------------------------------
# 结果导出
# ---------------------------------------------------------------------------

_STATUS_CN = {"single": "唯一匹配", "multi": "多候选", "none": "无匹配"}


def _fmt_num(v):
    """数值 → 无科学计数法的短字符串（Excel 公式/追溯标注用）。"""
    if v is None:
        return ""
    try:
        if float(v) == int(v):
            return str(int(v))
        s = f"{float(v):.6f}".rstrip("0").rstrip(".")
        return s or "0"
    except (TypeError, ValueError):
        return str(v)


def _grade_tag(item, entry_name):
    """混凝土强度等级（工程数量表材料列 C25/C30…）→ 定额名称后备注；
    定额名称已含同等级时不重复标注。"""
    m = re.search(r"C\d{2}", _norm_text(item.get("mat") or ""))
    if m and m.group() not in _norm_text(entry_name or ""):
        return f"{m.group()} 混凝土"
    return ""


def result_rows(results) -> list[dict]:
    rows = []
    for r in results:
        item = r["item"]
        if item.get("kind") == "桩基参数":
            mat_label = f"桩基参数 φ{item['dia']:g}/h{item['depth']:g}m×{item['count']:g}根"
        else:
            mat_label = item.get("mat", "")
        base = {"序号": r["idx"] + 1, "部位": item.get("sheet", ""),
                "分项名": item.get("fenxiang", ""),
                "材料": mat_label,
                "清单数量": (_fmt_num(item.get("value")) if item.get("value") is not None
                             else item.get("count", "")),
                "状态": _STATUS_CN.get(r["status"], r["status"])}
        # single：只输出入选定额一行（参考候选仅页面 expander 展示）
        cands = r["matches"][:1] if r["status"] == "single" else r["matches"]
        if cands:
            for c in cands:
                e = c["entry"]
                g = _grade_tag(item, e["name"])
                rows.append({**base,
                             "定额编号": e["code"],
                             "定额名称": e["name"] + (f"（{g}）" if g else ""),
                             "定额单位": e["unit"],
                             "定额工程量": c.get("qty_de", ""),
                             "单价": e["price"] if e["price"] is not None else "",
                             "合价": c.get("amount", ""),
                             "匹配理由": "；".join(c.get("reasons") or [])})
        else:
            ref = "；".join(f"{c['entry']['code']} {c['entry']['name']}"
                            for c in r.get("matches") or [])
            reason = "；".join(r.get("checklist") or [])
            if ref:
                reason += ("；库内较接近子目（仅供参考，不得直接套用）：" + ref) if reason else ref
            rows.append({**base, "定额编号": "", "定额名称": "", "定额单位": "",
                         "定额工程量": "", "单价": "", "合价": "", "匹配理由": reason})
    return rows


def build_result_excel(results, lib=None, waste_haul_km=5, tpl_ratios=None,
                       soil_haul_km=5, earth_reuse=False) -> bytes:
    """导出：
    Sheet1 分部分项工程项目预算计价表——仿表格1013 第二张表格式（两行表头、
    分部/子分部行、子分部内序号重排、单位/工程量/单价/合价/其中人工费）；
    只放最优候选（multi 在名称后标注候选数，无匹配条目不进表）；
    有计算过程的量写成 Excel 公式（工程量/单价换算/合价）。
    Sheet2 候选明细与理由——每候选一行（最优标★），含得分/理由/定额工程量换算。
    Sheet3 补充信息清单——无匹配条目 + 样例组合需补充项。
    lib 可选：传入定额库后，桩基础子分部末尾补 3-10-44 凿除桩顶
    （凿桩头按 0.5m/根计，分部内所有桩合计）+ 1-6-71/1-6-72 弃运
    （建筑垃圾按虚方 ×2，1-6-72 套 (waste_haul_km-1) 次）。
    tpl_ratios：混凝土→模板定额系数覆盖（{系数键: 系数}，缺省用参考表样例比例），
    模板量=混凝土量×系数，公式写入 F 列。
    """
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = openpyxl.Workbook()
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    # ---------------- Sheet1 计价表 ----------------
    ws = wb.active
    ws.title = "分部分项工程项目预算计价表"
    for col, w in (("A", 7), ("B", 11.5), ("C", 19.8), ("D", 7.2), ("E", 12.8),
                   ("F", 12.5), ("G", 7.3), ("H", 7.5), ("I", 16.7), ("J", 13.2)):
        ws.column_dimensions[col].width = w
    ws.merge_cells("A1:J1")
    ws["A1"] = "分部分项工程项目预算计价表"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = center
    ws["A2"] = "专业工程名称:"
    ws["H2"] = "金额单位：元"
    for k, v in (("A3", "序号"), ("B3", "编号"), ("C3", "项 目 名 称"),
                 ("E3", "单位"), ("F3", "工程量"), ("G3", "金  额")):
        ws[k] = v
    ws["G4"], ws["I4"], ws["J4"] = "单价", "合价", "其中：人工费"
    for rng in ("A3:A4", "B3:B4", "C3:D4", "E3:E4", "F3:F4", "G3:H3"):
        ws.merge_cells(rng)
    for cell in ("A3", "B3", "C3", "E3", "F3", "G3", "G4", "I4", "J4"):
        ws[cell].alignment = center
    # 表头两行（序号/编号/项目名称/单位/工程量/金额…）也加边框
    # （用户 2026-10-27 口径；合并单元格逐个描边，防合并后边框缺角）
    for _rr in (3, 4):
        for _cc in "ABCDEFGHIJ":
            ws[f"{_cc}{_rr}"].border = border

    def style_row(row, end="J"):
        for col in "ABCDEFGHIJ"[: "ABCDEFGHIJ".index(end) + 1]:
            ws[f"{col}{row}"].border = border

    def _write_name(text):
        """分部/分项/子分项/子组名称行：C、D 两列合并（用户 2026-10-27 口径）。"""
        nonlocal row
        ws[f"C{row}"] = text
        ws.merge_cells(f"C{row}:D{row}")
        style_row(row)
        row += 1

    # 分部/分项/子分项/子组 层级（对齐单位工程.xlsx 模板结构）：
    # 下部结构（灌注桩/承台/支撑梁/系梁/桥台/桥墩——桥台、桥墩各分
    # 桥台/台帽/台盖梁、桥墩/墩帽/墩盖梁 子分项，子分项下按
    # 成孔/混凝土/钢筋/施工措施/钢绞线 等子组分明细）→ 上部结构 → 附属结构 →
    # 土方（第 4 分部，位于最后）。未上传模板（旧库）时回退旧参考表体系
    # （基础工程→下部→上部→附属，平铺无子组）。
    tpl_divs = (lib or {}).get("tpl_divs") or []
    hier = []
    for d in tpl_divs:
        secs = d.get("sections") or [
            {"name": g, "group": g, "subs": []} for g in d.get("groups", [])]
        hier.append({"div": d["div"], "sections": secs})
    if not hier:
        hier = [
            {"div": "基础工程", "sections": [
                {"name": g, "group": g, "subs": []}
                for g in ("土方", "桩基础", "承台")]},
            {"div": "下部结构", "sections": [
                {"name": g, "group": g, "subs": []}
                for g in ("支撑梁", "系梁", "桥台", "台帽", "台盖梁",
                          "桥墩", "墩帽", "墩盖梁")]},
            {"div": "上部结构", "sections": [
                {"name": g, "group": g, "subs": []}
                for g in ("预应力砼现浇箱梁", "预应力混凝土小箱梁", "预应力板梁",
                          "钢箱梁", "桥面铺装")]},
            {"div": "附属结构", "sections": [
                {"name": g, "group": g, "subs": []}
                for g in ("支座", "伸缩缝", "人行道板", "防撞护栏", "隔声屏障",
                          "搭板", "排水", "沥青铺装", "防水铺装", "防水层", "护坡")]},
        ]
    # 梯道并入上部结构（用户 2026-10-27 口径）：与主梁同构，分「混凝土/钢筋」
    # 两个子组展示；上传模板未含梯道分项时自动补入（内置层级已含）
    for _d in hier:
        if _d["div"] == "上部结构" and not any(
                (s.get("group") or s.get("name")) == "梯道" for s in _d["sections"]):
            _d["sections"].append({"name": "梯道", "group": "梯道",
                                   "subs": ["混凝土", "钢筋"]})
    _GROUP_POS = {}   # 规则组 → (分部序, 分项序, 子分项序)
    _SEC_SUBS = {}    # 规则组 → 子组顺序列表
    for di, d in enumerate(hier):
        for si, s in enumerate(d["sections"]):
            chs = s.get("children") or [{"name": s["name"], "group": s["group"],
                                         "subs": s.get("subs") or []}]
            for ci, ch in enumerate(chs):
                _GROUP_POS[ch["group"]] = (di, si, ci)
                _SEC_SUBS[ch["group"]] = ch.get("subs") or []
    _DIV_RANK = {dv: i for i, dv in enumerate([d["div"] for d in hier] + ["梯道下部"])}
    _EARTH_KWS = ("土方", "挖方", "填方", "锥坡", "护坡", "回填")

    def _div_of(it):
        fx = it.get("fenxiang", "")
        if _norm_text(fx).startswith("梯道下部"):
            return "梯道下部"
        sub = _map_group(fx) or ""
        if sub in _GROUP_POS:
            return hier[_GROUP_POS[sub][0]]["div"]
        # 模板层级未覆盖的土方类条目（挖方/填方/锥坡…）归土方分部
        if tpl_divs and any(k in _norm_text(fx) for k in _EARTH_KWS):
            return "土方"
        return "附属结构"

    # 凿桩头 3-10-44：按 0.5m/根 计，分部内所有桩合计（体积=桩数×0.5×π/4×φ²）
    head_entry = None
    if lib:
        head_entry = next((e for e in lib.get("entries") or []
                           if str(e["code"]) == "3-10-44"), None)
    # 凿出的桩头属建筑垃圾：弃运套 1-6-71（运距1km）+ 1-6-72（每增运1km）。
    # 建筑垃圾（石/混凝土类）弃运量按虚方 ×2；1-6-72 套 (运距-1) 次（运距 _HAUL_KM km）
    waste_entries = []
    if lib:
        for code, nm_kw in (("1-6-71", "运(弃)土、废料 运距1km"),
                            ("1-6-72", "运(弃)土、废料 每增运1km")):
            e = next((x for x in lib.get("entries") or []
                      if str(x["code"]) == code), None)
            if e is None:
                rr = next((x for x in (lib.get("rules") or {}).get("桩基础") or []
                           if str(x["code"]) == code), None)
                e = {"code": code, "name": rr["name"] if rr else nm_kw,
                     "unit": rr["unit"] if rr else "m3", "price": None, "row": 0}
            waste_entries.append(e)
    # 土方规则（用户 2026-10-26 口径，天津市政预算基价系数）：挖方=自然方、
    # 填方=压实方（清单量原样套用）；土方不利用 → 买土 1-6-20=填方量×1.22（自然方）；
    # 余土外运 1-6-71/72=挖方量×1.32（松方），1-6-72 套 (土方运距-1) 次
    def _is_dig_fill(it, kw):
        if it.get("kind") != "材料":
            return False
        t = _norm_text((it.get("fenxiang") or "") + " " + (it.get("mat") or ""))
        return kw in t
    dig_vals = [v for r in results if _is_dig_fill(r["item"], "挖方")
                for v in [_to_float(r["item"].get("value")) or 0] if v]
    fill_vals = [v for r in results if _is_dig_fill(r["item"], "填方")
                 for v in [_to_float(r["item"].get("value")) or 0] if v]
    dig_total, fill_total = sum(dig_vals), sum(fill_vals)
    earth_entries = {}
    if lib:
        for code, nm_kw in (("1-6-20", "买土"),
                            ("1-6-71", "运(弃)土、废料 运距1km"),
                            ("1-6-72", "运(弃)土、废料 每增运1km")):
            e = next((x for x in lib.get("entries") or []
                      if str(x["code"]) == code), None)
            if e is None:
                rr = next((x for x in (lib.get("rules") or {}).get("土方") or []
                           if str(x["code"]) == code), None)
                e = {"code": code, "name": rr["name"] if rr else nm_kw,
                     "unit": rr["unit"] if rr else "m3", "price": None, "row": 0}
            earth_entries[code] = e
    # 模板系数（模板量=混凝土量×系数）：默认参考表样例比例，页面可改
    ratios = dict(TPL_RATIOS)
    if tpl_ratios:
        ratios.update({k: float(v) for k, v in tpl_ratios.items() if v is not None})
    head_by_div = {}
    last_pile_by_div = {}
    if head_entry is not None:
        for r in results:
            it = r["item"]
            if it.get("kind") != "桩基参数":
                continue
            div = _div_of(it)
            h = head_by_div.setdefault(div, {"parts": [], "total": 0.0})
            v = it["count"] * 0.5 * 3.14159265 / 4 * (it["dia"] ** 2)
            h["parts"].append((it["count"], it["dia"]))
            h["total"] += v
            last_pile_by_div[div] = r["idx"]

    # ---- 灌注桩按桩径分组（用户 2026-10-26 口径）----
    # 灌注桩 → φX 灌注桩（子分项）→ 成孔/混凝土/钢筋（子组）：
    # 混凝土工程量 = 3.1415926×r²×桩长×桩数（桩长=工程数量表总长度÷桩数）；
    # 钢筋 = 钢筋总量×本桩径混凝土÷量表总桩基混凝土（每方混凝土含筋量指标，kg→t）。
    # 量表桩基材料行（桥台/桥墩/梯道下部 桩基）只作为分配总量来源，明细行由本分组重算写入。
    def _is_pile_mat_item(it):
        if it.get("kind") != "材料":
            return False
        if "桩基" not in _norm_text(it.get("fenxiang") or ""):
            return False
        mat = _norm_text(it.get("mat") or "")
        return bool(re.match(r"^C\d{2}", mat)) or mat in ("HPB300", "HRB400", "Q235")

    pile_param_rs = [r for r in results if r["item"].get("kind") == "桩基参数"]
    pile_mat_rs = [r for r in results if _is_pile_mat_item(r["item"])]
    shengguan_rs = [r for r in results
                    if "检测管" in _norm_text(r["item"].get("fenxiang") or "")]
    # 泄水管（泄水孔）内按重量计的钢材类：合并一行套预埋铁件 3-2-8
    # （用户 2026-10-26 口径；判定与 match_quotas 内同规则保持同步）
    def _is_drain_steel_kg(it):
        fx = _norm_text(it.get("fenxiang") or "")
        if "泄水" not in fx and "排水" not in fx:
            return False
        if _unit_parts(it.get("mat_unit") or "")[0] != "kg":
            return False
        mat = _norm_text(it.get("mat") or "")
        if _PLASTIC_PIPE_RE.search(mat):
            return False
        return any(k in mat for k in ("钢", "铁", "扣件", "栅盖", "法兰", "管卡"))

    drain_steel_rs = [r for r in results if _is_drain_steel_kg(r["item"])]
    pile_groups = {}
    if pile_param_rs:
        for r in pile_param_rs:
            it = r["item"]
            g = pile_groups.setdefault(float(it["dia"]),
                                       {"count": 0.0, "parts": [],
                                        "param_result": None})
            g["count"] += float(it["count"])
            g["parts"].append((float(it["count"]), float(it["depth"])))
            if g["param_result"] is None:
                g["param_result"] = r
        # 分配总量与定额条目：沿用桩基材料行的匹配结果（缺则按编号找库兜底）
        # 量表每位置矩阵行（桥台桩基/桥墩桩基/梯道下部桩基）→ pos → {材料: 值}，
        # 用户 2026-10-26 口径③：梯道桩基单独取量表原值；主体桩基（桥台+桥墩等
        # 非梯道位置）钢筋合并后按混凝土比例分配
        pos_map = {}
        conc_total = 0.0
        steel_totals = {}
        conc_entry = None
        steel_entries = {}
        conc_grade = ""
        for r in pile_mat_rs:
            it = r["item"]
            mat = _norm_text(it.get("mat") or "")
            v = _to_float(it.get("value")) or 0
            pos_map.setdefault(_norm_text(it.get("fenxiang") or ""), {})[mat] = v
            if re.match(r"^C\d{2}", mat):
                conc_total += v
                if conc_entry is None and r["matches"]:
                    conc_entry = r["matches"][0]["entry"]
                if not conc_grade:
                    m_c = re.search(r"C\d{2}", mat)
                    conc_grade = m_c.group() if m_c else ""
            else:
                steel_totals[mat] = steel_totals.get(mat, 0) + v
                if mat not in steel_entries and r["matches"]:
                    steel_entries[mat] = r["matches"][0]["entry"]
        # 主体桩基合并总量：非梯道位置的混凝土/钢筋求和，供 φ 组按比例分配
        body_pos = [p for p in pos_map if "梯道" not in p]
        body_conc = 0.0
        body_steel = {}
        for p in body_pos:
            for m, v in pos_map[p].items():
                if re.match(r"^C\d{2}", m):
                    body_conc += v
                else:
                    body_steel[m] = body_steel.get(m, 0) + v
        if conc_entry is None:
            conc_entry = next((e for e in lib.get("entries") or []
                               if str(e["code"]) == "3-1-60"), None)
        mud_entries = [next((e for e in lib.get("entries") or []
                             if str(e["code"]) == c), None) for c in ("3-1-58", "3-1-59")]
        cage_entry = next((e for e in lib.get("entries") or []
                           if str(e["code"]) == "3-2-10"), None)
        for dia, g in pile_groups.items():
            r2 = dia / 2
            terms = "+".join(f"{_fmt_num(c)}*{_fmt_num(d)}" for c, d in g["parts"])
            g["conc_expr"] = f"3.1415926*{_fmt_num(r2)}^2*({terms})"
            g["conc_val"] = 3.1415926 * r2 * r2 * sum(c * d for c, d in g["parts"])
            g["head_expr"] = (f"{_fmt_num(g['count'])}*0.5*3.1415926/4*{_fmt_num(dia)}^2")
            # 用本桩径计算混凝土量反查量表位置行（混凝土量精确吻合即同一位置）
            for pos, pm in pos_map.items():
                cc = next((v for m, v in pm.items() if re.match(r"^C\d{2}", m)), None)
                if (cc is not None and g["conc_val"] > 0
                        and abs(cc - g["conc_val"]) <= g["conc_val"] * 1e-5):
                    g["pos_mat"] = pm
                    g["pos_name"] = pos
                    break

    # 模板样例编号 → 子组（供条目归类；模板未出现的编号按名称兜底）
    _CODE_SUB = {}
    for recs in ((lib or {}).get("rules") or {}).values():
        for rr in recs:
            if rr.get("subgroup"):
                _CODE_SUB[str(rr["code"])] = rr["subgroup"]

    def _subgroup_of(entry):
        sg = _CODE_SUB.get(str(entry.get("code") or ""))
        if sg:
            return sg
        nm = _norm_text(entry.get("name") or "")
        code = str(entry.get("code") or "")
        if code.startswith("3-2"):
            return "钢筋"
        if code.startswith("3-5"):
            return "钢绞线"
        if any(k in nm for k in ("护筒", "钻孔", "回旋钻机")):
            return "成孔"
        if any(k in nm for k in ("泥浆", "黏土", "声测管", "凿除桩顶", "混凝土", "模板")):
            return "混凝土"
        if any(k in nm for k in ("支架", "垫层", "拆除")):
            return "施工措施"
        if "四氟板式橡胶支座" in nm:
            return "四氟板式橡胶支座"
        if "板式橡胶支座" in nm:
            return "板式橡胶支座"
        if "细粒" in nm:
            return "细粒式沥青混凝土"
        if "中粒" in nm:
            return "中粒式沥青混凝土"
        if "黏层" in nm or "结合油" in nm:
            return "粘层"
        return None

    buckets = {}   # (分部序, 分项序, 子分项序, 子组序) → [r]
    tail = []      # [(div, sub, r)]：模板层级未覆盖的条目，挂所属分部末尾
    for r in results:
        if r["status"] == "none" or not r["matches"]:
            continue  # 无匹配条目不进计价表（见 补充信息清单）
        it = r["item"]
        # 桩基/声测管条目不进通用分桶，由 _write_pile_section 按桩径分组重算写入
        if pile_groups and (it.get("kind") == "桩基参数" or _is_pile_mat_item(it)
                            or "检测管" in _norm_text(it.get("fenxiang") or "")):
            continue
        # 泄水管钢材类（按重量计）不进通用分桶，由 _write_drain_steel 合并一行写入
        if _is_drain_steel_kg(it):
            continue
        sub = (_group_of(it, (lib or {}).get("rules") or {})
               or _map_group(it.get("fenxiang", ""))
               or (it.get("fenxiang", "") or "其他"))
        # 桥面铺装的钢筋条目并入防水混凝土铺装的「钢筋」子组展示
        # （2026-10-27 用户口径：名称就叫钢筋，不再单列「桥面铺装」；仅影响导出）
        if (sub == "桥面铺装" and "防水铺装" in _GROUP_POS
                and re.search(r"钢筋|CRB|HPB|HRB|冷轧",
                              _norm_text(it.get("mat") or ""))):
            sub = "防水铺装"
        div = _div_of(it)
        if sub in _GROUP_POS and div == hier[_GROUP_POS[sub][0]]["div"]:
            di, si, ci = _GROUP_POS[sub]
            sg = _subgroup_of(r["matches"][0]["entry"])
            subs = _SEC_SUBS[sub]
            sgi = subs.index(sg) if sg in subs else (99 if subs else 0)
            buckets.setdefault((di, si, ci, sgi), []).append(r)
        else:
            tail.append((div, sub, r))
    tail_by_div = {}
    for t in tail:
        tail_by_div.setdefault(t[0], []).append(t)

    def _write_head(div):
        """桩基础子分部收尾：写 3-10-44 凿除桩顶合计行 + 1-6-71/1-6-72 弃运行
        （每分部一次；建筑垃圾弃运按虚方 ×2，1-6-72 套 (运距-1) 次）。"""
        nonlocal row
        h = head_by_div.get(div)
        if head_entry is None or not h or h.get("done"):
            return
        h["done"] = True
        cnt_terms = "+".join(f"{c:g}×φ{d:g}" for c, d in h["parts"])
        vol_terms = "+".join(f"{c:g}*0.5*3.14159265/4*{d:g}^2" for c, d in h["parts"])
        hm = _unit_parts(head_entry["unit"])[1]
        ws[f"B{row}"] = head_entry["code"]
        ws[f"C{row}"] = (head_entry["name"]
                         + f"〔凿桩头按 0.5m/根计，本分部桩合计：{cnt_terms}〕")
        ws.merge_cells(f"C{row}:D{row}")
        ws[f"E{row}"] = head_entry["unit"]
        ws[f"F{row}"] = f"=({vol_terms})/{hm:g}"
        if head_entry["price"] is not None:
            ws[f"G{row}"] = round(head_entry["price"], 4)
            ws[f"I{row}"] = f"=F{row}*G{row}"
        style_row(row)
        row += 1
        # 凿出的桩头属建筑垃圾：1-6-71 弃运1km、1-6-72 每增运1km（弃运量×2 虚方）
        for k, we in enumerate(waste_entries):
            if k == 1 and waste_haul_km <= 1:
                continue  # 运距≤1km 不套每增运1km
            wm = _unit_parts(we["unit"])[1]
            if k == 0:
                wf = f"=({vol_terms})*2"
                note = (f"〔凿桩头弃运×2（建筑垃圾按虚方）；运距 {waste_haul_km:g}km"
                        + ("；单价库中无此子目，需补充" if we["price"] is None else "")
                        + "〕")
            else:
                wf = f"=({vol_terms})*2*{waste_haul_km - 1:g}"
                note = (f"〔每增运1km，套 {waste_haul_km - 1:g} 次"
                        + ("；单价库中无此子目，需补充" if we["price"] is None else "")
                        + "〕")
            if wm != 1:
                wf += f"/{wm:g}"
            ws[f"B{row}"] = we["code"]
            ws[f"C{row}"] = we["name"] + note
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"E{row}"] = we["unit"]
            ws[f"F{row}"] = wf
            if we["price"] is not None:
                ws[f"G{row}"] = round(we["price"], 4)
                ws[f"I{row}"] = f"=F{row}*G{row}"
            style_row(row)
            row += 1

    row = 5

    def _write_item(r):
        """一条匹配结果的明细行组：主行 + 泥浆/钢筋笼/模板/护筒 随同套用行。"""
        nonlocal row
        it = r["item"]
        c = r["matches"][0]
        e = c["entry"]
        mult = _unit_parts(e["unit"])[1]
        ws[f"B{row}"] = e["code"]
        name = e["name"]
        tags = []
        if it.get("syn_tags"):
            tags.extend(it["syn_tags"])
        if r["status"] == "multi" and len(r["matches"]) > 1:
            tags.append(f"另有{len(r['matches']) - 1}个候选")
        # 混凝土强度等级（工程数量表材料列）：定额名称未含该等级时备注在名称后
        g = _grade_tag(it, e["name"])
        if g:
            tags.append(g)
        # 清单原值追溯标注
        if it.get("kind") == "桩基参数":
            tags.append(f"清单 {_fmt_num(it['count'])}根×φ{_fmt_num(it['dia'])}"
                        f"×h{_fmt_num(it['depth'])}m")
        elif it.get("value") is not None:
            tags.append(f"清单 {_fmt_num(it['value'])} {_norm_text(it.get('mat_unit') or '')}".rstrip())
        if c.get("asphalt_conv"):
            ac = c["asphalt_conv"]
            tag = f"清单 m3→t ×{ac['dens']:g}（{ac['grain']}式，定额说明第8条）"
            if ac.get("thick"):
                tag += (f"；设计厚度{ac['thick']:g}m、另加{ac['extra'] * 100:g}cm（第9条）"
                        if ac.get("extra")
                        else f"；设计厚度{ac['thick']:g}m")
            tags.append(tag)
        if c.get("conv_pipe"):
            tags.append(f"按 {c['conv_pipe']['w_kgm']:g} kg/m 换算长度")
        if e.get("price") is None:
            tags.append("单价库中无此子目，需补充")
        # 桩长超库内最大深度档的借用说明（理由标注"最大档"时）
        if any("最大档" in (x or "") for x in c.get("reasons") or []):
            tags.append("桩长超库内最大深度档，按最大档套用")
        if tags:
            name += "〔" + "；".join(tags) + "〕"
        ws[f"C{row}"] = name
        ws.merge_cells(f"C{row}:D{row}")
        ws[f"E{row}"] = e["unit"]  # 单位与定额单位一致
        # 工程量（定额单位口径）：有计算过程的写成公式
        if it.get("f_formula"):  # 按桩径分组计算的合成行（混凝土/钢筋分配/凿桩头…）
            ws[f"F{row}"] = it["f_formula"]
        elif c.get("asphalt_conv"):  # 沥青 m3→t：方量/设计厚度×(厚度+附加)×密度（定额说明第8/9条）
            ac = c["asphalt_conv"]
            if ac.get("thick"):
                t2 = round(ac["thick"] + ac["extra"], 4)
                ws[f"F{row}"] = (f"={_fmt_num(it['value'])}/{ac['thick']:g}*{t2:g}"
                                 f"*{ac['dens']:g}"
                                 + (f"/{mult:g}" if mult != 1 else ""))
            else:
                ws[f"F{row}"] = (f"={_fmt_num(it['value'])}*{ac['dens']:g}"
                                 + (f"/{mult:g}" if mult != 1 else ""))
        elif it.get("kind") == "桩基参数":
            ws[f"F{row}"] = f"={_fmt_num(it['count'])}*{_fmt_num(it['depth'])}/{mult:g}"
        elif c.get("note_conv"):  # kg→t
            ws[f"F{row}"] = (f"={_fmt_num(it['value'])}/1000/{mult:g}" if mult != 1
                             else f"={_fmt_num(it['value'])}/1000")
        elif c.get("conv_pipe"):  # kg→m（钢管 φD×t 按 W kg/m 换算声测管长度）
            cp = c["conv_pipe"]
            ws[f"F{row}"] = (f"={_fmt_num(cp['val'])}/{_fmt_num(cp['w_kgm'])}/{mult:g}"
                             if mult != 1
                             else f"={_fmt_num(cp['val'])}/{_fmt_num(cp['w_kgm'])}")
        elif mult != 1 and it.get("value") is not None:
            ws[f"F{row}"] = f"={_fmt_num(it['value'])}/{mult:g}"
        else:
            v = it.get("value")
            ws[f"F{row}"] = round(v, 4) if v is not None else 0
        # 单价 = 定额原单价
        price = e["price"]
        if price is not None:
            ws[f"G{row}"] = round(price, 4)
            ws[f"I{row}"] = f"=F{row}*G{row}"
        style_row(row)
        row += 1
        # 桩基混凝土条目：泥浆制作/运输随同套用（工程量=桩基混凝土量）
        for me in r.get("mud") or []:
            mmult = _unit_parts(me["unit"])[1]
            ws[f"B{row}"] = me["code"]
            ws[f"C{row}"] = (me["name"]
                             + f"〔工程量=桩基混凝土 {_fmt_num(me['qty_natural'])} m3〕")
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"E{row}"] = me["unit"]
            ws[f"F{row}"] = f"={_fmt_num(me['qty_natural'])}/{mmult:g}"
            if me["price"] is not None:
                ws[f"G{row}"] = round(me["price"], 4)
                ws[f"I{row}"] = f"=F{row}*G{row}"
            style_row(row)
            row += 1
        # 桩基钢筋条目（HPB300 行携带）：钻孔桩钢筋笼入孔 3-2-10，
        # 工程量=同部位 HPB300+HRB400 钢筋量之和（kg→t）
        cage = r.get("cage")
        if cage:
            cmult = _unit_parts(cage["unit"])[1]
            hpb_kg, hrb_kg = cage.get("cage_parts", (0, 0))
            ws[f"B{row}"] = cage["code"]
            ws[f"C{row}"] = (cage["name"]
                             + f"〔工程量=HPB300+HRB400 钢筋量 "
                               f"{_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)} kg〕")
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"E{row}"] = cage["unit"]
            ws[f"F{row}"] = (f"=({_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)})/1000/{cmult:g}"
                             if cmult != 1
                             else f"=({_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)})/1000")
            if cage["price"] is not None:
                ws[f"G{row}"] = round(cage["price"], 4)
                ws[f"I{row}"] = f"=F{row}*G{row}"
            style_row(row)
            row += 1
        # 混凝土条目：相应模板定额随同套用（模板量=混凝土量×系数，默认参考表比例、页面可改）
        tpl = r.get("tpl")
        if tpl and tpl.get("conc_val") is not None:
            tmult = _unit_parts(tpl["unit"])[1]
            ratio = ratios.get(tpl.get("tpl_key"), tpl.get("tpl_default", 0))
            ratio_changed = abs(ratio - tpl.get("tpl_default", 0)) > 1e-9
            ws[f"B{row}"] = tpl["code"]
            ws[f"C{row}"] = (tpl["name"]
                             + f"〔模板系数 {ratio:g}"
                             + (f"（默认 {tpl['tpl_default']:g}）" if ratio_changed else "")
                             + "；模板量=混凝土量×系数〕")
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"E{row}"] = tpl["unit"]
            ws[f"F{row}"] = f"={_fmt_num(tpl['conc_val'])}*{ratio:g}/{tmult:g}"
            if tpl["price"] is not None:
                ws[f"G{row}"] = round(tpl["price"], 4)
                ws[f"I{row}"] = f"=F{row}*G{row}"
            style_row(row)
            row += 1
        # 桩基参数：钻孔行后再套一行陆地埋设钢护筒（每根桩 2m，工程量=桩数×2）
        liner = r.get("pile_liner")
        if it.get("kind") == "桩基参数" and liner:
            lmult = _unit_parts(liner["unit"])[1]
            ws[f"B{row}"] = liner["code"]
            ws[f"C{row}"] = (liner["name"]
                             + f"〔清单 {_fmt_num(it['count'])}根×2m 护筒〕")
            ws.merge_cells(f"C{row}:D{row}")
            ws[f"E{row}"] = liner["unit"]
            ws[f"F{row}"] = f"={_fmt_num(it['count'])}*2/{lmult:g}"
            if liner["price"] is not None:
                ws[f"G{row}"] = round(liner["price"], 4)
                ws[f"I{row}"] = f"=F{row}*G{row}"
            style_row(row)
            row += 1

    def _syn(entry, f_formula, tags, reasons, item_extra=None):
        """分组计算/合并计价的合成行：F 列写公式。"""
        it = {"sheet": "工程数量表", "kind": "材料",
              "fenxiang": "桩基类型", "pos": "", "remark": "",
              "count": None, "mat": "", "mat_unit": entry["unit"],
              "value": None, "row": 0,
              "f_formula": f_formula, "syn_tags": tags}
        if item_extra:
            it.update(item_extra)
        return {"idx": -1, "item": it,
                "status": "single",
                "matches": [{"entry": entry, "score": 40, "unit_ok": True,
                             "reasons": reasons}],
                "rule_combo": [], "checklist": [], "notes": [],
                "mud": [], "cage": None, "tpl": None, "pile_liner": None}

    def _write_pile_section(s):
        """灌注桩分项：分项行 → φX 灌注桩（子分项）→ 成孔/混凝土/钢筋（子组）→ 声测管。"""
        nonlocal row
        _write_name(s["name"])
        for dia in sorted(pile_groups):
            g = pile_groups[dia]
            pr = g["param_result"]
            _write_name(f"φ{dia:g} 灌注桩")
            # 成孔：钻孔 + 护筒（由桩基参数匹配结果携带，工程量=桩数×桩长）
            if pr is not None and pr["matches"]:
                _write_name("成孔")
                _write_item(pr)
            # 混凝土：灌注桩混凝土（3.1415926×r²×桩长×桩数）+ 泥浆制作/运输
            # + 凿桩头（0.5m/根，本桩径各算各的）+ 弃运（×2 虚方）
            if conc_entry is not None:
                _write_name("混凝土")
                cmult = _unit_parts(conc_entry["unit"])[1]
                _write_item(_syn(conc_entry, f"={g['conc_expr']}/{cmult:g}",
                                 [f"{conc_grade} 混凝土；工程量按 3.1415926×r²×桩长×桩数"
                                  f"={_fmt_num(g['conc_val'])} m³ 计算"
                                  + (f"（=量表{g['pos_name']}原值）" if g.get("pos_name") else "")],
                                 ["按 3.1415926×r²×桩长×桩数 计算本桩径混凝土工程量"]))
                for me in mud_entries:
                    if me is None:
                        continue
                    mmult = _unit_parts(me["unit"])[1]
                    _write_item(_syn(me, f"={g['conc_expr']}/{mmult:g}",
                                     ["工程量=本桩径桩基混凝土"],
                                     ["泥浆制作/运输随灌注桩混凝土套用，工程量=桩基混凝土量"]))
                if head_entry is not None:
                    hm = _unit_parts(head_entry["unit"])[1]
                    _write_item(_syn(head_entry, f"={g['head_expr']}/{hm:g}",
                                     [f"凿桩头按 0.5m/根计，本桩径 {_fmt_num(g['count'])}根"],
                                     ["凿桩头按 0.5m/根计，本桩径各算各的"]))
                    for k, we in enumerate(waste_entries):
                        if k == 1 and waste_haul_km <= 1:
                            continue  # 运距≤1km 不套每增运1km
                        wm = _unit_parts(we["unit"])[1]
                        if k == 0:
                            wf = f"={g['head_expr']}*2"
                            wtag = f"凿桩头弃运×2（建筑垃圾按虚方）；运距 {waste_haul_km:g}km"
                            wreason = "凿桩头弃运：建筑垃圾按虚方×2"
                        else:
                            wf = f"={g['head_expr']}*2*{waste_haul_km - 1:g}"
                            wtag = f"每增运1km，套 {waste_haul_km - 1:g} 次"
                            wreason = f"每增运1km：套 {waste_haul_km - 1:g} 次"
                        if wm != 1:
                            wf += f"/{wm:g}"
                        _write_item(_syn(we, wf, [wtag], [wreason]))
            # 钢筋：梯道桩基单独取量表原值（用户 2026-10-26 口径③）；主体桩基
            # （桥台+桥墩等非梯道位置）钢筋合并后按混凝土量比例分配
            # （=主体钢总量×本桩径混凝土÷主体总混凝土）+ 钢筋笼入孔
            if conc_total > 0 and steel_entries:
                _write_name("钢筋")
                pm = g.get("pos_mat")
                is_tidao = "梯道" in (g.get("pos_name") or "")
                for mat in ("HPB300", "HRB400", "Q235"):
                    ent = steel_entries.get(mat)
                    stot = steel_totals.get(mat, 0)
                    if ent is None or not stot:
                        continue
                    if is_tidao and pm and (pm.get(mat) or 0) > 0:
                        pv = pm[mat]
                        ind2 = round(pv / g["conc_val"], 2)
                        _write_item(_syn(
                            ent, f"={_fmt_num(pv)}/1000",
                            [f"{mat}；量表{g['pos_name']}原值（梯道单独计算，指标 {_fmt_num(ind2)} kg/m³）"],
                            [f"梯道桩基 {mat} 单独取量表原值，不参与主体分配"]))
                    elif body_conc > 0:
                        btot = body_steel.get(mat, 0)
                        ind = round(btot / body_conc, 2)
                        _write_item(_syn(
                            ent, f"={_fmt_num(btot)}*{g['conc_expr']}/{_fmt_num(body_conc)}/1000",
                            [f"{mat}；主体桩基合并按混凝土比例分配（主体指标 {_fmt_num(ind)} kg/m³）"],
                            [f"主体桩基 {mat} 合并按混凝土比例分配：主体总量×本桩径混凝土÷主体总混凝土"]))
                    else:
                        ind = stot / conc_total
                        _write_item(_syn(
                            ent, f"={_fmt_num(stot)}*{g['conc_expr']}/{_fmt_num(conc_total)}/1000",
                            [f"{mat}；按混凝土量比例分配（指标 {_fmt_num(ind)} kg/m³）"],
                            [f"{mat} 按混凝土量比例分配：总量×本桩径混凝土÷总桩基混凝土"]))
                if cage_entry is not None:
                    if is_tidao and pm is not None and (pm.get("HPB300") or pm.get("HRB400")):
                        _write_item(_syn(
                            cage_entry,
                            f"=({_fmt_num(pm.get('HPB300') or 0)}"
                            f"+{_fmt_num(pm.get('HRB400') or 0)})/1000",
                            ["工程量=本桩径 HPB300+HRB400 量表原值之和"],
                            ["钢筋笼入孔工程量=HPB300+HRB400 钢筋量之和（量表原值）"]))
                    elif body_conc > 0 and (body_steel.get("HPB300") or body_steel.get("HRB400")):
                        _write_item(_syn(
                            cage_entry,
                            f"=({_fmt_num(body_steel.get('HPB300') or 0)}"
                            f"+{_fmt_num(body_steel.get('HRB400') or 0)})*{g['conc_expr']}"
                            f"/{_fmt_num(body_conc)}/1000",
                            ["工程量=本桩径 HPB300+HRB400 主体分配量之和"],
                            ["钢筋笼入孔工程量=HPB300+HRB400 主体分配量之和"]))
                    else:
                        hpb_t = steel_totals.get("HPB300", 0)
                        hrb_t = steel_totals.get("HRB400", 0)
                        if hpb_t or hrb_t:
                            _write_item(_syn(
                                cage_entry,
                                f"=({_fmt_num(hpb_t)}+{_fmt_num(hrb_t)})*{g['conc_expr']}"
                                f"/{_fmt_num(conc_total)}/1000",
                                ["工程量=本桩径 HPB300+HRB400 分配量之和"],
                                ["钢筋笼入孔工程量=HPB300+HRB400 钢筋量之和（按分配量）"]))
        # 声测管：量表总量横跨全部桩，不按桩径拆分，挂灌注桩分项末尾
        if shengguan_rs:
            _write_name("声测管")
            for r in shengguan_rs:
                _write_item(r)

    def _write_drain_steel(s):
        """泄水孔分项：钢材类（按重量计）合并一行套预埋铁件 3-2-8（用户 2026-10-26 口径）。"""
        nonlocal row
        _write_name(s["name"])
        entry = None
        parts = []
        total = 0.0
        for r in drain_steel_rs:
            it = r["item"]
            v = _to_float(it.get("value")) or 0
            parts.append(f"{it.get('mat')} {_fmt_num(v)}kg")
            total += v
            if entry is None and r["matches"]:
                entry = r["matches"][0]["entry"]
        if entry is None:
            entry = next((e for e in lib.get("entries") or []
                          if str(e["code"]) == "3-2-8"), None)
        if entry is not None and total > 0:
            mult = _unit_parts(entry["unit"])[1]
            wf = f"=({'+'.join(_fmt_num(_to_float(r['item'].get('value')) or 0)
                               for r in drain_steel_rs)})/1000"
            if mult != 1:
                wf += f"/{mult:g}"
            _write_item(_syn(
                entry, wf,
                [f"泄水管钢材类合并计价：{' + '.join(parts)} = {_fmt_num(total)} kg"],
                ["泄水管钢材类按重量合计，统一套预埋铁件 3-2-8（用户口径）"],
                {"sheet": "附属", "fenxiang": "泄水孔"}))

    _earth_written = {"done": False}

    def _write_earth_derived():
        """土方分项收尾（用户 2026-10-26 口径）：
        土方不利用 → 买土 1-6-20 = 填方（压实方）×1.22（自然方）；
        余土外运 1-6-71/72 = 挖方（自然方）×1.32（松方），1-6-72 套 (土方运距-1) 次。"""
        nonlocal row
        if _earth_written["done"]:
            return
        _earth_written["done"] = True
        if not earth_reuse and fill_total > 0:
            e = earth_entries.get("1-6-20")
            if e is not None:
                mult = _unit_parts(e["unit"])[1]
                f = f"=({'+'.join(_fmt_num(v) for v in fill_vals)})*1.22"
                if mult != 1:
                    f += f"/{mult:g}"
                _write_item(_syn(
                    e, f,
                    [f"买土（自然方）=填方（压实方）×1.22（土方不利用，外购土）"],
                    [f"买土定额：填方 {_fmt_num(fill_total)} m3 压实方 → 自然方 ×1.22"
                     "（用户口径，天津市政预算基价系数）"],
                    {"sheet": "工程数量表", "fenxiang": "土方",
                     "mat": "买土", "mat_unit": "m3"}))
        if dig_total > 0:
            dig_terms = "+".join(_fmt_num(v) for v in dig_vals)
            if earth_reuse and fill_total > 0:
                # 土方利用：余土（自然方）= 挖方 − 填方×1.22（填方折回自然方），再 ×1.32 → 松方
                fill_terms = "+".join(_fmt_num(v) for v in fill_vals)
                base = f"(({dig_terms})-({fill_terms})*1.22)*1.32"
                base_tag = ("余土外运（松方）=(挖方−填方×1.22)×1.32"
                            "（土方利用，挖方扣除回填利用后外运）；运距1km")
                base_reason = (f"余土外运：挖方 {_fmt_num(dig_total)} m3 自然方，扣除回填利用 "
                               f"{_fmt_num(fill_total)}×1.22 m3 后 ×1.32 → 松方"
                               "（用户口径，天津市政预算基价道路工程系数）")
            else:
                base = f"({dig_terms})*1.32"
                base_tag = "余土外运（松方）=挖方（自然方）×1.32；运距1km"
                base_reason = (f"余土外运：挖方 {_fmt_num(dig_total)} m3 自然方 → 松方 ×1.32"
                               "（土方不利用，挖方全部外运；用户口径）")
            for k, code in enumerate(("1-6-71", "1-6-72")):
                if k == 1 and soil_haul_km <= 1:
                    continue  # 运距≤1km 不套每增运1km
                e = earth_entries.get(code)
                if e is None:
                    continue
                mult = _unit_parts(e["unit"])[1]
                f = f"={base}"
                if k == 1:
                    f += f"*{soil_haul_km - 1:g}"
                if mult != 1:
                    f += f"/{mult:g}"
                tag = (base_tag if k == 0
                       else f"每增运1km×{soil_haul_km - 1:g} 次（运距 {soil_haul_km:g}km）")
                _write_item(_syn(e, f, [tag], [base_reason],
                                 {"sheet": "工程数量表", "fenxiang": "土方",
                                  "mat": "余土外运", "mat_unit": "m3"}))

    # 按 分部→分项→子分项→子组 层级顺序写入（严格对齐模板排布）
    # 表内只有桩基条目时（无其他分项进 buckets/tail），桩基础所在分部也要写
    pile_div = (hier[_GROUP_POS["桩基础"][0]]["div"]
                if pile_groups and "桩基础" in _GROUP_POS else None)
    for di, d in enumerate(hier):
        div = d["div"]
        ts = tail_by_div.get(div) or []
        if not any(k[0] == di for k in buckets) and not ts and not (div == pile_div):
            continue
        _write_name(div)
        for si, s in enumerate(d["sections"]):
            chs = s.get("children") or [{"name": s["name"], "group": s["group"],
                                         "subs": s.get("subs") or []}]
            if s.get("group") == "排水" and drain_steel_rs:
                _write_drain_steel(s)   # 泄水孔钢材类合并一行套预埋铁件
                continue
            if s.get("group") == "桩基础" and pile_groups:
                _write_pile_section(s)   # 灌注桩按桩径分组写入（成孔/混凝土/钢筋 + 声测管）
                continue
            if not any(k[:2] == (di, si) for k in buckets):
                continue
            _write_name(s["name"])      # 分项行（灌注桩/承台/现浇箱梁/桥台/桥墩…）
            for ci, ch in enumerate(chs):
                if not any(k[:3] == (di, si, ci) for k in buckets):
                    continue
                if s.get("children"):
                    _write_name(ch["name"])  # 子分项行（桥台/台帽/台盖梁/墩帽…）
                subs = ch.get("subs") or []
                for sgi, sgn in enumerate(subs):
                    rs = buckets.get((di, si, ci, sgi))
                    if not rs:
                        continue
                    _write_name(sgn)          # 子组行（成孔/混凝土/钢筋/施工措施…）
                    for r in rs:
                        _write_item(r)
                    # 灌注桩/混凝土 子组收尾：凿除桩顶合计 + 弃运（与模板同位置）
                    if ch["group"] == "桩基础" and sgn == "混凝土":
                        _write_head(div)
                rs = buckets.get((di, si, ci, 0))
                if rs and not subs:              # 无子组的分项（钢箱梁/伸缩缝/土方…）
                    for r in rs:
                        _write_item(r)
                    if ch["group"] == "土方":
                        _write_earth_derived()   # 买土/余土外运（×1.22 / ×1.32）
                rs = buckets.get((di, si, ci, 99))
                if rs:                           # 子组未覆盖条目挂在分项末尾
                    for r in rs:
                        _write_item(r)
                if ch["group"] == "桩基础":
                    _write_head(div)             # 兜底：子组缺失时也在分项收尾补
        # 分部末尾：模板层级未覆盖的条目（按子分部聚组，保持结果顺序）
        last_tsub = None
        for _, tsub, r in ts:
            if tsub != last_tsub:
                _write_name(tsub)
                last_tsub = tsub
            _write_item(r)
            if tsub == "土方":
                _write_earth_derived()   # 兜底：土方条目走 tail 时也补买土/余土外运
    # 模板之外的分部：梯道下部按「桥墩/承台 → 混凝土/钢筋」子组写入
    # （用户 2026-10-26 口径，与下部结构模板同构）；其余分部平铺。
    # 分部收尾凿桩头/弃运。
    _TD_SECTIONS = ("桥墩", "承台")
    _TD_SUBS = ("混凝土", "钢筋")
    for div, ts in tail_by_div.items():
        if div in [d["div"] for d in hier]:
            continue
        _write_name(div)
        if div == "梯道下部":
            by_sec, leftover = {}, []
            for _, tsub, r in ts:
                if tsub in _TD_SECTIONS:
                    by_sec.setdefault(tsub, []).append(r)
                else:
                    leftover.append((tsub, r))
            for sec in _TD_SECTIONS:
                rs = by_sec.get(sec)
                if not rs:
                    continue
                _write_name(sec)                 # 分项行（桥墩/承台）
                for sgn in _TD_SUBS:
                    srs = [r for r in rs
                           if _subgroup_of(r["matches"][0]["entry"]) == sgn]
                    if not srs:
                        continue
                    _write_name(sgn)             # 子组行（混凝土/钢筋）
                    for r in srs:
                        _write_item(r)
                rest = [r for r in rs
                        if _subgroup_of(r["matches"][0]["entry"]) not in _TD_SUBS]
                for r in rest:                      # 子组归类不了的挂分项末尾
                    _write_item(r)
            last_tsub = None
            for tsub, r in leftover:
                if tsub != last_tsub:
                    _write_name(tsub)
                    last_tsub = tsub
                _write_item(r)
        else:
            last_tsub = None
            for _, tsub, r in ts:
                if tsub != last_tsub:
                    _write_name(tsub)
                    last_tsub = tsub
                _write_item(r)
                if tsub == "土方":
                    _write_earth_derived()   # 兜底：土方条目走 tail 时也补买土/余土外运
        _write_head(div)
    # 序号：子分部内从 1 重排（B 列有编号的是明细行）
    seq = 0
    for rr in range(5, row):
        if ws[f"B{rr}"].value:
            seq += 1
            ws[f"A{rr}"] = seq
        else:
            seq = 0

    # ---------------- Sheet2 候选明细与理由 ----------------
    ws2 = wb.create_sheet("候选明细与理由")
    headers2 = ["序号", "部位", "分项名", "材料", "清单数量", "状态", "候选",
                "定额编号", "定额名称", "定额单位", "定额工程量", "定额单价",
                "合价", "得分", "匹配理由"]
    ws2.append(headers2)
    for col, w in (("A", 6), ("B", 12), ("C", 14), ("D", 22), ("E", 10),
                   ("F", 9), ("G", 8), ("H", 11), ("I", 40), ("J", 9),
                   ("K", 12), ("L", 10), ("M", 12), ("N", 6), ("O", 50)):
        ws2.column_dimensions[col].width = w
    r2 = 2
    for r in results:
        it = r["item"]
        if it.get("kind") == "桩基参数":
            mat_label = f"桩基参数 φ{it['dia']:g}/h{it['depth']:g}m×{it['count']:g}根"
        else:
            mat_label = it.get("mat", "")
        qty_label = (_fmt_num(it.get("value")) if it.get("value") is not None
                     else _fmt_num(it.get("count")) if it.get("count") is not None else "")
        base_row = [r["idx"] + 1, it.get("sheet", ""), it.get("fenxiang", ""),
                    mat_label, qty_label,
                    _STATUS_CN.get(r["status"], r["status"])]
        if r["matches"]:
            for j, c in enumerate(r["matches"]):
                e = c["entry"]
                mult = _unit_parts(e["unit"])[1]
                ws2.append(base_row + ["★最优" if j == 0 else "",
                                       e["code"], e["name"], e["unit"],
                                       None, e["price"], None,
                                       c.get("score", ""),
                                       "；".join(c.get("reasons") or [])])
                # 定额工程量（扩大单位）：有计算过程的写成公式
                if it.get("kind") == "桩基参数":
                    ws2[f"K{r2}"] = f"={_fmt_num(it['count'])}*{_fmt_num(it['depth'])}/{mult:g}"
                elif c.get("note_conv"):
                    ws2[f"K{r2}"] = (f"={_fmt_num(it['value'])}/1000/{mult:g}" if mult != 1
                                     else f"={_fmt_num(it['value'])}/1000")
                else:
                    ws2[f"K{r2}"] = (f"={_fmt_num(it['value'])}/{mult:g}" if mult != 1
                                     else round(it["value"], 4) if it.get("value") is not None else 0)
                if e["price"] is not None:
                    ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                r2 += 1
            liner = r.get("pile_liner")
            if it.get("kind") == "桩基参数" and liner:
                lmult = _unit_parts(liner["unit"])[1]
                ws2.append(base_row + ["★套用",
                                       liner["code"], liner["name"], liner["unit"],
                                       None, liner["price"], None, "",
                                       "陆地埋设钢护筒：每根桩 2m，工程量=桩数×2m"])
                ws2[f"K{r2}"] = f"={_fmt_num(it['count'])}*2/{lmult:g}"
                if liner["price"] is not None:
                    ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                r2 += 1
            if head_entry is not None and it.get("kind") == "桩基参数":
                g = pile_groups.get(float(it["dia"])) if pile_groups else None
                if g is not None and g.get("param_result") is r:
                    # 按桩径分组：凿桩头/弃运本桩径各算各的（与 Sheet1 同口径）
                    hm = _unit_parts(head_entry["unit"])[1]
                    ws2.append(base_row + ["★套用",
                                           head_entry["code"], head_entry["name"],
                                           head_entry["unit"], None, head_entry["price"],
                                           None, "",
                                           f"凿桩头按 0.5m/根计，本桩径 {_fmt_num(g['count'])}根"])
                    ws2[f"K{r2}"] = f"={g['head_expr']}/{hm:g}"
                    if head_entry["price"] is not None:
                        ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                    r2 += 1
                    # 凿出的桩头属建筑垃圾：弃运 1-6-71 ×2、1-6-72 ×2×(运距-1)
                    for k, we in enumerate(waste_entries):
                        if k == 1 and waste_haul_km <= 1:
                            continue
                        wm = _unit_parts(we["unit"])[1]
                        if k == 0:
                            wf = f"={g['head_expr']}*2"
                            reason = (f"凿桩头弃运：×2（建筑垃圾按虚方），运距 {waste_haul_km:g}km"
                                      + ("；单价库中无此子目，需补充"
                                         if we["price"] is None else ""))
                        else:
                            wf = f"={g['head_expr']}*2*{waste_haul_km - 1:g}"
                            reason = (f"每增运1km：套 {waste_haul_km - 1:g} 次"
                                      f"（运距 {waste_haul_km:g}km）"
                                      + ("；单价库中无此子目，需补充"
                                         if we["price"] is None else ""))
                        if wm != 1:
                            wf += f"/{wm:g}"
                        ws2.append(base_row + ["★套用",
                                               we["code"], we["name"], we["unit"],
                                               None, we["price"], None, "", reason])
                        ws2[f"K{r2}"] = wf
                        if we["price"] is not None:
                            ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                        r2 += 1
                elif (not pile_groups
                      and r["idx"] == last_pile_by_div.get(_div_of(it))):
                    # 无桩基类型块时保持旧口径：分部内所有桩合计套用一次
                    h = head_by_div.get(_div_of(it))
                    if h and h.get("parts"):
                        cnt_terms = "+".join(f"{c:g}×φ{d:g}" for c, d in h["parts"])
                        vol_terms = "+".join(f"{c:g}*0.5*3.14159265/4*{d:g}^2"
                                             for c, d in h["parts"])
                        hm = _unit_parts(head_entry["unit"])[1]
                        ws2.append(base_row + ["★套用",
                                               head_entry["code"], head_entry["name"],
                                               head_entry["unit"], None, head_entry["price"],
                                               None, "",
                                               "凿桩头按 0.5m/根计，本分部所有桩合计套用一次"
                                               f"（{cnt_terms}）"])
                        ws2[f"K{r2}"] = f"=({vol_terms})/{hm:g}"
                        if head_entry["price"] is not None:
                            ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                        r2 += 1
                        for k, we in enumerate(waste_entries):
                            if k == 1 and waste_haul_km <= 1:
                                continue
                            wm = _unit_parts(we["unit"])[1]
                            if k == 0:
                                wf = f"=({vol_terms})*2"
                                reason = (f"凿桩头弃运：×2（建筑垃圾按虚方），运距 {waste_haul_km:g}km"
                                          + ("；单价库中无此子目，需补充"
                                             if we["price"] is None else ""))
                            else:
                                wf = f"=({vol_terms})*2*{waste_haul_km - 1:g}"
                                reason = (f"每增运1km：套 {waste_haul_km - 1:g} 次"
                                          f"（运距 {waste_haul_km:g}km）"
                                          + ("；单价库中无此子目，需补充"
                                             if we["price"] is None else ""))
                            if wm != 1:
                                wf += f"/{wm:g}"
                            ws2.append(base_row + ["★套用",
                                                   we["code"], we["name"], we["unit"],
                                                   None, we["price"], None, "", reason])
                            ws2[f"K{r2}"] = wf
                            if we["price"] is not None:
                                ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                            r2 += 1
            for me in r.get("mud") or []:
                mmult = _unit_parts(me["unit"])[1]
                ws2.append(base_row + ["★套用",
                                       me["code"], me["name"], me["unit"],
                                       None, me["price"], None, "",
                                       "随灌注桩混凝土套用：工程量=桩基混凝土量（÷10 换算）"])
                ws2[f"K{r2}"] = f"={_fmt_num(me['qty_natural'])}/{mmult:g}"
                if me["price"] is not None:
                    ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                r2 += 1
            cage = r.get("cage")
            if cage:
                cmult = _unit_parts(cage["unit"])[1]
                hpb_kg, hrb_kg = cage.get("cage_parts", (0, 0))
                ws2.append(base_row + ["★套用",
                                       cage["code"], cage["name"], cage["unit"],
                                       None, cage["price"], None, "",
                                       "钢筋笼入孔随桩基钢筋套用：工程量=HPB300+HRB400 钢筋量之和（÷1000 换 t）"])
                ws2[f"K{r2}"] = (f"=({_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)})/1000/{cmult:g}"
                                 if cmult != 1
                                 else f"=({_fmt_num(hpb_kg)}+{_fmt_num(hrb_kg)})/1000")
                if cage["price"] is not None:
                    ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                r2 += 1
            tpl = r.get("tpl")
            if tpl and tpl.get("conc_val") is not None:
                tmult = _unit_parts(tpl["unit"])[1]
                ratio = ratios.get(tpl.get("tpl_key"), tpl.get("tpl_default", 0))
                ratio_changed = abs(ratio - tpl.get("tpl_default", 0)) > 1e-9
                ws2.append(base_row + ["★套用",
                                       tpl["code"], tpl["name"], tpl["unit"],
                                       None, tpl["price"], None, "",
                                       "模板随混凝土套用：模板量=混凝土量×系数 "
                                       f"{ratio:g}"
                                       + (f"（默认 {tpl['tpl_default']:g}，页面可改）"
                                          if ratio_changed else "（页面可改）")])
                ws2[f"K{r2}"] = f"={_fmt_num(tpl['conc_val'])}*{ratio:g}/{tmult:g}"
                if tpl["price"] is not None:
                    ws2[f"M{r2}"] = f"=K{r2}*L{r2}"
                r2 += 1
        else:
            ref = "；".join(f"{c['entry']['code']} {c['entry']['name']}"
                            for c in r.get("matches") or [])
            reason = "；".join(r.get("checklist") or [])
            if ref:
                reason += ("；库内较接近子目（仅供参考，不得直接套用）：" + ref) if reason else ref
            ws2.append(base_row + ["", "", "", "", "", "", "", "", reason])
            r2 += 1

    # ---------------- Sheet3 补充信息清单 ----------------
    ws3 = wb.create_sheet("补充信息清单")
    ws3.append(["序号", "部位", "分项名", "材料", "缺失维度/补充说明"])
    for col, w in (("A", 6), ("B", 12), ("C", 14), ("D", 24), ("E", 60)):
        ws3.column_dimensions[col].width = w
    for r in results:
        if r["status"] == "none":
            ws3.append([r["idx"] + 1, r["item"].get("sheet", ""),
                        r["item"].get("fenxiang", ""), r["item"].get("mat", ""),
                        "；".join(r.get("checklist") or [])])
        for rr in r.get("rule_combo") or []:
            if rr.get("status") == "需补充" and rr.get("need"):
                ws3.append([r["idx"] + 1, r["item"].get("sheet", ""),
                            r["item"].get("fenxiang", ""), rr["code"] + " " + rr["name"],
                            rr["need"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
