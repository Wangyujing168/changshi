"""
植物表规格写入名称 — 把 胸（地）径/高度/冠幅 拼接进名称列（B 列）。

规则（2026-10-09 用户确认）：
- 乔木 → 「胸径Xcm」；灌木 → 「地径Xcm」；拿不准 → 「胸（地）径Xcm」
- C 列已标「地径」的照标（如 白皮松「地径8」→ 地径8cm）
- 有哪项写哪项（缺高度/冠幅只写有的），三项全空保留原名
- 只处理 C 列表头含「径」的苗木区块；景石（规格长X宽X高）等区块不动
"""
import io
import re
from pathlib import Path

import openpyxl

# 大乔木（未标地径时写 胸径）
_TREE_KEYS = ["油松", "白皮松", "华山松", "云杉", "桧柏", "侧柏", "雪松", "乔松",
              "法桐", "悬铃木", "白蜡", "栾树", "银杏", "国槐", "刺槐", "杨", "柳",
              "椿", "枫", "槭", "榆", "朴", "皂荚", "合欢", "玉兰", "樟", "泡桐",
              "楸", "梓", "水杉", "落羽杉", "楝", "乌桕", "重阳木", "七叶树", "女贞"]
# 灌木/小乔木（按地径计）
_SHRUB_KEYS = ["木槿", "紫薇", "榆叶梅", "山桃", "山杏", "杏梅", "黄栌", "迎春",
               "连翘", "红瑞木", "接骨木", "月季", "卫矛", "小檗", "紫荆", "海棠",
               "丁香", "珍珠梅", "绣线菊", "蔷薇", "玫瑰", "黄刺玫", "棣棠", "猥实",
               "锦带", "金银木", "忍冬", "荚蒾", "海桐", "火棘", "沙地柏", "地柏",
               "冬青", "黄杨", "委陵菜", "鸢尾", "麦冬", "鼠尾草", "萱草", "金鸡菊",
               "玉簪", "景天", "石竹", "石榴", "紫叶李", "樱花", "碧桃", "桂花"]


def _fmt(v) -> str:
    """数值格式化：18.0 → '18'，0.8 → '0.8'，字符串原样。"""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _jing_part(name: str, c_val) -> str | None:
    """径字段：'胸径14-16cm' / '地径12cm' / '胸（地）径Xcm'，无值返回 None。"""
    if c_val is None:
        return None
    s = str(c_val).strip()
    if not s:
        return None
    if "地径" in s:  # 表里已标地径 → 照标
        return f"地径{_fmt(s.replace('地径', ''))}cm"
    if any(k in name for k in _SHRUB_KEYS):
        tag = "地径"
    elif any(k in name for k in _TREE_KEYS):
        tag = "胸径"
    else:
        tag = "胸（地）径"
    return f"{tag}{_fmt(s)}cm"


def _spec_part(label: str, val) -> str | None:
    if val is None or str(val).strip() == "":
        return None
    return f"{label}{_fmt(val)}m"


def rewrite_name(name: str, c_val, d_val, e_val) -> str:
    """拼接新名称；三项全空返回原名称。"""
    parts = [_jing_part(name, c_val),
             _spec_part("高度", d_val),
             _spec_part("冠幅", e_val)]
    parts = [p for p in parts if p]
    if not parts:
        return name
    return name + " " + " ".join(parts)


def _is_row_num(v) -> bool:
    if isinstance(v, (int, float)):
        return True
    s = str(v).strip() if v is not None else ""
    return bool(re.fullmatch(r"\d+", s))


def _rewrite_sheet(ws) -> list[dict]:
    changes = []
    spec_mode = False
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        a, b, c = row[0].value, row[1].value, row[2].value
        # 区块表头行：A=序号 B=名称 → 根据 C 列表头是否含「径」决定是否处理该区块
        if (a is not None and str(a).strip() == "序号"
                and b is not None and str(b).strip() == "名称"):
            spec_mode = bool(c and "径" in str(c))
            continue
        if not spec_mode or b is None:
            continue
        if not _is_row_num(a):
            continue
        old = str(b).strip()
        new = rewrite_name(old, c, row[3].value, row[4].value)
        if new != old:
            row[1].value = new
            changes.append({"行号": row[0].row, "原名称": old, "新名称": new})
    return changes


def process_plant_table(file_or_path):
    """加载植物表并改写所有工作表的名称列，返回 (changes, wb)。

    file_or_path: 文件路径（str/Path）或 Streamlit UploadedFile。
    """
    if isinstance(file_or_path, (str, Path)):
        wb = openpyxl.load_workbook(str(file_or_path))
    else:
        wb = openpyxl.load_workbook(io.BytesIO(file_or_path.getvalue()))
    changes = []
    for ws in wb.worksheets:
        changes += _rewrite_sheet(ws)
    return changes, wb


def to_bytes(wb) -> bytes:
    """工作簿另存为 bytes（供 st.download_button）。"""
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
